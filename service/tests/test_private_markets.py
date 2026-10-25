"""D200: the private markets minimums.

The mandate each private markets product needs, by allocation type (Full,
ex-HFs, ex-RA, ex-HFs ex-RA), risk level and base currency, in millions of
dollars. An allocation holding private markets is offered at a risk level and
in a currency only from the lowest minimum any product has there; above it,
a product below its own minimum stops the download, as a position below its
product's minimum does. A custom portfolio is checked product by product at
its profile's Full. Loaded from the desk's long-form file or edited on
screen; every change a revision, with what moved. Before the first table, the
$20m rule stands.
"""

import base64
import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from cyrus_pmg.pmgService import dashboardRouter
from cyrus_pmg.pmgService.scenario import (customPortfolio, privateMarkets as pm, products,
                                           repositoryFeed, rules, scenarioStore, sleeveRepo)
from cyrus_pmg.pmgService.scenario.types import (BasisInput, MandateInput, PortfolioKey,
                                                 ValidationError)

MA = 'PMG Multi-Asset Portfolio'
USD = BasisInput(currency='USD', hedging='Hedged')
APS = 'gs-alternatives-portfolio-services'
BUYOUT = 'gs-private-markets-buyout'
#: the desk's APS figures under Full (CHF, EUR, GBP, USD)
APS_FULL = {'Low Vol': [96, 111, 101, 84], 'Conservative': [59, 65, 48, 44],
            'Conservative-Moderate': [47, 41, 36, 34], 'Moderate': [35, 29, 27, 26],
            'Moderate-Aggressive': [28, 24, 22, 22], 'Aggressive': [22, 20, 18, 18],
            'Higher Risk': [18, 17, 16, 15]}


@pytest.fixture()
def store(tmp_path, monkeypatch):
    monkeypatch.setenv('SCENARIO_SLEEVES_DB', str(tmp_path / 'sleeves.db'))
    monkeypatch.setenv('SCENARIO_REGISTER_DB', str(tmp_path / 'proposals.db'))
    monkeypatch.setenv('SCENARIO_STORE_DIR', str(tmp_path / 'store'))
    monkeypatch.setenv('PMG_ALLOWED_KERBEROS', 'alice,bob')
    products.reload()
    sleeveRepo.describe()
    yield tmp_path
    products.reload()


def _caller(kerberos='alice'):
    from cyrus_pmg.pmgService.core import accessControl
    return accessControl.UserData(kerberos, accessControl._rolesFor(kerberos))


def _rows(pid=APS, label='APS', scale=1.0, skip=()):
    out = []
    for a in pm.ALLOCATIONS:
        for risk, values in APS_FULL.items():
            for cur, v in zip(('CHF', 'EUR', 'GBP', 'USD'), values):
                if (a, risk, cur) in skip:
                    continue
                out.append([a, risk, cur, round(v * scale, 2) if a == 'Full' else round(v * scale * 1.1, 2),
                            label, pid])
    return out


def _csv(rows, head=('allocation_type', 'risk_level', 'currency', 'value', 'Type', 'productId')):
    text = '\n'.join(','.join(str(x) for x in r) for r in [list(head)] + rows)
    return base64.b64encode(text.encode()).decode()


def _xlsx(rows):
    book = Workbook()
    book.active.title = 'Consolidated'
    book.active.append(['a summary, not the table'])
    sheet = book.create_sheet('Long Form')
    sheet.append(['allocation_type', 'risk_level', 'currency', 'value', 'Type', 'productId'])
    for r in rows:
        sheet.append(r)
    book.create_sheet('Variables').append(['allocation_type', 'risk_level', 'currency', 'productId'])
    buf = io.BytesIO()
    book.save(buf)
    return base64.b64encode(buf.getvalue()).decode()


def _load(rows, base=0, name='mins.csv'):
    content = _xlsx(rows) if name.endswith('.xlsx') else _csv(rows)
    return pm.importFile(name, content, user='alice', base=base)


def _scenario(mandate=1e9):
    from test_scenario_backend import BASIS, PORT, _confirmHedgeFunds, _sleeveMap
    key = PortfolioKey('USD', 'Moderate', 'Full', False)
    state = scenarioStore.createScenario(
        MandateInput(topAccountSize=2e9, mandateSize=mandate, primaryPwa='A. Castellanos — Madrid'),
        BASIS, createdBy='alice')
    result = PORT.resolve_portfolio(BASIS, key)
    scenarioStore.updateScenario(state['id'], variant=MA)
    scenarioStore.recordColumn(state['id'], key, 'base')
    scenarioStore.updateScenario(state['id'], sleeves=_sleeveMap(result['categories'], MA))
    _confirmHedgeFunds(state['id'])
    return state['id']


# ------------------------------------------------------- before a table ---

def test_before_a_table_the_twenty_million_rule_stands(store):
    assert pm.current() == {'products': [], 'revision': 0, 'inForce': False, 'updatedBy': '',
                            'updatedAt': ''}
    assert rules.PRIVATE_ASSETS_MINIMUM == pm.FALLBACK_MINIMUM == 20_000_000
    assert pm.gate('USD|Moderate|Full|0') == {'minimum': 20e6, 'gap': False}
    assert pm.gate('USD|Moderate|Core|0') is None and pm.gate('USD|All Equity|NA|NA') is None
    offered = rules.availability(USD, 15e6, MA)
    assert 'USD|Moderate|Core|0' in offered and 'USD|Moderate|Full|0' not in offered
    assert 'USD|Moderate|Full|0' in rules.availability(USD, 20e6, MA)
    # the allocation is still listed, so the page can grey it out with what it needs
    assert 'Full' in rules.allocationsFor(15e6, MA)
    with pytest.raises(ValidationError) as refused:
        rules.validateKey(PortfolioKey('USD', 'Moderate', 'Full', False), MA, 15e6)
    assert refused.value.field == 'mandateSize'
    assert 'Full needs a mandate of $20,000,000 in USD Moderate' in refused.value.message
    assert customPortfolio.forCaller(True)['privateMinimum'] == 20_000_000


# --------------------------------------------------------------- the file ---

def test_the_desks_workbook_is_read_from_its_long_form_sheet(store):
    preview = pm.preview('private-market-mins.xlsx', _xlsx(_rows()))
    assert preview['rows'] == 112 and preview['problems'] == [] and preview['gaps'] == []
    assert preview['products'] == [{'productId': APS, 'type': 'APS', 'values': 112}]
    assert preview['unknown'] == [APS] and preview['changed'] == 1 and preview['base'] == 0
    done = _load(_rows(), name='private-market-mins.xlsx')
    assert done['revision'] == 1 and done['inForce'] and done['updatedBy'] == 'alice'
    p = done['products'][0]
    assert p['minimums']['Full|LowVol|EUR'] == 111 and p['minimums']['Full|Moderate|USD'] == 26
    assert p['minimums']['ex-HFs ex-RA|Higher Risk|USD'] == 16.5
    assert pm.history()[0]['action'] == 'imported' and pm.history()[0]['note'] == 'Loaded private-market-mins.xlsx'


def test_all_covers_every_one_and_the_row_naming_more_wins(store):
    rows = [['All', 'Moderate', 'All', 40, 'BUY', BUYOUT], ['Full', 'Moderate', 'USD', 30, 'BUY', BUYOUT],
            ['ex-HFs ex-RAs', 'moderate', 'gbp', 45, '', BUYOUT]]
    read = pm.readFile('m.csv', _csv(rows))
    assert read['problems'] == [] and read['unknown'] == []
    got = read['table']['products'][0]['minimums']
    assert len(got) == 16 and got['Full|Moderate|USD'] == 30 and got['ex-HFs ex-RA|Moderate|GBP'] == 45
    assert got['ex-RA|Moderate|CHF'] == 40
    assert len(read['gaps'][0]['cells']) == 112 - 16


@pytest.mark.parametrize('rows, words', [
    ([['Full', 'Moderate', 'USD', 30, '', APS], ['Full', 'Moderate', 'USD', 31, '', APS]],
     'Rows 2 and 3 both give'),
    ([['Fully', 'Moderate', 'USD', 30, '', APS]], 'allocation_type "Fully" is not one the table takes'),
    ([['Full', 'Medium', 'USD', 30, '', APS]], 'risk_level "Medium" is not one'),
    ([['Full', 'Moderate', 'JPY', 30, '', APS]], 'currency "JPY" is not one'),
    ([['Full', 'Moderate', 'USD', 'lots', '', APS]], '"lots" is not a number of millions of dollars'),
    ([['Full', 'Moderate', 'USD', -1, '', APS]], 'between 0 and'),
    ([['Full', 'Moderate', 'USD', 30, '', '']], 'name the product by its ProductId'),
])
def test_a_file_with_a_problem_is_shown_and_refused(store, rows, words):
    preview = pm.preview('m.csv', _csv(rows))
    assert any(words in p for p in preview['problems']), preview['problems']
    with pytest.raises(ValidationError) as refused:
        pm.importFile('m.csv', _csv(rows), user='alice', base=0)
    assert words in refused.value.message and pm.current()['revision'] == 0


def test_a_file_missing_a_column_or_not_a_file_is_refused(store):
    with pytest.raises(ValidationError, match='missing column'):
        pm.readFile('m.csv', _csv([['Full', 'Moderate', 'USD', APS]],
                                  head=('allocation_type', 'risk_level', 'currency', 'productId')))
    with pytest.raises(ValidationError, match='No sheet in m.xlsx has the columns'):
        book = Workbook()
        book.active.append(['nothing'])
        buf = io.BytesIO()
        book.save(buf)
        pm.readFile('m.xlsx', base64.b64encode(buf.getvalue()).decode())
    with pytest.raises(ValidationError, match='a .csv or an .xlsx'):
        pm.readFile('m.txt', _csv([]))


def test_the_download_is_the_desks_long_form_and_loads_back_unchanged(store):
    _load(_rows(skip={('Full', 'Moderate', 'GBP')}))
    rows = list(load_workbook(io.BytesIO(pm.exportWorkbook())).worksheets[0].iter_rows(values_only=True))
    assert rows[0] == pm.COLUMNS and len(rows) == 1 + 111
    assert ('Full', 'Low Vol', 'CHF', 96, 'APS', APS) in rows
    again = pm.preview('back.xlsx', base64.b64encode(pm.exportWorkbook()).decode())
    assert again['changed'] == 0 and again['problems'] == []


# ------------------------------------------------------------- the gates ---

def test_an_allocation_opens_from_the_lowest_minimum_of_any_product(store):
    _load(_rows() + _rows(BUYOUT, 'BUY', scale=2.0))
    assert pm.gate('USD|Moderate|Full|0') == {'minimum': 26e6, 'gap': False}
    assert pm.gate('USD|Moderate|Full|1') == {'minimum': 28.6e6, 'gap': False}     # ex-RA
    assert pm.gate('GBP|LowVol|ex-HFs|0')['minimum'] == 111.1e6
    offered = rules.availability(USD, 30e6, MA)
    assert 'USD|Moderate|Full|0' in offered and 'USD|Conservative|Full|0' not in offered
    assert 'USD|Conservative|Core|0' in offered and 'USD|Higher Risk|Full|0' in offered
    with pytest.raises(ValidationError, match=r'Full needs a mandate of \$44,000,000 in USD Conservative'):
        rules.validateKey(PortfolioKey('USD', 'Conservative', 'Full', False), MA, 30e6)
    served = rules.schemaPayload(USD, 30e6, {}, {}, MA)['rules']['privateMarkets']
    assert served['inForce'] and served['gates']['USD|Moderate|Full|0'] == 26e6
    assert served['minimums'][BUYOUT]['USD|Moderate|Full|0'] == 52e6
    assert 'USD|Moderate|Core|0' not in served['gates'] and served['types'][BUYOUT] == 'BUY'
    assert customPortfolio.forCaller(True)['privateMinimum'] is None


def test_a_gap_is_never_read_as_no_minimum(store):
    _load(_rows(skip={('Full', 'Moderate', 'USD')}))
    assert pm.gate('USD|Moderate|Full|0') == {'minimum': None, 'gap': True}
    assert 'USD|Moderate|Full|0' not in rules.availability(USD, 1e9, MA)
    assert pm.closedWords('USD|Moderate|Full|0') == 'No private markets minimum is set for USD Moderate Full yet.'
    assert pm.gaps()[0]['cells'] == ['Full|Moderate|USD']
    assert rules.schemaPayload(USD, 1e9, {}, {}, MA)['rules']['privateMarkets']['gates']['USD|Moderate|Full|0'] is None


def test_each_product_below_its_minimum_or_without_one_is_a_shortfall(store):
    _load(_rows() + _rows(BUYOUT, 'BUY', scale=2.0, skip={('Full', 'Aggressive', 'USD')}))
    groups = [{'category': 'Private Equity', 'items': [
        {'productId': APS, 'name': 'APS', 'printedPct': 3.0},
        {'productId': BUYOUT, 'name': 'Buyout', 'printedPct': 0.0, 'initialPct': 2.0},
        {'productId': 'gs-not-in-the-table', 'name': 'Other', 'printedPct': 5.0}]}]
    at = pm.cellOf('USD|Moderate|Full|0')
    assert [(s['name'], s['minimum']) for s in pm.shortfalls(groups, at, 30e6)] == [('Buyout', 52e6)]
    assert pm.shortfalls(groups, at, 60e6) == []
    assert pm.shortfalls(groups, pm.cellOf('USD|Moderate|Core|0'), 1e6) == []
    gap = pm.shortfalls(groups, pm.cellOf('USD|Agg|Full|0'), 1e9)
    assert [(s['name'], s['gap']) for s in gap] == [('Buyout', True)]
    assert 'no minimum is set for Full · Aggressive · USD' in pm.shortfallText(gap, pm.cellOf('USD|Agg|Full|0'), 1e9)
    assert pm.customCell('USD', 'Moderate') == 'Full|Moderate|USD' and pm.customCell('USD', 'All Equity') is None


def test_the_download_refuses_a_product_below_its_minimum(store):
    scenarioId = _scenario()
    first = dashboardRouter.exportScenario(scenarioId, caller=_caller())
    assert first.status_code == 200, first.body
    _load(_rows() + _rows(BUYOUT, 'BUY', scale=50.0))              # Buyout needs $1.3bn here
    refused = dashboardRouter.exportScenario(scenarioId, caller=_caller())
    said = json.loads(refused.body)
    assert refused.status_code == 422 and said['field'] == 'privateMarkets'
    assert said['error'].startswith('Below the private markets minimum for Full · Moderate · USD at a $1,000,000,000 mandate: ')
    assert '(needs a $1,300,000,000 mandate)' in said['error']
    _load(_rows() + _rows(BUYOUT, 'BUY', scale=1.0), base=1)        # now $26m
    assert dashboardRouter.exportScenario(scenarioId, caller=_caller()).status_code == 200


# ------------------------------------------------------- editing, record ---

def test_an_edit_on_screen_is_a_revision_with_each_value_that_moved(store):
    _load(_rows())
    table = pm.current()
    edited = json.loads(json.dumps(table['products']))
    edited[0]['minimums']['Full|Moderate|USD'] = 28
    del edited[0]['minimums']['Full|Moderate|GBP']
    edited[0]['type'] = 'APS2'
    edited.append({'productId': BUYOUT, 'type': 'BUY', 'minimums': {'Full|Moderate|USD': 40}})
    body = dashboardRouter.savePrivateMarketMinimums({'products': edited, 'note': 'desk review', 'base': 1},
                                                     caller=_caller('bob'))
    now = body['minimums']
    assert now['revision'] == 2 and now['updatedBy'] == 'bob' and len(now['products']) == 2
    assert [g['productId'] for g in now['gaps']] == [APS, BUYOUT]
    top = dashboardRouter.getPrivateMarketMinimumsHistory(caller=_caller())['history'][0]
    assert top['note'] == 'desk review' and top['current'] and top['actor'] == 'bob'
    assert top['changes'] == [f'{APS}: short name APS → APS2', 'APS2 · Full · Moderate · USD: 26 → 28',
                              'APS2 · Full · Moderate · GBP: 27 → —', 'Added BUY (1 values)']
    same = dashboardRouter.savePrivateMarketMinimums({'products': edited, 'base': 2}, caller=_caller())
    assert same.status_code == 422 and 'already in force' in json.loads(same.body)['error']
    stale = dashboardRouter.savePrivateMarketMinimums({'products': [], 'base': 1}, caller=_caller())
    assert stale.status_code == 409
    back = dashboardRouter.revertPrivateMarketMinimums({'revision': 1, 'base': 2}, caller=_caller())
    assert back['minimums']['revision'] == 3 and back['minimums']['products'] == table['products']
    assert pm.history()[0]['action'] == 'reverted' and pm.history()[0]['note'] == 'Restored revision 1'
    feed = [e for e in repositoryFeed.activity(limit=50)['entries'] if e['area'] == 'minimums']
    assert [e['revision'] for e in feed] == [3, 2, 1] and feed[1]['note'] == 'desk review'


@pytest.mark.parametrize('products, words', [
    ([{'productId': APS, 'minimums': {'Full|Moderate|USD': 'x'}}], 'not a number of millions'),
    ([{'productId': APS, 'minimums': {'Full|Moderate|JPY': 1}}], 'is not a cell of the table'),
    ([{'productId': 'two words', 'minimums': {}}], 'is not a ProductId'),
    ([{'productId': APS, 'minimums': {}}, {'productId': APS, 'minimums': {}}], 'is in the table twice'),
])
def test_a_table_that_is_not_one_is_refused(store, products, words):
    with pytest.raises(ValidationError, match=words):
        pm.save({'products': products}, user='alice', base=0)


def test_the_console_routes_preview_load_and_download(store):
    preview = dashboardRouter.previewPrivateMarketMinimums({'name': 'm.csv', 'content': _csv(_rows())},
                                                           caller=_caller())
    assert preview['changed'] == 1
    done = dashboardRouter.importPrivateMarketMinimums({'name': 'm.csv', 'content': _csv(_rows()), 'base': 0},
                                                       caller=_caller())
    assert done['minimums']['revision'] == 1
    assert dashboardRouter.getPrivateMarketMinimums(caller=_caller())['cells'] == 112
    sheet = dashboardRouter.exportPrivateMarketMinimums(caller=_caller())
    assert sheet.headers['content-disposition'] == 'attachment; filename="private-market-mins.xlsx"'
    bad = dashboardRouter.importPrivateMarketMinimums({'name': 'm.csv', 'content': _csv([['Full', 'x', 'USD', 1, '', APS]]),
                                                       'base': 1}, caller=_caller())
    assert bad.status_code == 422 and json.loads(bad.body)['field'] == 'file'
