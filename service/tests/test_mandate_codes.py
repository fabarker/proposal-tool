"""D205, D211, D215: the mandate codes - the code the account opening
request names for the mandate's own row of its SMA Details, for each
portfolio a PWA can be offered. Its benchmark is the portfolio benchmarks'
(D211).

A portfolio is an implementation type, a currency, a risk level, an
allocation (with its form without real assets, all equity and custom) and a
hedging policy. The codes are given by group (D215): a name, a code and the
portfolios it holds - All where it chooses none - in order, a portfolio
taking the code of the first group that holds it. Until the first save the
groups are the desk's two: PMG ESG takes G3, everything else 7L. A
portfolio no group holds is a gap.
"""

import base64
import csv
import io

import pytest

from cyrus_pmg.pmgService import dashboardRouter
from cyrus_pmg.pmgService.scenario import mandateCodes, repositoryFeed, sleeveRepo
from cyrus_pmg.pmgService.scenario.types import ValidationError
from test_proposal_register import _caller

MA = 'PMG Multi-Asset Portfolio'
ESG = 'PMG ESG'


@pytest.fixture()
def store(tmp_path, monkeypatch):
    monkeypatch.setenv('SCENARIO_SLEEVES_DB', str(tmp_path / 'sleeves.db'))
    monkeypatch.setenv('SCENARIO_REGISTER_DB', str(tmp_path / 'proposals.db'))
    monkeypatch.setenv('SCENARIO_STORE_DIR', str(tmp_path / 'store'))
    sleeveRepo.describe()
    yield tmp_path


def _upload(tmp_path, rows, name='codes.csv'):
    path = tmp_path / name
    with open(path, 'w', newline='', encoding='utf-8') as fh:
        csv.writer(fh).writerows(rows)
    return base64.b64encode(path.read_bytes()).decode('ascii')


def _group(name, code, **chosen):
    return dict({'name': name, 'code': code}, **{d: chosen.get(d, []) for d in mandateCodes.GROUP_DIMENSIONS})


HEAD = list(mandateCodes.COLUMNS)
#: the file before groups (D205): a row a portfolio, any of its five All
BEFORE = ['ImplementationType', 'Currency', 'RiskLevel', 'AllocationType', 'Hedging', 'MandateCode']


def at(*key):
    return mandateCodes.codeFor(mandateCodes.cell(*key))


def test_the_portfolios_on_offer_are_each_types_own(store):
    forms = {t['type']: [f['form'] for f in t['forms']] for t in mandateCodes.layout()}
    assert forms[MA] == ['Full', 'Full ex-RA', 'Core', 'ex-Alts', 'ex-HFs', 'ex-HFs ex-RA', 'All Equity', 'Custom']
    assert forms[ESG] == ['ex-Alts', 'ex-HFs ex-RA', 'All Equity', 'Custom'], 'ESG leaves real assets out'
    assert forms['US Onshore'] == ['ex-Alts', 'All Equity', 'Custom']
    assert mandateCodes.forms() == forms[MA]
    cells = mandateCodes.cells()
    assert len(cells) == len(set(cells)) == (51 + 23 + 16 + 16) * 4 * 4
    assert mandateCodes.cell(MA, 'Full', 'Moderate', 'USD', 'Hedged') in cells
    assert mandateCodes.cell(ESG, 'ex-HFs', 'Moderate', 'USD', 'Hedged') not in cells
    # a proposal finds its cell by its key, or a custom portfolio by its profile
    assert mandateCodes.keyOf(MA, 'USD|Moderate|Full|1', hedging='ISG Hedged') == \
        'PMG Multi-Asset Portfolio|Full ex-RA|Moderate|USD|ISG Hedged'
    assert mandateCodes.keyOf(MA, 'GBP|All Equity|NA|NA', hedging='Unhedged') == \
        'PMG Multi-Asset Portfolio|All Equity|All Equity|GBP|Unhedged'
    assert mandateCodes.keyOf(ESG, profile='ModAgg', currency='CHF', hedging='Hedged') == \
        'PMG ESG|Custom|ModAgg|CHF|Hedged'


def test_before_a_save_pmg_esg_takes_g3_and_everything_else_7l(store):
    now = mandateCodes.current()
    assert now['revision'] == 0
    assert [(g['name'], g['code'], g['types']) for g in now['groups']] == [
        ('PMG ESG', 'G3', [ESG]), ('Everything else', '7L', [])]
    esg = sum(1 for c in mandateCodes.cells() if c.startswith(ESG + '|'))
    assert now['given'] == [{'gives': esg, 'holds': esg},
                            {'gives': len(mandateCodes.cells()) - esg, 'holds': len(mandateCodes.cells())}]
    assert at(ESG, 'ex-Alts', 'Moderate', 'USD', 'Hedged') == 'G3'
    assert at(ESG, 'Custom', 'All Equity', 'EUR', 'Equity Not Hedged') == 'G3'
    assert at(MA, 'Full', 'Moderate', 'USD', 'Hedged') == '7L'
    assert at('US Onshore', 'Custom', 'LowVol', 'CHF', 'Hedged') == '7L'
    assert mandateCodes.gaps() == {'noCode': []}
    assert mandateCodes.history() == [], 'a read never writes'
    assert mandateCodes.exportRows() == [HEAD, ['PMG ESG', 'G3', ESG, 'All', 'All', 'All', 'All'],
                                         ['Everything else', '7L', 'All', 'All', 'All', 'All', 'All']]


def test_the_first_group_that_holds_a_portfolio_gives_its_code(store):
    groups = [_group('GBP Core', 'MA-GBP-CORE', types=[MA], forms=['Core', 'Core'], currencies=['GBP']),
              _group('PMG ESG', 'G3', types=[ESG]),
              _group('Everything else', '7L')]
    saved = mandateCodes.save({'groups': groups}, 'Split out GBP Core', user='alice', base=0)
    assert saved['revision'] == 1
    assert saved['groups'][0]['forms'] == ['Core'], 'one of each, in the console\'s order'
    assert at(MA, 'Core', 'Moderate', 'GBP', 'Hedged') == 'MA-GBP-CORE'
    assert at(MA, 'Core', 'Moderate', 'USD', 'Hedged') == '7L'
    assert at(ESG, 'ex-Alts', 'Moderate', 'GBP', 'Hedged') == 'G3'
    # every one of a five ticked is All
    every = mandateCodes.options()['currencies']
    out = mandateCodes.normalise({'groups': [_group('A', 'X', currencies=list(reversed(every)))]})
    assert out['groups'][0]['currencies'] == []
    words = mandateCodes.history()[0]['changes']
    assert words[0] == ('Group GBP Core added: MA-GBP-CORE · PMG Multi-Asset Portfolio · Core · All risk levels · '
                        'GBP · All hedging policies')
    assert words[-1].endswith('portfolios take another code')
    # without a last group of All, what no group holds is a gap
    mandateCodes.save({'groups': groups[:2]}, user='alice', base=1)
    gaps = mandateCodes.gaps()['noCode']
    assert gaps and not any(c.startswith(ESG + '|') for c in gaps)
    assert mandateCodes.codeFor(mandateCodes.cell(MA, 'Full', 'Moderate', 'USD', 'Hedged')) is None


@pytest.mark.parametrize('groups, field, words', [
    ([_group('', 'X')], 'group0name', 'needs a name'),
    ([_group('A', 'X', types=[ESG]), _group('a', 'Y')], 'group1name', 'both called a'),
    ([_group('A', '')], 'group0code', 'needs a mandate code'),
    ([_group('A', 'X' * 41)], 'group0code', 'keep the code to 40 characters'),
    ([_group('A', 'X', currencies=['JPY'])], 'group0currencies', "'JPY' is not one of the currencies"),
    ([_group('All of it', 'X'), _group('PMG ESG', 'G3', types=[ESG])], 'group1name',
     'taken by a group above it, All of it'),
    ([_group('ESG Full', 'X', types=[ESG], forms=['Full'])], 'group0name', 'holds no portfolio on offer'),
])
def test_a_group_that_cannot_be_is_refused_on_its_box(store, groups, field, words):
    with pytest.raises(ValidationError, match=words) as refused:
        mandateCodes.save({'groups': groups}, user='alice')
    assert refused.value.field == field
    assert mandateCodes.current()['revision'] == 0


def test_a_file_a_row_a_group_in_order(store, tmp_path):
    rows = [HEAD,
            ['ESG hedged', 'G3-H', 'PMG ESG', '', 'All', 'USD, GBP', 'Hedged; ISG Hedged'],
            ['PMG ESG', 'G3', 'PMG ESG', 'All', 'All', 'All', 'All'],
            ['Multi-Asset Full', 'MA-F', 'PMG Multi', 'Full, Full · ex Real Assets', 'Moderate-Aggressive, Low Vol', '', ''],
            ['Everything else', '7L', 'All', 'All', 'All', 'All', 'All']]
    preview = mandateCodes.preview('codes.csv', _upload(tmp_path, rows))
    assert preview['problems'] == [] and preview['rows'] == 4 and preview['groups'] == 4 and not preview['before']
    assert preview['noCode'] == 0 and preview['cells'] == len(mandateCodes.cells())
    mandateCodes.importFile('codes.csv', _upload(tmp_path, rows), user='alice', base=preview['base'])
    groups = mandateCodes.current()['groups']
    assert [g['name'] for g in groups] == ['ESG hedged', 'PMG ESG', 'Multi-Asset Full', 'Everything else']
    assert groups[0]['currencies'] == ['USD', 'GBP'] and groups[0]['hedgings'] == ['Hedged', 'ISG Hedged']
    assert groups[2]['forms'] == ['Full', 'Full ex-RA'] and groups[2]['risks'] == ['LowVol', 'ModAgg']
    assert at(ESG, 'ex-Alts', 'Moderate', 'GBP', 'ISG Hedged') == 'G3-H'
    assert at(ESG, 'ex-Alts', 'Moderate', 'EUR', 'Hedged') == 'G3'
    assert at(MA, 'Full ex-RA', 'LowVol', 'CHF', 'Unhedged') == 'MA-F'
    assert at(MA, 'Full', 'Moderate', 'CHF', 'Unhedged') == '7L'
    top = mandateCodes.history()[0]
    assert top['action'] == 'imported' and top['note'] == 'Loaded codes.csv' and top['current']
    # the export loads back with nothing to do
    out = mandateCodes.exportRows()
    assert out[0] == HEAD and out[3] == ['Multi-Asset Full', 'MA-F', MA, 'Full, Full ex-RA', 'Low Vol, Moderate-Aggressive',
                                         'All', 'All']
    again = mandateCodes.preview('again.csv', _upload(tmp_path, out, 'again.csv'))
    assert again['problems'] == [] and again['changed'] == 0
    feed = [e for e in repositoryFeed.activity(limit=50)['entries'] if e['area'] == 'mandateCodes']
    assert [e['name'] for e in feed] == ['Mandate codes']


@pytest.mark.parametrize('row, words', [
    (['A', '', 'All', 'All', 'All', 'All', 'All'], 'no MandateCode'),
    (['', 'X', 'All', 'All', 'All', 'All', 'All'], 'no Group name'),
    (['A', 'X', 'All', 'All', 'All', 'JPY', 'All'], 'Currency "JPY" is not one the table takes'),
    (['A', 'X', 'Somewhere', 'All', 'All', 'All', 'All'], 'ImplementationType "Somewhere"'),
    (['A', 'X', ESG, 'Full', 'All', 'All', 'All'], 'holds no portfolio on offer'),
    (['A', 'X' * 41, 'All', 'All', 'All', 'All', 'All'], 'keep the code to 40 characters'),
])
def test_a_bad_row_is_a_problem_and_nothing_loads(store, tmp_path, row, words):
    preview = mandateCodes.preview('codes.csv', _upload(tmp_path, [HEAD, row]))
    assert any(words in p for p in preview['problems']), preview['problems']
    with pytest.raises(ValidationError):
        mandateCodes.importFile('codes.csv', _upload(tmp_path, [HEAD, row]), user='alice')
    assert mandateCodes.current()['revision'] == 0


def test_a_group_below_one_that_takes_all_it_holds_is_a_problem(store, tmp_path):
    rows = [HEAD, ['Everything', '7L', 'All', 'All', 'All', 'All', 'All'],
            ['PMG ESG', 'G3', ESG, 'All', 'All', 'All', 'All'], ['Everything', '7M', 'All', 'All', 'All', 'All', 'All']]
    problems = mandateCodes.preview('codes.csv', _upload(tmp_path, rows))['problems']
    assert 'Rows 2 and 4 are both the group Everything.' in problems
    assert any(p.startswith('Row 3: Every portfolio PMG ESG holds is taken by a group above it, Everything') for p in problems)


def test_a_file_from_before_groups_reads_a_row_a_group_the_row_naming_more_first(store, tmp_path):
    """D205's file: All for every one, the row naming more winning - each
    row a group, the narrower ones first; a Benchmark column is not read
    (D211)."""
    rows = [BEFORE + ['Benchmark'],
            ['All', 'All', 'All', 'All', 'All', 'PMG-ANY', ''],
            [MA, 'USD', 'All', 'Full', 'All', 'PMG-MA-USD-FULL', ''],
            [MA, 'usd', 'Moderate-Aggressive', 'full', 'ISG Hedged', 'PMG-MA-USD-MA-FULL-I', 'Not read'],
            [ESG, 'GBP', 'All Equity', '', 'All', 'ESG-GBP-AE', ''],
            ['PMG Multi', 'EUR', 'Conservative', 'ex-RA', 'Hedged', 'PMG-MA-EUR-CON-FRA', '']]
    preview = mandateCodes.preview('codes.csv', _upload(tmp_path, rows))
    assert preview['problems'] == [] and preview['before'] and preview['benchmarkColumn'] and preview['groups'] == 5
    mandateCodes.importFile('codes.csv', _upload(tmp_path, rows), user='alice', base=preview['base'])
    assert at(MA, 'Full', 'Moderate', 'USD', 'Hedged') == 'PMG-MA-USD-FULL'
    assert at(MA, 'Full', 'ModAgg', 'USD', 'ISG Hedged') == 'PMG-MA-USD-MA-FULL-I'
    assert at(MA, 'Full', 'ModAgg', 'USD', 'Hedged') == 'PMG-MA-USD-FULL'
    assert at(MA, 'Full ex-RA', 'Conservative', 'EUR', 'Hedged') == 'PMG-MA-EUR-CON-FRA'
    assert at(ESG, 'All Equity', 'All Equity', 'GBP', 'Unhedged') == 'ESG-GBP-AE'
    assert at('US Onshore', 'Custom', 'LowVol', 'CHF', 'Hedged') == 'PMG-ANY'
    groups = mandateCodes.current()['groups']
    assert groups[-1]['name'] == 'Every portfolio'
    assert groups[0]['name'] == 'PMG Multi-Asset Portfolio · Full · Moderate-Aggressive · USD · ISG Hedged'
    # two rows naming as much that give one portfolio two codes are named
    clash = [BEFORE, [MA, 'USD', 'All', 'Core', 'Hedged', 'A'], [MA, 'USD', 'Moderate', 'All', 'Hedged', 'B'],
             [MA, 'GBP', 'Moderate', 'Core', 'Hedged', 'C'], [MA, 'GBP', 'Moderate', 'Core', 'Hedged', 'C']]
    problems = mandateCodes.preview('codes.csv', _upload(tmp_path, clash))['problems']
    assert problems == ['Rows 2 and 3 both give PMG Multi-Asset Portfolio · Core · Moderate · USD · Hedged: codes A and B.']


def test_a_revision_from_before_groups_reads_a_group_a_portfolio(store):
    """A table saved before groups (D205) - each portfolio's code, an older
    one [code, benchmark] (D211) - reads as groups giving each portfolio its
    code and no other, those of one code folded together, and its history
    from nothing."""
    c1 = mandateCodes.cell(MA, 'Core', 'Moderate', 'USD', 'Hedged')
    c2 = mandateCodes.cell(ESG, 'Custom', 'All Equity', 'EUR', 'Equity Not Hedged')
    conn = sleeveRepo.connection()
    with conn:
        body = '{"codes":{"%s":["KEPT","Old benchmark"],"%s":"ESG-CU","nope|x|y|z|w":"GONE"}}' % (c2, c1)
        conn.execute("INSERT INTO policies (kind, scope, body, revision, updatedBy, updatedAt) "
                     "VALUES ('mandateCodes', '*', ?, 1, 'alice', '2026-10-24T10:00:00')", (body,))
        conn.execute("INSERT INTO policyHistory (kind, scope, revision, action, body, note, actor, at) "
                     "VALUES ('mandateCodes', '*', 1, 'updated', ?, '', 'alice', '2026-10-24T10:00:00')", (body,))
    conn.close()
    now = mandateCodes.current()
    assert [(g['name'], g['code'], g['hedgings']) for g in now['groups']] == [
        ('ESG-CU', 'ESG-CU', ['Hedged']), ('KEPT', 'KEPT', ['Equity Not Hedged'])]
    assert mandateCodes.codeFor(c2) == 'KEPT' and mandateCodes.codeFor(c1) == 'ESG-CU'
    assert len(mandateCodes.gaps()['noCode']) == len(mandateCodes.cells()) - 2
    assert mandateCodes.history()[0]['changes'][:1] == [
        'Group ESG-CU added: ESG-CU · PMG Multi-Asset Portfolio · Core · Moderate · USD · Hedged']
    # one code across the table folds into a few groups, each portfolio still its code
    codes = {c: 'ONE' for c in mandateCodes.cells()}
    codes[c1] = 'TWO'
    folded = mandateCodes._fromCodes(codes)
    assert len(folded) < 20 and [g['name'] for g in folded].count('TWO') == 1
    assert mandateCodes._codesOf(folded) == codes
    assert all(sum(1 for g in folded if mandateCodes.holds(g, k)) == 1 for k in codes), 'no two hold one'
    # the desk's two are revision 0, and can be put back
    back = mandateCodes.revert(0, user='bob', base=1)
    assert [g['code'] for g in back['groups']] == ['G3', '7L'] and back['revision'] == 2


def test_an_edit_on_screen_a_restore_and_the_workbook(store):
    seed = mandateCodes.current()['groups']
    with pytest.raises(ValidationError, match='already in force'):
        mandateCodes.save({'groups': seed}, user='alice')
    moved = [dict(seed[0], code=' G4 ')] + seed[1:]
    mandateCodes.save({'groups': moved}, 'G4 for ESG', user='alice', base=0)
    assert mandateCodes.current()['groups'][0]['code'] == 'G4'
    with pytest.raises(sleeveRepo.StaleError):
        mandateCodes.save({'groups': seed}, user='bob', base=0)
    with pytest.raises(ValidationError, match='list of groups'):
        mandateCodes.save({'codes': {}}, user='bob', base=1)
    mandateCodes.save({'groups': [_group('ESG', 'G3', types=[ESG]), dict(seed[1], name='Everything else')]},
                      user='bob', base=1)
    words = mandateCodes.history()[0]['changes']
    assert 'Group ESG added: G3 · PMG ESG · All allocations · All risk levels · All currencies · All hedging policies' in words
    assert 'Group PMG ESG taken off (G4 · PMG ESG · All allocations · All risk levels · All currencies · All hedging policies)' in words
    assert words[-1] == '368 portfolios take another code'
    mandateCodes.revert(1, user='bob', base=2)
    assert [g['code'] for g in mandateCodes.current()['groups']] == ['G4', '7L']
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(mandateCodes.exportWorkbook()))
    assert book.sheetnames == ['Mandate Codes'] and book.active.max_row == 3
    assert [c.value for c in book.active[1]] == HEAD


def test_the_tab_is_the_admins_and_answers_over_its_routes(store, tmp_path):
    got = dashboardRouter.getMandateCodes(caller=_caller('alice'))
    assert got['revision'] == 0 and got['cells'] == len(mandateCodes.cells())
    assert [t['type'] for t in got['layout']][:2] == [MA, ESG] and got['codeMax'] == 40 and got['nameMax'] == 120
    assert [o['value'] for o in got['options']['types']] == mandateCodes.types()
    assert {'value': 'Full ex-RA', 'label': 'Full · ex Real Assets'} in got['options']['forms']
    assert [g['code'] for g in got['seed']] == ['G3', '7L']
    groups = [_group('GBP', 'X', currencies=['GBP']), _group('Rest', 'Y')]
    saved = dashboardRouter.saveMandateCodes({'groups': groups, 'base': 0}, caller=_caller('alice'))
    assert saved['mandateCodes']['revision'] == 1
    assert dashboardRouter.saveMandateCodes({'groups': groups, 'base': 0}, caller=_caller('alice')).status_code == 409
    assert dashboardRouter.saveMandateCodes({'groups': [_group('', 'Z')], 'base': 1},
                                            caller=_caller('alice')).status_code == 422
    assert dashboardRouter.getMandateCodesHistory(caller=_caller('alice'))['history'][0]['revision'] == 1
    rows = [HEAD, ['All', 'W', 'All', 'All', 'All', 'All', 'All']]
    preview = dashboardRouter.previewMandateCodes({'name': 'codes.csv', 'content': _upload(tmp_path, rows)},
                                                  caller=_caller('alice'))
    assert preview['changed'] > 0
    done = dashboardRouter.importMandateCodes({'name': 'codes.csv', 'content': _upload(tmp_path, rows),
                                               'base': 1}, caller=_caller('alice'))
    assert set(done['mandateCodes']['codes'].values()) == {'W'}
    back = dashboardRouter.revertMandateCodes({'revision': 1, 'base': 2}, caller=_caller('alice'))
    assert back['mandateCodes']['revision'] == 3
    assert dashboardRouter.exportMandateCodes(caller=_caller('alice')).body[:2] == b'PK'
    from cyrus_pmg.pmgService.core.accessControl import requireAdmin
    for route in dashboardRouter.router.routes:
        if route.path.startswith('/scenario/repository/mandate-codes'):
            assert requireAdmin in [d.call for d in route.dependant.dependencies], route.path


def test_the_console_mirrors_which_group_gives_a_portfolio_its_code():
    """The page's own reading of the groups (aocHolds, aocResolved) gives
    each portfolio the code the server does."""
    import json
    import os
    import shutil
    import subprocess
    node = shutil.which('node')
    if not node:
        pytest.skip('node not available')
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', 'generator', 'js', 'repository.js')
    with open(path, encoding='utf-8') as fh:
        source = fh.read()

    def fn(name):
        start = source.index('function ' + name + '(')
        return source[start:source.index('\n}\n', start)] + '\n}\n'
    groups = [_group('GBP Core', 'A', types=[MA], forms=['Core'], currencies=['GBP']),
              _group('ESG', 'G3', types=[ESG]), _group('Rest', '7L', risks=['Moderate'])]
    cells = mandateCodes.cells()
    script = ('var AOC_DIMS = %s; var aoc = { draft: %s, data: { _cells: %s } };\n' % (
        json.dumps(list(mandateCodes.GROUP_DIMENSIONS)), json.dumps(groups), json.dumps(cells))
        + fn('aocAllCells') + fn('aocGroups') + fn('aocHolds') + fn('aocResolved') + '''
var r = aocResolved(), out = {};
Object.keys(r.given).forEach(function (k) { out[k] = aoc.draft[r.given[k]].code; });
process.stdout.write(JSON.stringify({ codes: out, counts: r.counts }));
''')
    got = json.loads(subprocess.run([node, '-e', script], capture_output=True, text=True, check=True).stdout)
    given = mandateCodes.resolve(groups)
    assert got['codes'] == {k: groups[n]['code'] for k, n in given.items()}
    assert [c['gives'] for c in got['counts']] == [sum(1 for n in given.values() if n == i) for i in range(3)]
