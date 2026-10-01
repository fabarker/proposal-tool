"""D136, D141: the initial allocation of a private-markets book.

A commitment to private markets is called over years; until it is, the money
earmarked for it is held one third in Investment Grade Fixed Income and two
thirds in Public Equity. The model carries that initial allocation beside the
long-term one; the page shows it on demand, in columns paired with their
long-term twins (D136); the workbook gives it a sheet of its own and the deck
a slide of its own, after the long-term table's, for a book that holds
private markets and never for one that does not (D141).
"""

import io
import json
import os
import shutil
import subprocess

import pytest

from cyrus_pmg.pmgService.scenario import assetEstimates, rules, sheetDoc, pptWriter
from cyrus_pmg.pmgService.scenario.bakedAdapter import BakedScenarioPort
from cyrus_pmg.pmgService.scenario.proposalRegister import implementedPicture
from cyrus_pmg.pmgService.scenario.sleeves import listSleeves
from cyrus_pmg.pmgService.scenario.types import BasisInput, MandateInput, PortfolioKey
from cyrus_pmg.pmgService.scenario.workbook import (
    buildImplementationRows, implColumns, writeWorkbook)

HERE = os.path.dirname(os.path.abspath(__file__))
PORT = BakedScenarioPort()
IG, EQ = 'Investment Grade Fixed Income', 'Public Equity'
PRIVATE = rules.PRIVATE_FUNDED_FROM
VARIANT = 'PMG Multi-Asset Portfolio'


def _book(currency, key):
    basis = BasisInput(currency=currency, hedging='Hedged')
    return basis, PORT.resolve_portfolio(basis, PortfolioKey.fromStr(key))


def _sleeves(result):
    """The first sleeve the library offers for every category the book holds."""
    baseKey = PortfolioKey.fromStr(result['keyStr'])
    chosen = {}
    for category in result['categories']:
        name = rules.sleeveCategory(category['name'])
        library = listSleeves(name, VARIANT, baseKey)
        if library:
            chosen[name] = library[0]['name']
    return chosen


def _model(currency, key, tilt=False, vol=False, schedule='CASP', mandate=50e6):
    basis, result = _book(currency, key)
    return buildImplementationRows(result, _sleeves(result), rules.AUTO_SLEEVE_CATEGORIES,
                                   mandate, VARIANT, tilt, schedule,
                                   'PMG Target' if schedule else None, 50e6, vol, currency)


def _group(model, name):
    return next((g for g in model['groups'] if g['category'] == name), None)


# --------------------------------------------------------------- the rule ---

def test_the_rule_parks_the_private_line_a_third_and_two_thirds():
    lines = [{'name': IG, 'weightPct': 24.0}, {'name': 'Other Fixed Income', 'weightPct': 9.0},
             {'name': EQ, 'weightPct': 38.0}, {'name': PRIVATE, 'weightPct': 9.0},
             {'name': 'Hedge Funds', 'weightPct': 20.0}]
    parked = rules.initialLines(lines)
    by = {line['name']: line['weightPct'] for line in parked}
    assert by[IG] == pytest.approx(27.0) and by[EQ] == pytest.approx(44.0)
    assert by[PRIVATE] == 0.0 and by['Hedge Funds'] == 20.0
    assert sum(by.values()) == pytest.approx(100.0)
    assert lines[3]['weightPct'] == 9.0, 'the caller keeps its lines'
    assert [line['name'] for line in parked] == [line['name'] for line in lines], 'order kept'
    assert rules.initialLines([{'name': IG, 'weightPct': 60.0}, {'name': EQ, 'weightPct': 40.0}]) is None
    # the house split in force: D136's thirds, kept to four places (D148)
    assert rules.privateFunding() == [(IG, pytest.approx(1 / 3, abs=1e-6)),
                                      (EQ, pytest.approx(2 / 3, abs=1e-6))]
    # and any split handed in parks by that split instead
    halves = rules.initialLines(lines, [('Other Fixed Income', 0.5), (EQ, 0.5)])
    by = {line['name']: line['weightPct'] for line in halves}
    assert by['Other Fixed Income'] == pytest.approx(13.5) and by[EQ] == pytest.approx(42.5)
    assert by[IG] == 24.0 and by[PRIVATE] == 0.0


def test_the_schema_serves_the_rule_so_the_page_never_restates_it():
    schema = PORT.get_schema(BasisInput(currency='USD', hedging='Hedged'),
                             MandateInput(topAccountSize=5e7, mandateSize=5e7, primaryPwa='x'))
    served = schema['rules']['privateFunding']
    assert served['from'] == PRIVATE
    assert [(t['category'], t['share']) for t in served['to']] == rules.privateFunding()
    # the split as entered, and which revision of which scope is in force (D148)
    assert [t['weightPct'] for t in served['to']] == [33.3333, 66.6667]
    assert served['scope'] == '*' and served['revision'] >= 1 and not served['inherited']


# -------------------------------------------------------------- the model ---

def test_the_model_carries_the_initial_allocation_beside_the_long_term_one():
    model = _model('USD', 'USD|Moderate|Full|0', tilt=True, vol=True)
    initial = model['initial']
    assert initial is not None
    items = [i for g in model['groups'] for i in g['items']]
    assert round(initial['total']['weightPct'], 6) == 100.0, 'the column closes on 100.00'
    assert initial['total']['notional'] == 50e6
    committed = initial['commitment']
    private = _group(model, PRIVATE)
    assert committed['weightPct'] == pytest.approx(sum(i['printedPct'] for i in private['items']))
    assert committed['notional'] == sum(i['notional'] for i in private['items'])
    held = {h['category']: h for h in committed['held']}
    assert held[IG]['weightPct'] == pytest.approx(committed['weightPct'] / 3, abs=0.02)
    assert held[EQ]['weightPct'] == pytest.approx(committed['weightPct'] * 2 / 3, abs=0.02)
    assert held[IG]['notional'] + held[EQ]['notional'] == pytest.approx(committed['notional'], abs=200)
    for item in private['items']:
        assert item['initialPct'] == 0.0 and item['initialNotional'] == 0.0
    for group in model['groups']:
        if group['category'] in (IG, EQ, PRIVATE):
            continue
        for item in group['items']:
            # nothing else moves, bar where the remainders fall
            assert item['initialPct'] == pytest.approx(item['printedPct'], abs=0.011), item['name']
    for item in items:
        fee = (float(item['productCost']) + item['managementFee']) * item['initialPct']
        assert item['initialWtdFeeBp'] == pytest.approx(fee)
        assert item['initialNotional'] == round(50e6 * item['initialPct'] / 1e4) * 100
    assert initial['total']['wtdFeeBp'] < model['total']['wtdFeeBp'], 'parking costs less'
    # the export gate reads the long-term positions only
    assert {b['name'] for b in model['breaches']} == {
        i['name'] for g in model['groups'] for i in g['items'] if i['belowMinimum']}


def test_a_book_without_private_markets_is_exactly_today():
    model = _model('USD', 'USD|Moderate|ex-Alts|0', tilt=True, vol=True)
    assert model['initial'] is None
    assert not any('initialPct' in i for g in model['groups'] for i in g['items'])


def test_the_tilt_that_empties_igfi_still_parks_a_third_there():
    """32 of the 448 private-markets books hold exactly 8% IGFI, which the
    tactical tilt takes entirely. The IGFI row goes from the long-term table
    but the initial allocation still parks P/3 there, through the sleeve's
    shares; the empty long-term lines are not positions, so they breach
    nothing, and the long-term gate is unchanged."""
    model = _model('EUR', 'EUR|Higher Risk|Full|0', tilt=True)
    ig = _group(model, IG)
    assert ig is not None, 'the row stays for its initial figures'
    assert ig['weightPct'] == 0.0
    assert ig['initialWeightPct'] == pytest.approx(3.0)
    assert sum(i['initialPct'] for i in ig['items']) == pytest.approx(3.0)
    assert all(i['printedPct'] == 0.0 and not i['belowMinimum'] for i in ig['items'])
    assert not any(b['category'] == IG for b in model['breaches'])
    # at $50m each IG product is bought below its $5m minimum on day one:
    # said, not gated (the gate stays on the long-term commitments)
    assert {b['name'] for b in model['initial']['breaches']} >= {i['name'] for i in ig['items']}


def test_a_private_line_with_no_sleeve_yet_commits_its_own_weight():
    """Before a private-markets sleeve is chosen the band has no items; the
    commitment is then the band's own weight and notional, not nothing."""
    basis, result = _book('USD', 'USD|Moderate|Full|0')
    sleeves = {name: sleeve for name, sleeve in _sleeves(result).items() if name != PRIVATE}
    model = buildImplementationRows(result, sleeves, rules.AUTO_SLEEVE_CATEGORIES, 50e6,
                                    VARIANT, False, 'CASP', 'PMG Target', 50e6, False, 'USD')
    private = _group(model, PRIVATE)
    assert private is not None and not private['items']
    committed = model['initial']['commitment']
    assert committed['weightPct'] == pytest.approx(9.0)
    assert committed['notional'] == 4.5e6
    held = {h['category']: h['weightPct'] for h in committed['held']}
    assert held[IG] + held[EQ] == pytest.approx(9.0, abs=0.02)


def test_an_unpriced_book_has_no_initial_fee():
    model = _model('USD', 'USD|Moderate|Full|0', schedule=None)
    assert model['initial']['total']['wtdFeeBp'] is None
    assert all(i['initialWtdFeeBp'] is None for g in model['groups'] for i in g['items'])
    # the slide compares allocation and notional alone; the workbook weights
    # the product cost in the fee's place (D152)
    assert [key for _, key in sheetDoc.initialMeasures(False, deck=True)] == ['pct', 'notional']
    assert [key for _, key in sheetDoc.initialMeasures(False)] == ['pct', 'notional', 'cost']
    doc, header, _ = sheetDoc.buildInitialAllocationDoc(model, includeFees=False, deck=True)
    assert [doc.rows[header].cells[c].value for c in sorted(doc.rows[header].cells)] == \
        ['Categories & Asset Classes', 'Products', 'Long-term', 'Initial', 'Change',
         'Long-term', 'Initial', 'Change']
    doc, header, total = sheetDoc.buildInitialAllocationDoc(model, includeFees=False)
    assert doc.rows[header - 1].cells[7].value == 'Weighted cost (bp)'
    cost = lambda field: sum(i[field] * float(i['productCost'])
                             for g in model['groups'] for i in g['items'])
    assert doc.rows[total].cells[7].value == pytest.approx(cost('printedPct'))
    assert doc.rows[total].cells[8].value == pytest.approx(cost('initialPct'))


def test_the_slide_drops_the_green_the_workbook_keeps():
    """D152: the deck's initial allocation carries no green; the sheet does."""
    model = _model('USD', 'USD|Moderate|Full|0')
    greens = {sheetDoc.INITIAL_HEAD, sheetDoc.INITIAL_BAND, sheetDoc.INITIAL_TINT}
    def fills(doc):
        return {c.fill for r in doc.rows.values() for c in r.cells.values() if c.fill}
    deck, _, _ = sheetDoc.buildInitialAllocationDoc(model, deck=True)
    sheet, _, _ = sheetDoc.buildInitialAllocationDoc(model)
    assert not fills(deck) & greens
    assert fills(sheet) & greens


# ------------------------------------------------------------- the files ---

def _golden():
    import test_deck
    results, implementation = test_deck._goldenCase()
    return results, implementation


def _workbook(results, implementation):
    return writeWorkbook(
        BasisInput(currency='USD', hedging='Hedged'),
        MandateInput(topAccountSize=5e7, mandateSize=5e7, primaryPwa='x'),
        results, implementation['sleeves'], rules.AUTO_SLEEVE_CATEGORIES,
        implementation['variant'], implementation['tacticalTilt'],
        implementation['feeSchedule'], implementation['feeLevel'],
        implementation['includeFees'], implementation['volPremium'],
        assets=assetEstimates.forSlice('USD', 'Hedged'))


def _sheetTable(sheet, first='Categories & Asset Classes'):
    header = next(r for r in range(1, 14) if sheet.cell(row=r, column=1).value == first)
    names = [n for n in (sheet.cell(row=header, column=c).value
                         for c in range(1, sheet.max_column + 1)) if n]
    return header, names


def test_the_workbook_gives_the_initial_allocation_a_sheet_of_its_own():
    """D141. The Implementation sheet is the long-term target's alone - the
    columns of a book without private markets - and an Initial Allocation
    sheet follows it: each measure's name over a Long-term and an Initial
    column, teal-headed where initial, the commitment in its preamble, both
    totals closing on 100% and the mandate, the initial fee the lower."""
    from openpyxl import load_workbook
    results, implementation = _golden()
    book = load_workbook(io.BytesIO(_workbook(results, implementation)))
    assert book.sheetnames[:5] == ['portfolios', 'risk_dashboard', 'assumptions',
                                   'Implementation', 'Initial Allocation']
    impl = book['Implementation']
    _, names = _sheetTable(impl)
    assert names == implColumns(True)
    assert not any(str(impl.cell(row=r, column=1).value).startswith('Private Markets')
                   for r in range(1, 14))
    sheet = book['Initial Allocation']
    header, names = _sheetTable(sheet)
    # long-term first in each pair (D154; D142 put initial first)
    assert names == ['Categories & Asset Classes', 'Products'] + ['Long-term', 'Initial'] * 3
    measures = [sheet.cell(row=header - 1, column=c).value for c in (3, 5, 7)]
    assert measures == ['Allocation (%)', 'Notional', 'Weighted fee (bp)']
    merged = {str(m) for m in sheet.merged_cells.ranges}
    assert {'{0}{2}:{1}{2}'.format(a, b, header - 1)
            for a, b in (('C', 'D'), ('E', 'F'), ('G', 'H'))} <= merged
    for c, name in enumerate(names, start=1):
        fill = (sheet.cell(row=header, column=c).fill.fgColor.rgb or '')[-6:]
        assert fill == (sheetDoc.INITIAL_HEAD if name == 'Initial' else sheetDoc.HEADER_NAVY)
    labels = {sheet.cell(row=r, column=1).value: sheet.cell(row=r, column=2).value
              for r in range(1, header - 1)}
    assert labels['Private Markets Commitment'].endswith('(9.00%), invested as capital is called')
    assert '3.00% Investment Grade Fixed Income' in labels['Held Until Called']
    assert '6.00% Public Equity' in labels['Held Until Called']
    total = next(r for r in range(header, sheet.max_row + 1)
                 if sheet.cell(row=r, column=1).value == 'Total')
    assert [sheet.cell(row=total, column=c).value for c in (3, 4)] == [pytest.approx(1.0)] * 2
    assert [sheet.cell(row=total, column=c).value for c in (5, 6)] == [5e7, 5e7]
    assert sheet.cell(row=total, column=8).value < sheet.cell(row=total, column=7).value
    private = next(r for r in range(header, total) if sheet.cell(row=r, column=1).value == PRIVATE)
    assert sheet.cell(row=private, column=4).value == 0.0
    assert sheet.cell(row=private, column=6).value == 0
    # the measures' names in the header's size, and a light silver rule
    # between the measures from their names down to the total (D142, D147)
    assert {sheet.cell(row=header - 1, column=c).font.sz for c in (3, 5, 7)} == {12}
    silver = ('thin', 'C0C0C0')

    def edge(side):
        if side is None or side.style is None:
            return None
        return (side.style, (side.color.rgb if side.color is not None else '')[-6:])
    for row in range(header - 1, total + 1):
        for last in (4, 6):
            assert edge(sheet.cell(row=row, column=last).border.right) == silver, (row, last)
            assert edge(sheet.cell(row=row, column=last + 1).border.left) == silver, (row, last)
        if row != header - 1:                              # the names are merged over a pair
            assert edge(sheet.cell(row=row, column=3).border.right) is None, 'inside a pair'
    assert edge(sheet.cell(row=header - 1, column=3).border.right) == silver, 'the merge origin'
    assert sheet.freeze_panes == 'A{}'.format(header + 1)
    assert sheet.protection.sheet, 'locked like every other sheet (D98)'


def test_a_book_without_private_markets_writes_todays_sheets():
    from openpyxl import load_workbook
    basis, plain = _book('USD', 'USD|Moderate|ex-Alts|0')
    content = writeWorkbook(basis, MandateInput(topAccountSize=5e7, mandateSize=5e7, primaryPwa='x'),
                            [plain], _sleeves(plain), rules.AUTO_SLEEVE_CATEGORIES, VARIANT,
                            False, 'CASP', 'PMG Target', True, False,
                            assets=assetEstimates.forSlice('USD', 'Hedged'))
    book = load_workbook(io.BytesIO(content))
    assert 'Initial Allocation' not in book.sheetnames
    sheet = book['Implementation']
    header, names = _sheetTable(sheet)
    assert names == implColumns(True)
    assert not any(str(sheet.cell(row=r, column=1).value).startswith('Private Markets')
                   for r in range(1, header))


def test_the_deck_gives_the_initial_allocation_a_slide_of_its_own():
    """D141. The implementation slide carries the long-term target's columns
    alone; the slide after it, 'Implemented Portfolio: Initial Allocation',
    compares: the names, then each measure as Long-term and Initial under
    its name, which repeats with the header - at the cap, one line, one row
    height, across the slide; the commitment in its subheading, the day-one
    holdings and the fees in its notes. A book without private markets has
    no such slide."""
    import test_deck
    plan = test_deck._plan()
    headings = [e.get('heading') for e in plan]
    at = headings.index('Implemented Model')
    assert headings[at + 1] == 'Implemented Portfolio: Initial Allocation'
    assert headings[at + 2] == 'Composition of the Implemented Model'
    (impl,) = plan[at]['tables']
    head = impl.doc.rows[impl.header[0]].cells
    assert [head[c].value for c in sorted(head)] == implColumns(True)
    assert plan[at]['notes'][1].endswith('is on the next slide.')
    entry = plan[at + 1]
    (tp,) = entry['tables']
    assert len(tp.header) == 2, 'the measures repeat with the header'
    measures = tp.doc.rows[tp.header[0]].cells
    # on the slide each measure carries a Change column, and no green (D152)
    assert [measures[c].value for c in (3, 6, 9)] == ['Allocation (%)', 'Notional',
                                                      'Weighted fee (bp)']
    heads = tp.doc.rows[tp.header[1]].cells
    assert [heads[c].value for c in sorted(heads)] == \
        ['Categories & Asset Classes', 'Products'] + ['Long-term', 'Initial', 'Change'] * 3
    assert tp.natural and tp.font == pptWriter.BASE_PT
    assert pptWriter._measure(tp, tp.font)[1], 'every figure fits'
    # the measures' names set as large as the header (D142), and the light
    # silver rule between the measures drawn from both sides, from the
    # names' row down - the names' merged cells carrying it from their
    # origins (D147)
    assert {measures[c].font['size'] for c in (3, 6, 9)} == {12}
    edges = pptWriter.sharedEdges(tp)
    for r in tp.rows:
        for last in (5, 8):
            assert edges[(r, last)].get('right') == sheetDoc.MEASURE_RULE, (r, last)
            assert edges[(r, last + 1)].get('left') == sheetDoc.MEASURE_RULE, (r, last)
    for origin in (3, 6):
        assert edges[(tp.header[0], origin)].get('right') == sheetDoc.MEASURE_RULE
    assert len(set(tp.heights)) == 1
    assert sum(tp.widths()) == pytest.approx(tp.box[2])
    assert entry['sub'] == ('Until capital is called, the 9.00% ($2,250,000) committed to '
                            'private markets is held a third in Investment Grade Fixed Income '
                            'and two thirds in Public Equity')
    assert entry['notes'][0].endswith('3.00% in Investment Grade Fixed Income and 6.00% in '
                                      'Public Equity.')
    assert entry['notes'][1].startswith('Fees are charged on what is invested: 93.0bp')
    _, plain = _book('USD', 'USD|Moderate|ex-Alts|0')
    model = buildImplementationRows(plain, _sleeves(plain), rules.AUTO_SLEEVE_CATEGORIES,
                                    25e6, VARIANT, False, 'CASP', 'PMG Target', 1e7, False, 'USD')
    others = pptWriter.planDeck(BasisInput(currency='USD', hedging='Hedged'),
                                test_deck.MANDATE, [plain], model,
                                assetEstimates.forSlice('USD', 'Hedged'), True, VARIANT,
                                'CASP', 'PMG Target', None)
    assert 'Implemented Portfolio: Initial Allocation' not in [e.get('heading') for e in others]


def test_the_register_keeps_the_initial_figures():
    import test_deck
    results, implementation = _golden()
    model = test_deck._model(results, implementation)
    picture = implementedPicture(model, implementation['variant'])
    items = [i for g in picture for i in g['items']]
    assert all('initialPct' in i and 'initialNotional' in i for i in items)
    private = next(g for g in picture if g['category'] == PRIVATE)
    assert all(i['initialPct'] == 0.0 for i in private['items'])
    assert sum(i['initialPct'] for i in items) == pytest.approx(100.0)


# -------------------------------------------------------------- the page ---

def _js():
    with open(os.path.join(HERE, '..', '..', 'generator', 'js', 'implementation.js'),
              encoding='utf-8') as handle:
        return handle.read()


def test_js_initial_lines_mirror_agrees_with_python():
    """The page parks the private line in JavaScript and the files in
    Python. Across private-markets books, with the tilt and the premium on
    and off, the two must land on the same lines."""
    node = shutil.which('node')
    if not node:
        pytest.skip('node not available')
    source = _js()
    fn = source[source.index('function privateFunding('):source.index('/* end of the initial overlay */')]
    cases, expected = [], []
    for currency, key in (('USD', 'USD|Moderate|Full|0'), ('USD', 'USD|ModAgg|ex-HFs|1'),
                          ('EUR', 'EUR|Higher Risk|Full|0'), ('USD', 'USD|Moderate|ex-Alts|0')):
        _, result = _book(currency, key)
        for tilt in (False, True):
            for vol in (False, True):
                categories = rules.implementedCategories(result['categories'], tilt, vol, currency)
                lines = [{'name': s['name'], 'weightPct': s['weightPct']}
                         for s in rules.sleeveCategories(categories)]
                cases.append(lines)
                expected.append(rules.initialLines(lines))
    served = {'from': PRIVATE, 'to': [{'category': c, 'share': s} for c, s in rules.privateFunding()]}
    stub = ('var App = { opt: function (path, fallback) { return ({'
            "'rules.privateFunding': " + json.dumps(served) + '})[path]; } };\n')
    script = stub + fn + '\nconst cases = ' + json.dumps(cases) + ';\n' \
        + 'process.stdout.write(JSON.stringify(cases.map(initialLines)));\n'
    got = json.loads(subprocess.run([node, '-e', script], capture_output=True, text=True,
                                    check=True).stdout)
    assert len(got) == len(expected) == 16
    for js, py in zip(got, expected):
        if py is None:
            assert js is None
            continue
        assert [l['name'] for l in js] == [l['name'] for l in py]
        assert [l['weightPct'] for l in js] == pytest.approx([l['weightPct'] for l in py], abs=1e-12)


def test_js_screen_columns_pair_each_initial_twin_with_its_own():
    """Placement B on the page: Initial (%) beside Allocation (%), Initial
    notional beside Notional, Initial fee beside Wtd fee (only with fees) -
    and nothing at all for a book without private markets."""
    node = shutil.which('node')
    if not node:
        pytest.skip('node not available')
    source = _js()
    fn = source[source.index('var IMPL_TEXT_COLUMNS'):source.index('function implBandSpan(')]
    script = fn + '\nprocess.stdout.write(JSON.stringify([implScreenColumns(true, true), ' \
        'implScreenColumns(false, true), implScreenColumns(true, false)]));\n'
    priced, plain, today = json.loads(subprocess.run([node, '-e', script], capture_output=True,
                                                     text=True, check=True).stdout)
    for twin, of in (('Initial (%)', 'Allocation (%)'), ('Initial notional', 'Notional'),
                     ('Initial fee', 'Wtd fee')):
        assert priced.index(twin) == priced.index(of) + 1
    assert len(priced) == len(today) + 3
    assert 'Initial fee' not in plain and len(plain) == len(today) - 2 + 2
    assert not any(name.startswith('Initial') for name in today)


def test_a_reload_keeps_the_sleeve_chosen_under_a_group():
    """The private-markets sleeve is kept under its group's name, which no
    base category carries; the base resolving must not take it for a sleeve
    of a category the book no longer holds (found building D136)."""
    node = shutil.which('node')
    if not node:
        pytest.skip('node not available')
    with open(os.path.join(HERE, '..', '..', 'generator', 'js', 'core.js'), encoding='utf-8') as handle:
        source = handle.read()
    fn = source[source.index('function onBaseReady('):source.index('/* ---- scenario basis')]

    def run(categories):
        stub = ('var said = [], pushed = 0;\n'
                'var groups = ' + json.dumps(rules.SLEEVE_GROUPS) + ';\n'
                'function opt(path, fallback) { return path === "rules.sleeveGroups" ? groups : fallback; }\n'
                'function announce(level, text) { said.push(text); }\n'
                'function pushSleeves() { pushed += 1; }\n'
                'function revalidateSleeves() {}\n'
                'var state = { columns: [{ status: "ready", data: { categories: '
                + json.dumps([{'name': name} for name in categories]) + ' } }], sleeves: '
                + json.dumps({IG: 'a', PRIVATE: 'b', 'Hedge Funds': 'c'}) + ' };\n')
        script = stub + fn + '\nonBaseReady();\n' \
            'process.stdout.write(JSON.stringify([Object.keys(state.sleeves), pushed, said]));\n'
        return json.loads(subprocess.run([node, '-e', script], capture_output=True, text=True,
                                         check=True).stdout)

    kept, pushed, said = run([IG, EQ, 'Private Equity', 'Other Private Assets', 'Hedge Funds'])
    assert sorted(kept) == sorted([IG, PRIVATE, 'Hedge Funds']) and pushed == 0 and not said
    kept, pushed, said = run([IG, EQ])
    assert kept == [IG] and pushed == 1 and len(said) == 2, 'a book without them still drops them'


def test_the_screen_closes_each_pair_on_the_right_and_pins_the_world_row():
    """D146. On the Implementation screen the green rule sits on the RIGHT
    edge of each initial column, closing its Long-term / Initial pair; and
    the world row's cells over the two pinned columns pin with them, opaque
    and above the rest, so the labels slide under them when the table
    scrolls across - unpinned again where the product column is (narrow)."""
    root = os.path.join(HERE, '..', '..')
    with open(os.path.join(root, 'proposalTool', 'static', 'css', 'proposalTool.css'),
              encoding='utf-8') as handle:
        css = handle.read()
    assert '.tbl.impl.ini-on .ini{box-shadow:inset -2px 0 0 #7FB8A6}' in css
    assert 'inset 2px 0 0 #7FB8A6' not in css
    assert ('.tbl.impl thead tr.grp th.gpin{position:sticky;z-index:6;'
            'background:var(--surface,#fff)}') in css
    assert '.tbl.impl thead tr.grp th.gpin.g1{left:0}' in css
    assert '.tbl.impl thead tr.grp th.gpin.g2{left:var(--impl-c1,248px)}' in css
    assert '.tbl.impl thead tr.grp th.gpin.g2{position:static}' in css
    source = _js()
    assert "var pin = at === 0 ? ' gpin g1' : at === 1 ? ' gpin g2' : '';" in source
