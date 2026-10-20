"""Amounts in dollars whatever the base currency (D183): the PWA enters the
mandate and the top account size in dollars, and every product's minimum is
in dollars, so a GBP proposal's mandate, notionals, minimums, fee tiers,
charges and the messages that quote them are all printed in dollars - on
the screen, in the sheet and on the slides alike."""
import pytest

from cyrus_pmg.pmgService.scenario import (customPortfolio, fees, otherCharges, pptWriter,
                                           rules, sheetDoc, sleeves)
from cyrus_pmg.pmgService.scenario.types import BasisInput, MandateInput, PortfolioKey, ValidationError

from test_scenario_backend import PORT, TOP_ACCOUNT, _implementationFor, _sleeveMap


def _gbpExport(tmp_path):
    from openpyxl import load_workbook
    basis = BasisInput(currency='GBP', hedging='Hedged')
    result = PORT.resolve_portfolio(basis, PortfolioKey('GBP', 'Moderate', 'Core', False))
    variant = sleeves.VARIANTS[0]
    chosen = _sleeveMap(result['categories'], variant)
    mandate = MandateInput(TOP_ACCOUNT, 26e6, 'M. Aldridge — Zurich')
    path = tmp_path / 'wb.xlsx'
    path.write_bytes(PORT.build_export(basis, mandate, [result], {
        'sleeves': chosen, 'variant': variant, 'feeSchedule': 'RDR', 'feeLevel': 'PMG Floor'}))
    return load_workbook(path)['Implementation']


def test_a_gbp_sheet_writes_its_notionals_and_tier_in_dollars(tmp_path):
    sheet = _gbpExport(tmp_path)
    rows = list(sheet.iter_rows())
    tier = fees.tierFor(TOP_ACCOUNT)
    assert rows[3][1].value == '{} (${:g}m – ${:g}m)'.format(
        tier['id'], tier['min'] / 1e6, tier['max'] / 1e6)
    formats = {c.number_format for row in rows for c in row
               if isinstance(c.value, (int, float)) and c.value > 1000}
    assert formats == {'$#,##0'}


def test_the_messages_quote_amounts_in_dollars():
    with pytest.raises(ValidationError) as caught:
        PORT.validate_mandate(MandateInput(48e6, 1000, 'M. Aldridge — Zurich'))
    assert caught.value.message == 'Mandate size must be at least ${:,.0f}.'.format(
        rules.MANDATE_FLOOR)
    field = next(f for f in otherCharges.FIELDS if f['unit'] == 'usd' and 'fixed' not in f)
    with pytest.raises(ValidationError) as caught:
        otherCharges.clean({field['id']: field['max'] * 10})
    assert '$' in caught.value.message
    assert otherCharges.fieldValues({field['id']: 5000})[field['id']] == '$5,000'
    assert customPortfolio.checks({'profile': 'Moderate', 'products': [
        {'productId': 'x', 'category': rules.PRIVATE_FUNDED_FROM, 'weightPct': 100.0}]},
        sleeves.VARIANTS[0], 1e6, {})['problems'][-1] == (
        'Private markets need a mandate of ${:,.0f} or more.'.format(rules.PRIVATE_ASSETS_MINIMUM))


def test_a_gbp_proposals_slides_and_fee_ladder_are_in_dollars():
    model = _implementationFor(PortfolioKey('USD', 'Moderate', 'Full', False), currency='GBP')
    assert model['tier']['label'].startswith('$')
    mandate = MandateInput(TOP_ACCOUNT, 26e6, 'M. Aldridge — Zurich')
    assert pptWriter.fieldValues(model, mandate)['mandateSize'] == '$26m'
    labels = sheetDoc.ladderLabels(fees.tiers())
    assert labels[0].startswith('First $') and labels[-1].startswith('>$')
    assert fees.tiers()[0]['label'] == 'Under $10m'
    assert fees.tierLabel(100e6, None) == '$100m and up'
    for text in (str(v) for g in model['groups'] for i in g['items'] for v in i.values()):
        assert '£' not in text
