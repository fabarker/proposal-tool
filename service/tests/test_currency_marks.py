"""Amounts in the proposal's own base currency (D181): every amount a
proposal prints - the mandate, a notional, a minimum, a tier's bounds, a flat
charge, the messages that quote them - carries its currency's mark, $, £, €
or CHF, on the screen, in the sheet and on the slides alike. Where every
currency shares the figure - the fee card, the catalogue - it is bare."""
import pytest

from cyrus_pmg.pmgService.scenario import (customPortfolio, fees, money, otherCharges, pptWriter,
                                           rules, sheetDoc, sleeves)
from cyrus_pmg.pmgService.scenario.types import BasisInput, MandateInput, PortfolioKey, ValidationError

from test_scenario_backend import PORT, TOP_ACCOUNT, _sleeveMap


def test_each_currency_has_its_mark_and_the_sheet_and_the_slides_agree():
    assert [money.amount(1234567, c) for c in ('USD', 'GBP', 'EUR', 'CHF', None)] == [
        '$1,234,567', '£1,234,567', '€1,234,567', 'CHF 1,234,567', '1,234,567']
    assert money.amount(-250, 'GBP') == '-£250'
    assert money.numberFormat('USD') == '$#,##0', 'the sheet USD has always had'
    assert money.numberFormat('GBP') == '"£"#,##0' and money.numberFormat(None) == '#,##0'
    for currency in ('USD', 'GBP', 'EUR', 'CHF', None):
        fmt = money.numberFormat(currency)
        assert fmt in sheetDoc.NUMBER_FORMATS
        # what Excel prints is what the slide prints
        assert sheetDoc.renderNumber(26e6, fmt) == money.amount(26e6, currency)
    assert sheetDoc.renderNumber(-1.5e6, money.numberFormat('EUR')) == '-€1,500,000'


def _gbpExport(tmp_path):
    from openpyxl import load_workbook
    basis = BasisInput(currency='GBP', hedging='Hedged')
    result = PORT.resolve_portfolio(basis, PortfolioKey('GBP', 'Moderate', 'Core', False))
    variant = sleeves.VARIANTS[0]
    chosen = _sleeveMap(result['categories'], variant)
    mandate = MandateInput(TOP_ACCOUNT, 26e6, 'M. Aldridge — Zurich')
    path = tmp_path / 'wb.xlsx'
    path.write_bytes(PORT.build_export(basis, mandate, [result], {
        'sleeves': chosen, 'variant': variant, 'feeSchedule': 'RDR', 'feeLevel': 'PMG Floor'}))
    return load_workbook(path)['Implementation']


def test_a_gbp_sheet_writes_its_notionals_and_tier_in_pounds(tmp_path):
    sheet = _gbpExport(tmp_path)
    rows = list(sheet.iter_rows())
    tier = fees.tierFor(TOP_ACCOUNT)
    assert rows[3][1].value == '{} (£{:g}m – £{:g}m)'.format(
        tier['id'], tier['min'] / 1e6, tier['max'] / 1e6)
    formats = {c.number_format for row in rows for c in row
               if isinstance(c.value, (int, float)) and c.value > 1000}
    assert '"£"#,##0' in formats and '$#,##0' not in formats


def test_the_messages_quote_amounts_in_the_proposals_currency():
    with pytest.raises(ValidationError) as caught:
        rules.validateMandate(MandateInput(48e6, 1000, 'M. Aldridge — Zurich'), lambda n: True, 'CHF')
    assert caught.value.message == 'Mandate size must be at least CHF {:,.0f}.'.format(
        rules.MANDATE_FLOOR)
    with pytest.raises(ValidationError) as caught:
        PORT.validate_mandate(MandateInput(48e6, 1000, 'M. Aldridge — Zurich'), 'EUR')
    assert '€' in caught.value.message
    field = next(f for f in otherCharges.FIELDS if f['unit'] == 'usd' and 'fixed' not in f)
    with pytest.raises(ValidationError) as caught:
        otherCharges.clean({field['id']: field['max'] * 10}, 'GBP')
    assert '£' in caught.value.message and '$' not in caught.value.message
    values = otherCharges.fieldValues({field['id']: 5000}, 'EUR')
    assert values[field['id']] == '€5,000'


def test_the_slides_fields_and_the_fee_ladder_are_in_the_currency():
    from test_scenario_backend import _implementationFor
    model = _implementationFor(PortfolioKey('USD', 'Moderate', 'Full', False), currency='GBP')
    assert model['currency'] == 'GBP' and model['tier']['label'].startswith('£')
    mandate = MandateInput(TOP_ACCOUNT, 26e6, 'M. Aldridge — Zurich')
    assert pptWriter.fieldValues(model, mandate)['mandateSize'] == '£26m'
    labels = sheetDoc.ladderLabels(fees.tiers(), 'GBP')
    assert labels[0].startswith('First £') and labels[-1].startswith('>£')


def test_the_card_every_currency_shares_is_bare():
    assert fees.tiers()[0]['label'] == 'Under 10m'
    assert fees.tierLabel(10e6, 25e6) == '10m – 25m'
    assert fees.tierLabel(10e6, 25e6, 'USD') == '$10m – $25m'
    assert fees.tierLabel(100e6, None, 'GBP') == '£100m and up'


def test_a_custom_portfolio_too_small_for_private_markets_says_so_in_its_currency():
    assert customPortfolio.checks({'profile': 'Moderate', 'products': [
        {'productId': 'x', 'category': rules.PRIVATE_FUNDED_FROM, 'weightPct': 100.0}]},
        sleeves.VARIANTS[0], 1e6, {}, currency='EUR')['problems'][-1] == (
        'Private markets need a mandate of €{:,.0f} or more.'.format(rules.PRIVATE_ASSETS_MINIMUM))
