"""D125: the exports are set in GS Sans and GS Sans Condensed.

The font files themselves (named for Office, embeddable, tabular figures in
the condensed face), the faces every production sheet names - while the
golden comparison keeps the library's - and the deck carrying the faces
inside itself as Embedded OpenType parts.
"""

import io
import struct
import zipfile

from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation

from cyrus_pmg.pmgService.scenario import assetEstimates, houseFonts, rules, sheetDoc
from cyrus_pmg.pmgService.scenario.workbook import writeWorkbook

from test_deck import BASIS, MANDATE, UID, _deck, _goldenCase

HOUSE = {sheetDoc.SANS, sheetDoc.SANS_LIGHT, sheetDoc.CONDENSED}
LIBRARY = set(sheetDoc.HOUSE_FACES)
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}


def _faces(doc):
    return {cell.font['name'] for row in doc.rows.values()
            for cell in row.cells.values() if cell.font and 'name' in cell.font}


def _docs(engineParity):
    results, _ = _goldenCase()
    assets = assetEstimates.forSlice('USD', 'Hedged')
    return [sheetDoc.buildPortfoliosDoc(results, engineParity),
            sheetDoc.buildRiskDoc(results, engineParity),
            sheetDoc.buildAssumptionsDoc(assets, results, engineParity)]


def _workbook():
    results, implementation = _goldenCase()
    return writeWorkbook(
        BASIS, MANDATE, results, implementation['sleeves'],
        rules.AUTO_SLEEVE_CATEGORIES, implementation['variant'],
        implementation['tacticalTilt'], implementation['feeSchedule'],
        implementation['feeLevel'], implementation['includeFees'],
        implementation['volPremium'],
        assets=assetEstimates.forSlice('USD', 'Hedged'), proposalId=UID)


def _readEot(data):
    """The EOT fields a reader checks, parsed straight from the layout."""
    eotSize, fontSize, version, flags = struct.unpack_from('<IIII', data)
    at = 16 + 10 + 2 + 4 + 2
    magic, = struct.unpack_from('<H', data, at)
    at += 2 + 16 + 8 + 4 + 16
    names = []
    for _ in range(4):
        _, size = struct.unpack_from('<HH', data, at)
        at += 4
        names.append(data[at:at + size].decode('utf-16-le'))
        at += size
    _, rootSize = struct.unpack_from('<HH', data, at)
    at += 4 + rootSize
    return {'eotSize': eotSize, 'fontSize': fontSize, 'version': version,
            'flags': flags, 'magic': magic, 'names': names, 'font': data[at:]}


def test_the_font_files_are_named_for_office_and_may_be_embedded():
    for typeface, slots in houseFonts.FACES.items():
        for slot, filename in slots.items():
            info = houseFonts.details(houseFonts.fontBytes(filename))
            assert info['names'][1] == typeface, filename
            assert info['names'][2] == ('Bold' if slot == 'bold' else 'Regular')
            assert info['fsType'] == 0, 'installable: may be embedded and edited'
            assert info['weight'] == {'bold': 700}.get(
                slot, 300 if typeface == sheetDoc.SANS_LIGHT else 400)
    # the light weight is its own family, grouped under GS Sans
    light = houseFonts.details(houseFonts.fontBytes('GSSans-Light.ttf'))['names']
    assert (light[16], light[17]) == (sheetDoc.SANS, 'Light')


def test_production_sheets_are_set_in_the_house_faces_and_parity_keeps_the_librarys():
    for doc in _docs(engineParity=False):
        assert _faces(doc) <= HOUSE, (doc.name, _faces(doc) - HOUSE)
    for doc in _docs(engineParity=True):
        assert _faces(doc) <= LIBRARY, (doc.name, _faces(doc) - LIBRARY)
    # the strategic sheet's bold Light headings take GS Sans Bold: Light has
    # no bold of its own for Office to find
    portfolios = _docs(engineParity=False)[0]
    header = portfolios.cell(1, 2)
    assert header.font['name'] == sheetDoc.SANS and header.font['bold']
    assert header.align.get('wrap'), 'a long portfolio name wraps in its header'


def test_the_workbook_names_only_house_faces():
    book = load_workbook(io.BytesIO(_workbook()))
    for sheet in book.worksheets:
        if sheet.sheet_state == 'hidden':
            continue
        faces = {cell.font.name for row in sheet.iter_rows() for cell in row
                 if cell.value is not None}
        assert faces <= HOUSE, (sheet.title, faces - HOUSE)
    with zipfile.ZipFile(io.BytesIO(_workbook())) as archive:
        charts = [name for name in archive.namelist() if name.startswith('xl/charts/chart')]
        assert charts
        for name in charts:
            latin = etree.fromstring(archive.read(name)).findall('.//a:latin', NS)
            assert latin and {e.get('typeface') for e in latin} == {sheetDoc.SANS}, name


def test_the_deck_embeds_every_house_face_whole():
    content = _deck()
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        presentation = etree.fromstring(archive.read('ppt/presentation.xml'))
        rels = {rel.get('Id'): rel for rel in etree.fromstring(
            archive.read('ppt/_rels/presentation.xml.rels'))}
        assert presentation.get('embedTrueTypeFonts') == '1'
        assert presentation.get('saveSubsetFonts') is None
        children = [etree.QName(child).localname for child in presentation]
        at = children.index('embeddedFontLst')
        assert children[at - 1] == 'notesSz' and children[at + 1] == 'defaultTextStyle'
        listed = {}
        for entry in presentation.find('p:embeddedFontLst', NS):
            typeface = entry.find('p:font', NS).get('typeface')
            for slot in entry:
                if etree.QName(slot).localname == 'font':
                    continue
                rel = rels[slot.get('{%s}id' % NS['r'])]
                assert rel.get('Type').endswith('/relationships/font')
                listed[(typeface, etree.QName(slot).localname)] = \
                    archive.read('ppt/' + rel.get('Target'))
        assert set(listed) == {(t, s) for t, slots in houseFonts.FACES.items() for s in slots}
        for (typeface, slot), data in listed.items():
            eot = _readEot(data)
            original = houseFonts.fontBytes(houseFonts.FACES[typeface][slot])
            assert (eot['version'], eot['flags'], eot['magic']) == (0x00020001, 0, 0x504C)
            assert eot['eotSize'] == len(data) and eot['fontSize'] == len(original)
            assert eot['font'] == original, 'the whole file, not a subset'
            assert eot['names'][0] == typeface
        types = etree.fromstring(archive.read('[Content_Types].xml'))
        assert any(e.get('Extension') == 'fntdata' and
                   e.get('ContentType') == 'application/x-fontdata' for e in types)
        theme = etree.fromstring(archive.read('ppt/theme/theme1.xml'))
        for which in ('majorFont', 'minorFont'):
            assert theme.find('.//a:%s/a:latin' % which, NS).get('typeface') == sheetDoc.SANS


def test_every_run_in_the_deck_names_a_house_face():
    prs = Presentation(io.BytesIO(_deck()))
    faces = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            frames = []
            if shape.has_text_frame:
                frames.append(shape.text_frame)
            if shape.has_table:
                frames.extend(cell.text_frame for row in shape.table.rows for cell in row.cells)
            for frame in frames:
                faces.update(run.font.name for p in frame.paragraphs for run in p.runs)
    assert faces and faces <= HOUSE, faces - HOUSE


def test_a_font_that_forbids_embedding_is_refused():
    font = bytearray(houseFonts.fontBytes('GSSans-Regular.ttf'))
    tables = houseFonts._tables(bytes(font))
    offset, _ = tables['OS/2']
    struct.pack_into('>H', font, offset + 8, 0x0002)      # restricted licence
    try:
        houseFonts.eot(bytes(font))
    except ValueError:
        return
    raise AssertionError('a restricted font was embedded')
