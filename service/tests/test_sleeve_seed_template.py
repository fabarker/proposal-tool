"""The sleeve seed workbook (D159): the generated template loads through the
real importer - Sleeves and Rules read from one workbook by sheet name - and
the library exported into it loads back to the same library."""
import importlib
import os
import sys

import pytest
from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(ROOT, 'generator'))


@pytest.fixture
def repo(tmp_path, monkeypatch):
    monkeypatch.setenv('SCENARIO_SLEEVES_DB', str(tmp_path / 'sleeves.db'))
    from cyrus_pmg.pmgService.scenario import sleeveRepo
    return sleeveRepo


def _template(tmp_path):
    builder = importlib.import_module('build_sleeve_template')
    path = str(tmp_path / 'template.xlsx')
    builder.build(path)
    return builder, path


def _fill(path, sleeveRows, ruleRows):
    book = load_workbook(path)
    for name, rows in (('Sleeves', sleeveRows), ('Rules', ruleRows)):
        for i, row in enumerate(rows, start=2):
            for j, value in enumerate(row, start=1):
                book[name].cell(row=i, column=j, value=value)
    book.save(path)


def test_the_template_has_the_loader_headers(tmp_path, repo):
    _, path = _template(tmp_path)
    book = load_workbook(path)
    assert book.sheetnames[:3] == ['Instructions', 'Sleeves', 'Rules']
    assert [c.value for c in book['Sleeves'][1]][:6] == repo.SEED_COLUMNS_EDITIONS
    assert [c.value for c in book['Rules'][1]][:7] == repo.RULE_COLUMNS
    assert list(repo.readSeedRows(path)) == []          # empty: checks only, no rows
    assert repo.workbookRules(path)


def test_a_filled_template_with_an_edition_imports(tmp_path, repo):
    builder, path = _template(tmp_path)
    v = builder.vocabulary()
    a, b = v['products'][0]['productId'], v['products'][1]['productId']
    var, cat = v['variants'][0], 'Public Equity'
    _fill(path, [(var, cat, 'Seed Test', '', a, 0.6), (var, cat, 'Seed Test', '', b, 0.4),
                 (var, cat, 'Seed Test', 'GBP', a, 1.0)],
          [(var, cat, 'Seed Test', 'GBP', 'GBP', '', '')])
    rows = list(repo.readSeedRows(path))
    rules = list(repo.readRuleRows(path))
    assert len(rows) == 3 and len(rules) == 1
    assert repo.importRows(rows, replace=True, ruleRows=rules) == 2
    live = [s for s in repo.exportRows() if s[2] == 'Seed Test']
    assert sorted(r[3] for r in live) == ['', '', 'GBP']
    assert repo.exportRuleRows() == [(var, cat, 'Seed Test', 'GBP', 'GBP', '', '')]


def test_the_library_round_trips_through_the_workbook(tmp_path, repo):
    builder, _ = _template(tmp_path)
    repo.exportRows()                                   # seeds the temp store
    before, beforeRules = sorted(repo.exportRows()), sorted(repo.exportRuleRows())
    path = str(tmp_path / 'library.xlsx')
    builder.build(path, before, beforeRules)
    rows, rules = list(repo.readSeedRows(path)), list(repo.readRuleRows(path))
    repo.importRows(rows, replace=True, ruleRows=rules)
    assert sorted(repo.exportRows()) == before
    assert sorted(repo.exportRuleRows()) == beforeRules
