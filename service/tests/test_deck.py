"""Phase 3 of the PowerPoint export plan (D121): the deck renderer.

The same Docs the Excel renderer transcribes, rendered as slides - and held
to them cell for cell: every table cell's text must equal renderNumber of
the Doc's value under the Doc's format, the sign rule must resolve to the
CellIsRule's inks, the band fills and black rules must survive the trip,
and the doughnuts must carry the page's palette with the workbook's hole.
"""

import io
import json
import os

import pytest
from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.oxml.ns import qn

from cyrus_pmg.pmgService.scenario import assetEstimates, pptWriter, rules, sheetDoc
from cyrus_pmg.pmgService.scenario.pptWriter import writeDeck
from cyrus_pmg.pmgService.scenario.types import BasisInput, MandateInput
from cyrus_pmg.pmgService.scenario.workbook import (
    DONUT_PALETTE, buildImplementationRows, donutBreakdown, implColumns,
    stampedProposalId, _TEXT_COLUMNS, _WIDTHS)

HERE = os.path.dirname(os.path.abspath(__file__))
GOLDEN = os.path.join(HERE, 'golden')
UID = 'pr_0123456789ab'
BASIS = BasisInput(currency='USD', hedging='Hedged')
MANDATE = MandateInput(topAccountSize=10_000_000.0, mandateSize=25_000_000.0,
                       primaryPwa='J. Mercer')


def _goldenCase():
    with open(os.path.join(GOLDEN, 'engine_USD_Hedged_2col.results.json'),
              encoding='utf-8') as handle:
        results = json.load(handle)
    with open(os.path.join(GOLDEN, 'engine_USD_Hedged_2col.impl.json'),
              encoding='utf-8') as handle:
        implementation = json.load(handle)
    return results, implementation


def _model(results, implementation, includeFees=True):
    return buildImplementationRows(
        results[0], implementation['sleeves'], rules.AUTO_SLEEVE_CATEGORIES,
        MANDATE.mandateSize, implementation['variant'],
        implementation['tacticalTilt'],
        implementation['feeSchedule'] if includeFees else None,
        implementation['feeLevel'], MANDATE.topAccountSize,
        implementation['volPremium'], 'USD')


def _four(results):
    """Four portfolios: the golden pair and a renamed copy of it. The slide
    configuration turns on the count alone, so the figures may repeat."""
    return list(results) + [dict(r, name=r['name'] + ' (B)') for r in results]


def _deck(proposalId=UID, includeFees=True, cover=True, results=None):
    golden, implementation = _goldenCase()
    return writeDeck(BASIS, MANDATE, results or golden, implementation['sleeves'],
                     rules.AUTO_SLEEVE_CATEGORIES, implementation['variant'],
                     implementation['tacticalTilt'],
                     implementation['feeSchedule'], implementation['feeLevel'],
                     includeFees, implementation['volPremium'],
                     assets=assetEstimates.forSlice('USD', 'Hedged'),
                     proposalId=proposalId, cover=cover)


def _plan(results=None, includeFees=True):
    """The deck's own plan for the golden case, as writeDeck makes it."""
    golden, implementation = _goldenCase()
    return pptWriter.planDeck(
        BASIS, MANDATE, results or golden, _model(golden, implementation, includeFees),
        assetEstimates.forSlice('USD', 'Hedged'), includeFees,
        implementation['variant'],
        implementation['feeSchedule'] if includeFees else None,
        implementation['feeLevel'], UID)


def _tables(slide):
    return [shape.table for shape in slide.shapes if shape.has_table]


def _normal(text):
    lines = str(text).split('\n')
    return '\n'.join([lines[0]] + [line.strip() for line in lines[1:]])


def test_the_deck_is_the_plan_cell_for_cell():
    """The heart of parity: walk every planned table on every slide against
    the table drawn there. Every cell's text is renderNumber of the Doc's
    value under the Doc's format; every row sits at its planned height and
    every column at its planned width; every run at the planned size."""
    for results in (None, _four(_goldenCase()[0])):
        plan = _plan(results)
        prs = Presentation(io.BytesIO(_deck(results=results)))
        assert len(prs.slides) == len(plan)
        for slideIndex, entry in enumerate(plan):
            drawn = _tables(prs.slides[slideIndex])
            planned = entry.get('tables', [])
            assert len(drawn) == len(planned), (slideIndex, entry['kind'])
            for table, tp in zip(drawn, planned):
                columns = pptWriter._docColumns(tp.doc)
                assert (len(table.rows), len(table.columns)) == (len(tp.rows), columns)
                for i, width in enumerate(tp.widths()):
                    assert table.columns[i].width == pptWriter._pt(width)
                base = pptWriter._base(tp.doc)
                for tableRow, rowIndex in enumerate(tp.rows):
                    assert table.rows[tableRow].height == pptWriter._pt(tp.heights[tableRow])
                    spec = tp.doc.rows.get(rowIndex)
                    for column in range(1, columns + 1):
                        cell = spec.cells.get(column) if spec is not None else None
                        expected = ('' if cell is None
                                    else sheetDoc.renderNumber(cell.value, cell.fmt))
                        if column == 1:
                            expected = expected.strip()
                        if tp.oneLine:
                            expected = pptWriter._oneLineText(expected)     # D137
                        got = table.cell(tableRow, column - 1)
                        assert _normal(got.text) == _normal(expected), (
                            tp.doc.name, rowIndex, column)
                        font = (cell.font if cell is not None and cell.font else None) \
                            or pptWriter._DEFAULT_FONT
                        want = tp.font * font.get('size', base) / base
                        for paragraph in got.text_frame.paragraphs:
                            for run in paragraph.runs:
                                assert abs(run.font.size.pt - want) < 0.01


def test_one_or_two_portfolios_share_one_slide_without_the_repeated_rows():
    """D124. With no comparison or one: allocation and risk on ONE slide,
    allocation left and risk right; the risk table without the rows that
    repeat the allocation; both set in one size, both a little in from the
    full box. Each table's rows are at their own height - none opened to
    make the two end together (D138)."""
    plan = _plan()
    table = [entry for entry in plan if entry['kind'] == 'table']
    first = table[0]
    assert first['heading'] == 'Strategic Asset Allocation & Risk'
    alloc, risk = first['tables']
    assert (alloc.doc.name, risk.doc.name) == ('portfolios', 'risk_dashboard')
    assert risk.rows == pptWriter.headedRisk(sheetDoc.buildRiskDoc(_goldenCase()[0]))[1]
    labels = {str(risk.doc.rows[r].cells[1].value).strip()
              for r in risk.rows if 1 in risk.doc.rows[r].cells}
    for repeated in ('Public Equity', 'Estimated Mean Return', 'Sharpe Ratio', 'Volatility'):
        assert repeated not in labels, repeated
    assert 'Factor Based Risk Analytics' in labels and 'Financial Crisis' in labels
    assert alloc.font == risk.font, 'one size across the slide'
    full = pptWriter.FULL_BOX
    for tp in (alloc, risk):
        left, top, width, height = tp.box
        assert left > full[0] and top > full[1], 'more border than a lone table'
        assert left + width < full[0] + full[2] and top + height < full[1] + full[3]
        assert tp.heights == pptWriter._measure(tp, tp.font)[0], 'no row opened'
        assert sum(tp.heights) <= height + 1e-6
    assert alloc.box[0] + alloc.box[2] < risk.box[0], 'allocation left, risk right'
    # and no Risk Dashboard slide of its own
    assert not [e for e in table if e['heading'] == 'Risk Dashboard']


def test_three_or_more_portfolios_take_a_slide_each_at_full_width():
    """D124. With two comparisons or more: allocation on one slide and risk
    on the next, each filling the full box, and the risk table whole."""
    plan = _plan(_four(_goldenCase()[0]))
    headings = [entry.get('heading') for entry in plan]
    assert 'Strategic Asset Allocation & Risk' not in headings
    alloc = [e for e in plan if e.get('heading') == 'Strategic Asset Allocation']
    risk = [e for e in plan if e.get('heading') == 'Risk Dashboard']
    assert alloc and risk
    assert headings.index('Strategic Asset Allocation') + len(alloc) == \
        headings.index('Risk Dashboard'), 'risk follows allocation'
    for entry in alloc + risk:
        (tp,) = entry['tables']
        assert tp.box == pptWriter.FULL_BOX
    riskRows = [r for entry in risk for r in entry['tables'][0].rows]
    whole = risk[0]['tables'][0].doc
    assert set(riskRows) == set(whole.rows), 'the risk table keeps every row here'


def test_every_planned_table_fits_its_box_at_a_readable_size():
    """The fit's invariants, over both configurations and fees on and off:
    no table taller than its box, every size between MIN_PT and BASE_PT -
    9pt, the cap (D137) - or, for a one-line table its width holds under
    MIN_PT, no smaller than ONE_LINE_MIN_PT; every figure clear of its cell
    on one line, and one size per slide."""
    assert pptWriter.BASE_PT == 9.0
    golden = _goldenCase()[0]
    for results in (None, _four(golden)):
        for includeFees in (True, False):
            for entry in _plan(results, includeFees):
                plans = entry.get('tables', [])
                if len(plans) > 1:
                    assert len({tp.font for tp in plans}) == 1
                for tp in plans:
                    floor = pptWriter.ONE_LINE_MIN_PT if tp.oneLine else pptWriter.MIN_PT
                    assert floor <= tp.font <= pptWriter.BASE_PT
                    assert sum(tp.heights) <= tp.box[3] + 1e-6, (tp.doc.name, sum(tp.heights))
                    heights, fits = pptWriter._measure(tp, tp.font)
                    assert fits, (tp.doc.name, tp.font)


def _edge(tc, tag):
    """One edge of a table cell as written: ('none',) or (width, colour)."""
    line = tc.get_or_add_tcPr().find(qn(tag))
    if line is None or line.find(qn('a:noFill')) is not None:
        return ('none',)
    return (line.get('w'), line.find('.//' + qn('a:srgbClr')).get('val'))


@pytest.mark.parametrize('lineup', ['two', 'four'])
def test_every_rule_is_stated_on_both_cells_that_share_it(lineup):
    """D127. A rule written on one cell and contradicted by an explicit
    no-line on its neighbour is a rule PowerPoint does not draw. So every
    horizontal edge in every table reads the same from above and below, and
    the allocation table's rules - the header's, each class's, TOTAL's and
    the white ones in the metric block - are there from both sides. The one
    under the names and each class's are the deck's lighter weights (D137)."""
    golden = _goldenCase()[0]
    prs = Presentation(io.BytesIO(_deck(results=golden if lineup == 'two' else _four(golden))))
    ruled = 0
    for slide in prs.slides:
        for table in _tables(slide):
            rows = list(table.rows)
            for upper, lower in zip(rows, rows[1:]):
                for above, below in zip(upper.cells, lower.cells):
                    bottom, top = _edge(above._tc, 'a:lnB'), _edge(below._tc, 'a:lnT')
                    assert bottom == top, (above.text, below.text, bottom, top)
                    ruled += bottom != ('none',)
            # and left against right (D130): a merged cell's edge is its origin's
            for row in rows:
                cells, origin = list(row.cells), None
                for left, right in zip(cells, cells[1:]):
                    origin = origin if left.is_spanned else left
                    if right.is_spanned:
                        continue                         # inside one merged cell
                    assert _edge(origin._tc, 'a:lnR') == _edge(right._tc, 'a:lnL'), \
                        (origin.text, right.text)
    assert ruled > 0

    alloc = next(t for s in prs.slides for t in _tables(s)
                 if any(r.cells[0].text == 'TOTAL' for r in t.rows))
    labels = [r.cells[0].text for r in alloc.rows]
    thick = str(int(round(2.25 * 12700)))
    names, thin = str(int(round(1.25 * 12700))), str(int(round(0.5 * 12700)))
    expected = {labels.index('Other Fixed Income'): (thin, '000000'),
                labels.index('TOTAL'): (thick, '000000'),
                labels.index('Sharpe Ratio'): (thick, 'FFFFFF'),
                labels.index('Volatility'): (thick, 'FFFFFF')}
    for index, edge in expected.items():
        for column in range(len(alloc.columns)):
            assert _edge(alloc.cell(index - 1, column)._tc, 'a:lnB') == edge, labels[index]
    for column in range(len(alloc.columns)):
        assert _edge(alloc.cell(1, column)._tc, 'a:lnT') == (names, '000000'), 'under the names'
        assert _edge(alloc.cell(0, column)._tc, 'a:lnB') == (names, '000000'), 'under the names'


def test_the_combined_risk_table_names_its_band_in_the_header_row():
    """D128. On the combined slide the risk table opens on ONE navy row: the
    first band's title in the corner, the portfolio names beside it; the
    band's own row is gone, and the dotted rule that opened the stress block
    now runs under that header from both sides. The full-width slide, with
    four portfolios, keeps the band as its own row. The Doc is not touched."""
    golden = _goldenCase()[0]
    riskDoc = sheetDoc.buildRiskDoc(golden)
    band = min(riskDoc.sections)
    assert riskDoc.rows[band].cells[1].value == 'Factor Based Risk Analytics'
    view, rows = pptWriter.headedRisk(riskDoc)
    assert riskDoc.rows[1].cells[1].value is None, 'the sheet itself keeps its empty corner'
    assert band not in rows and rows[0] == 1
    corner = view.rows[1].cells[1]
    assert corner.value == 'Factor Based Risk Analytics'
    assert corner.fill == riskDoc.rows[1].cells[2].fill == sheetDoc.HEADER_NAVY
    assert [view.rows[1].cells[c].value for c in (2, 4)] == [r['name'] for r in golden]

    prs = Presentation(io.BytesIO(_deck(cover=False)))
    risk = [t for t in _tables(prs.slides[0]) if any('Oil Embargo' in c.text
                                                    for row in t.rows for c in row.cells)][0]
    first = [c.text for c in risk.rows[0].cells]
    assert first[0] == 'Factor Based Risk Analytics' and golden[0]['name'] in first
    assert 'Factor Based Risk Analytics' not in [row.cells[0].text for row in list(risk.rows)[1:]]
    dotted = (str(int(round(0.75 * 12700))), 'A9A9A9')
    for column in range(len(risk.columns)):
        assert _edge(risk.cell(0, column)._tc, 'a:lnB') == dotted, column
        assert _edge(risk.cell(1, column)._tc, 'a:lnT') == dotted, column

    four = Presentation(io.BytesIO(_deck(cover=False, results=_four(golden))))
    whole = next(t for s in four.slides for t in _tables(s)
                 if any('Oil Embargo' in c.text for row in t.rows for c in row.cells))
    labels = [row.cells[0].text for row in whole.rows]
    assert labels[0] == '' and 'Factor Based Risk Analytics' in labels[1:]


@pytest.mark.parametrize('lineup', ['two', 'four'])
def test_the_nominal_and_real_heads_are_underlined(lineup):
    """D129. Every Nominal and Real head in the deck's risk tables - the
    combined slide's and the full-width one's - is a single-underlined run;
    the stress block's label beside them and the figures under them are not."""
    golden = _goldenCase()[0]
    prs = Presentation(io.BytesIO(_deck(results=golden if lineup == 'two' else _four(golden))))
    heads = other = 0
    for slide in prs.slides:
        for table in _tables(slide):
            for row in table.rows:
                for cell in row.cells:
                    for run in (r for p in cell.text_frame.paragraphs for r in p.runs):
                        if not run.text:
                            continue
                        if run.text in ('Nominal', 'Real'):
                            heads += 1
                            assert run.font.underline is True, run.text
                        else:
                            other += 1
                            assert not run.font.underline, run.text
    assert heads == 2 * (len(golden) if lineup == 'two' else 2 * len(golden)) and other


def _noLabelWraps(tp):
    """Every unspanned column-A label of *tp* on one line at its size."""
    widths = tp.widths()
    assert widths[0] >= pptWriter._labelWidth(tp, tp.font) - 1e-6
    assert abs(sum(widths) - tp.box[2]) < 1e-6
    for r in tp.rows:
        spec = tp.doc.rows[r].cells.get(1)
        if r in tp.spans or spec is None or not spec.value:
            continue
        text = str(spec.value)
        size = tp.font * spec.font['size'] / pptWriter._base(tp.doc)
        inner = (widths[0] - 2 * pptWriter._PAD_H_EM * tp.font
                 - pptWriter._indentPt(spec, text, tp.font, 1, tp.doc)) * pptWriter._WIDTH_SLACK
        assert pptWriter._wrapLines(text.strip(), spec.font['name'], bool(spec.font.get('bold')),
                                    size, inner) == 1, text


@pytest.mark.parametrize('lineup', ['two', 'four'])
def test_no_allocation_label_wraps(lineup):
    """D130. The allocation's label column is fitted like the risk table's:
    the longest category or asset name never takes a second line, on the
    combined slide or the full-width one. And when a share is too narrow for
    its names, the column widens to them rather than wrapping one."""
    golden = _goldenCase()[0]
    plans = [tp for entry in _plan(golden if lineup == 'two' else _four(golden))
             for tp in entry.get('tables', []) if tp.doc.name == 'portfolios']
    assert plans and all(tp.fitLabel for tp in plans)
    for tp in plans:
        _noLabelWraps(tp)
        assert pptWriter._measure(tp, tp.font)[1], 'the figures still fit'

    doc = sheetDoc.buildPortfoliosDoc(golden)
    narrow = pptWriter.planTable(doc, pptWriter.FULL_BOX, pptWriter._shares(12.0, 3),
                                 fitLabel=True)[0]
    assert narrow.widths()[0] > pptWriter.FULL_BOX[2] * 0.12 + 1, 'the column widened'
    _noLabelWraps(narrow)


@pytest.mark.parametrize('lineup', ['two', 'four'])
def test_no_risk_label_wraps(lineup):
    """D128. The risk table's label column is fitted to its longest label at
    the size the table is set in - 'Predicted Performance Over Stress
    Periods' above all - so no label beside figures takes a second line, on
    the combined slide or the full-width one; the figures share the rest and
    still fit on one line each."""
    golden = _goldenCase()[0]
    plans = [tp for entry in _plan(golden if lineup == 'two' else _four(golden))
             for tp in entry.get('tables', []) if tp.doc.name == 'risk_dashboard']
    assert plans
    for tp in plans:
        widths = tp.widths()
        assert widths[0] >= pptWriter._labelWidth(tp, tp.font) - 1e-6
        assert abs(sum(widths) - tp.box[2]) < 1e-6
        _, fits = pptWriter._measure(tp, tp.font)
        assert fits, 'the figures still fit their narrower columns'
        seen = False
        for r in tp.rows:
            spec = tp.doc.rows[r].cells.get(1)
            if r in tp.spans or spec is None or not spec.value:
                continue
            text = str(spec.value)
            size = tp.font * spec.font['size'] / pptWriter._base(tp.doc)
            inner = (widths[0] - 2 * pptWriter._PAD_H_EM * tp.font
                     - pptWriter._indentPt(spec, text, tp.font, 1, tp.doc)) * pptWriter._WIDTH_SLACK
            assert pptWriter._wrapLines(text.strip(), spec.font['name'], bool(spec.font.get('bold')),
                                        size, inner) == 1, text
            seen = seen or text == 'Predicted Performance Over Stress Periods'
        assert seen


@pytest.mark.parametrize('lineup', ['two', 'four'])
def test_a_dotted_rule_parts_the_portfolios_on_the_risk_slides(lineup):
    """D130, D137. On every risk table in the deck, a rule runs between one
    portfolio's pair and the next from under the header to the last row -
    drawn exactly as the dotted grey rules across the table are, stated on
    both cells, carried by a merged cell's origin, and unbroken by the block
    headings, which span only to the first rule. Never across the header or
    the navy band, never inside a pair."""
    golden = _goldenCase()[0]
    count = len(golden) if lineup == 'two' else 2 * len(golden)
    prs = Presentation(io.BytesIO(_deck(results=golden if lineup == 'two' else _four(golden))))
    silver = (str(int(round(0.75 * 12700))), 'A9A9A9')
    tables = [t for s in prs.slides for t in _tables(s)
              if any('Oil Embargo' in c.text for row in t.rows for c in row.cells)]
    assert tables
    for table in tables:
        for index, row in enumerate(table.rows):
            cells = list(row.cells)
            navy = index == 0 or cells[0].text == 'Factor Based Risk Analytics'
            edges = [_edge(c._tc, 'a:lnR') for c in cells] + [_edge(c._tc, 'a:lnL') for c in cells]
            if navy:
                assert silver not in edges, cells[0].text
                continue
            for k in range(count - 1):
                last, following = 2 + 2 * k, 3 + 2 * k        # 0-based: a pair's Real, the next Nominal
                origin = last
                while cells[origin].is_spanned:
                    origin -= 1
                assert _edge(cells[origin]._tc, 'a:lnR') == silver, (cells[0].text, k)
                assert _edge(cells[following]._tc, 'a:lnL') == silver, (cells[0].text, k)
            if not cells[2].is_spanned and not cells[1].is_merge_origin:
                assert _edge(cells[1]._tc, 'a:lnR') != silver, 'inside a pair: ' + cells[0].text
        headings = [row for row in table.rows if row.cells[0].text.startswith('Value at Risk')]
        assert headings and headings[0].cells[0].is_merge_origin
        assert _edge(headings[0].cells[0]._tc, 'a:lnT') == silver, 'the same rule across'
        line = headings[0].cells[0]._tc.get_or_add_tcPr().find(qn('a:lnT'))
        assert line.find(qn('a:prstDash')).get('val') == 'sysDot'
        for row in table.rows:
            for cell in row.cells:
                right = cell._tc.get_or_add_tcPr().find(qn('a:lnR'))
                if _edge(cell._tc, 'a:lnR') == silver:
                    assert right.find(qn('a:prstDash')).get('val') == 'sysDot'
        assert not headings[0].cells[3].is_spanned, 'the heading stops at the first rule'


def test_the_deck_indents_a_quarter_less_and_follows_the_sheets_columns():
    """D134. An indented label - an asset line, a crisis, a horizon - steps in
    a quarter less than it did (1.35 em); and the implementation slide reads
    the sheet's own columns, Notional after the products and no Exposure
    ccy, while its Exposure currency doughnut is still drawn."""
    assert abs(pptWriter._INDENT_EM - 1.35 * 0.75) < 1e-9
    prs = Presentation(io.BytesIO(_deck(cover=False)))
    impl = next(t for s in prs.slides for t in _tables(s)
                if any(c.text == 'Categories & Asset Classes' for row in list(t.rows)[:2]
                       for c in row.cells))
    rows = list(impl.rows)
    head = [c.text for c in rows[1].cells]         # under the Long-term / Initial row (D136)
    assert head == implColumns(True, initial=True)
    assert head[:3] == ['Categories & Asset Classes', 'Products', 'Notional']
    assert 'Exposure ccy' not in head
    titles = [shape.chart.chart_title.text_frame.text for s in prs.slides
              for shape in s.shapes if shape.has_chart]
    assert 'Exposure currency' in titles and len(titles) == 5


def test_the_allocation_and_implementation_indent_a_third_less_than_the_risk_table():
    """D135. In the drawn deck an asset line on the allocation and a product
    line on the implementation step in by a third less than D134's step;
    a crisis or a horizon on the risk table keeps D134's; an asset on the
    assumptions steps in half as far (D137). Read off each cell's left
    margin: the padding plus the step, at the table's size."""
    third = pptWriter._INDENT_EM * 2 / 3
    assert pptWriter._INDENT_EM_BY_DOC == {'portfolios': third, 'Implementation': third,
                                           'assumptions': pptWriter._INDENT_EM / 2}
    plan = _plan()                                     # with its cover, as _deck() draws it
    prs = Presentation(io.BytesIO(_deck()))
    assert len(plan) == len(prs.slides)
    for entry, slide in zip(plan, prs.slides):
        for tp, table in zip(entry.get('tables', []), _tables(slide)):
            step = pptWriter._INDENT_EM_BY_DOC.get(tp.doc.name, pptWriter._INDENT_EM)
            checked = 0
            for position, r in enumerate(tp.rows):
                spec = tp.doc.rows[r].cells.get(1)
                if spec is None or not str(spec.value or '').startswith('  '):
                    continue
                want = pptWriter._pt((pptWriter._PAD_H_EM + step) * tp.font)
                assert table.cell(position, 0).margin_left == want, (tp.doc.name, spec.value)
                checked += 1
            assert checked, tp.doc.name


def test_the_risk_free_line_steps_in_like_an_asset_line():
    """D127. The allocation's risk-free line is set in by an indent level
    alone, with no leading spaces; the deck steps it in exactly as far as an
    asset line. The risk table's labels, which carry both, still step in
    once - by the risk table's own step (D135) - and a label with neither
    not at all."""
    results = _goldenCase()[0]
    doc = sheetDoc.buildPortfoliosDoc(results)
    labels = {row.cells[1].value: row.cells[1] for row in doc.rows.values()
              if 1 in row.cells and row.cells[1].value}
    step = pptWriter._indentPt(labels['  US Dollar Debt'], '  US Dollar Debt', 10.0, 1, doc)
    assert step > 0
    line = sheetDoc.RISK_FREE_LINE
    assert not line.startswith(' ') and labels[line].align['indent'] == 1
    assert pptWriter._indentPt(labels[line], line, 10.0, 1, doc) == step
    assert pptWriter._indentPt(labels['Estimated Mean Return'],
                               'Estimated Mean Return', 10.0, 1, doc) == 0
    risk = sheetDoc.buildRiskDoc(results)
    stress = [row.cells[1] for row in risk.rows.values()
              if 1 in row.cells and str(row.cells[1].value or '').startswith('  ')]
    assert stress and all(cell.align.get('indent') == 1 for cell in stress)
    riskStep = pptWriter._INDENT_EM * 10.0
    assert all(pptWriter._indentPt(cell, cell.value, 10.0, 1, risk) == riskStep
               for cell in stress)


def test_a_table_too_tall_for_one_slide_continues_with_its_header():
    """Under MIN_PT it paginates rather than shrinking further: cut at a
    section mark, the header repeated, one size on every page, every row on
    exactly one page."""
    golden = _goldenCase()[0]
    doc = sheetDoc.buildPortfoliosDoc(golden)
    short = (pptWriter.FULL_BOX[0], pptWriter.FULL_BOX[1], pptWriter.FULL_BOX[2], 220.0)
    plans = pptWriter.planTable(doc, short, pptWriter._shares(44.0, 3))
    assert len(plans) > 1
    # cut at the floor size, then every page re-fitted and set in the one size
    # the tallest page allows - never below the floor, never mixed
    assert len({tp.font for tp in plans}) == 1
    assert pptWriter.MIN_PT <= plans[0].font < pptWriter.BASE_PT
    body = [r for tp in plans for r in tp.rows if r != 1]
    assert body == [r for r in sorted(doc.rows) if r != 1]
    assert all(tp.rows[0] == 1 for tp in plans), 'the header repeats'
    assert all(tp.rows[1] in doc.sections for tp in plans[1:]), 'breaks at a category'
    assert all(sum(tp.heights) <= short[3] + 1e-6 for tp in plans)


def test_font_metrics_measure_the_faces_the_deck_embeds():
    """The committed metrics, spot-checked against the house faces' own
    figures (D125): GS Sans's figures are 0.58 em wide, GS Sans Condensed's
    tabular figures 0.494 em in both weights, and both lines 1.247 em."""
    sans, condensed = sheetDoc.SANS, sheetDoc.CONDENSED
    assert abs(pptWriter.textWidth('0', sans, False, 10.0) - 5.8) < 1e-9
    assert abs(pptWriter.textWidth('100.0%', sans, False, 10.0)
               - (4 * 5.8 + pptWriter.textWidth('.', sans, False, 10.0)
                  + pptWriter.textWidth('%', sans, False, 10.0))) < 1e-9
    for bold in (False, True):
        assert {pptWriter.textWidth(d, condensed, bold, 10.0) for d in '0123456789'} == {4.94}
    assert pptWriter._lineHeight(sans, False, 10.0) == pytest.approx(12.47, abs=0.01)
    # GS Sans Light has no bold cut and a face nothing measured has no
    # metrics: both are measured as GS Sans
    assert pptWriter.textWidth('Abc', sheetDoc.SANS_LIGHT, True, 11) == \
        pptWriter.textWidth('Abc', sans, True, 11)
    assert pptWriter.textWidth('Abc', 'Unmeasured', False, 11) == \
        pptWriter.textWidth('Abc', sans, False, 11)
    assert pptWriter._wrapLines('Conditional Value at Risk with 99% Confidence',
                                condensed, True, 10, 80) >= 2


def test_the_stress_rule_resolves_to_the_workbooks_inks():
    """Red below zero, green above, base ink outside the rule's range -
    resolved from the raw value exactly as the CellIsRule would."""
    prs = Presentation(io.BytesIO(_deck()))
    # the combined slide: allocation first, risk second (D124)
    table = _tables(prs.slides[1])[1]
    inks = {}
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    inks.setdefault(run.text, str(run.font.color.rgb))
    assert inks['-17.7%'] == '9C0006', 'a loss in the stress block is red'
    assert inks['10.7%'] == '006100', 'a gain in the stress block is green'
    # VaR is positive but OUTSIDE the rule's range: base ink, not green
    assert inks['17.4%'] == '000000'
    # and on the full-width risk slide, the Sharpe row too
    four = Presentation(io.BytesIO(_deck(results=_four(_goldenCase()[0]))))
    risk = [t for s in four.slides for t in _tables(s)
            if any('Oil Embargo' in c.text for row in t.rows for c in row.cells)][0]
    sharpe = [row for row in risk.rows if row.cells[0].text == 'Sharpe Ratio'][0]
    assert str(sharpe.cells[1].text_frame.paragraphs[0].runs[0].font.color.rgb) == '000000', \
        'Sharpe sits outside the rule'


def test_the_implementation_slide_keeps_the_bands_and_the_rules():
    """The category band's fill, the total's black rules, and no styling
    invented: left and right edges are explicit no-lines."""
    prs = Presentation(io.BytesIO(_deck()))
    table = _tables(prs.slides[2])[0]             # cover, combined, implementation
    bandCell = totalCell = None
    for row in table.rows:
        cells = list(row.cells)
        if cells[0].text == 'Public Equity':
            bandCell = cells[0]
        if cells[0].text == 'Total':
            totalCell = cells[0]
    assert bandCell is not None and totalCell is not None
    assert str(bandCell.fill.fore_color.rgb) == 'D3DDEA'
    tcPr = totalCell._tc.get_or_add_tcPr()
    for edge, lit in (('a:lnT', True), ('a:lnB', True), ('a:lnL', False)):
        lines = tcPr.findall(qn(edge))
        assert len(lines) == 1, edge
        fills = lines[0].findall(qn('a:solidFill'))
        if lit:
            assert fills and fills[0].find(qn('a:srgbClr')).get('val') == '000000'
        else:
            assert lines[0].find(qn('a:noFill')) is not None


def test_the_doughnuts_are_native_with_the_pages_palette_and_the_hole():
    """Five doughnut charts, slice colours by the alphabetical palette rule,
    a 50% hole written where python-pptx has no property (D137)."""
    results, implementation = _goldenCase()
    model = _model(results, implementation)
    items = [item for group in model['groups'] for item in group['items']]
    prs = Presentation(io.BytesIO(_deck()))
    chartSlide = prs.slides[3]                    # after the implementation
    charts = [shape.chart for shape in chartSlide.shapes if shape.has_chart]
    assert len(charts) == 5
    for chart in charts:
        assert chart.chart_type == XL_CHART_TYPE.DOUGHNUT
        holes = chart._chartSpace.findall('.//' + qn('c:holeSize'))
        assert holes and all(hole.get('val') == '50' for hole in holes)
    styleSlices = donutBreakdown(items, 'style')
    points = charts[0].series[0].points
    got = [str(points[index].format.fill.fore_color.rgb)
           for index in range(len(styleSlices))]
    assert got == [DONUT_PALETTE[entry['slot']] for entry in styleSlices]


def test_the_doughnuts_print_no_shares_on_their_slices():
    """D137, D138. Each chart's title and legend are 12pt; no slice carries
    a label (D137 printed each share, D138 took them off); the chart's data
    still holds the printed shares, as percentages to one place."""
    results, implementation = _goldenCase()
    items = [item for group in _model(results, implementation)['groups']
             for item in group['items']]
    prs = Presentation(io.BytesIO(_deck()))
    charts = [shape.chart for shape in prs.slides[3].shapes if shape.has_chart]
    assert len(charts) == 5
    for chart in charts:
        assert chart.chart_title.text_frame.paragraphs[0].runs[0].font.size.pt == 12
        assert chart.legend.font.size.pt == 12
        assert not chart.plots[0].has_data_labels
        assert not chart._chartSpace.findall('.//' + qn('c:dLbl'))
    series = charts[0].series[0]
    assert list(series.values) == pytest.approx(
        [e['pct'] / 100.0 for e in donutBreakdown(items, 'style')])


def _implementationPlan(includeFees=True):
    (entry,) = [e for e in _plan(includeFees=includeFees) if e.get('heading') == 'Implemented Model']
    (tp,) = entry['tables']
    return tp


@pytest.mark.parametrize('includeFees', [True, False])
def test_the_implementation_table_spans_the_slide_with_its_names_at_their_width(includeFees):
    """D137, D138. On the slide, the implementation table spans the full
    width, fees or not, and so sits centred: the category and product
    columns exactly as wide as their longest names, every other column its
    widest entry plus an equal share of the width left over; nothing on a
    second line, the heads included; every row one height; a little under
    the slide's rule; each head aligned as the figures or words beneath it."""
    tp = _implementationPlan(includeFees)
    assert tp.natural and tp.oneLine
    widths = tp.widths()
    needs = pptWriter._columnNeeds(tp, tp.font)
    assert sum(widths) == pytest.approx(tp.box[2]), 'the full width'
    assert tp.box[0] == pptWriter.FULL_BOX[0] and tp.box[2] == pptWriter.FULL_BOX[2]
    assert widths[:2] == needs[:2]
    extra = {round(w - n, 6) for w, n in zip(widths[2:], needs[2:])}
    assert len(extra) == 1 and extra.pop() >= 0, 'the rest shared equally'
    assert len(set(tp.heights)) == 1, 'every row one height'
    assert tp.box[1] == pytest.approx(pptWriter.FULL_BOX[1] + pptWriter._TABLE_DROP_PT)
    doc, base = tp.doc, pptWriter._base(tp.doc)
    for column in (1, 2):
        longest = max(pptWriter._oneLineNeed(doc.rows[r].cells[column], doc, tp.font, column)
                      for r in tp.rows if column in doc.rows[r].cells
                      and doc.rows[r].cells[column].value)
        assert widths[column - 1] == pytest.approx(longest)
    for r in tp.rows:
        for column, spec in doc.rows[r].cells.items():
            if spec.value in (None, ''):
                continue
            text = sheetDoc.renderNumber(spec.value, spec.fmt)
            font = spec.font or pptWriter._DEFAULT_FONT
            size = tp.font * font.get('size', base) / base
            inner = (widths[column - 1] - 2 * pptWriter._PAD_H_EM * tp.font
                     - pptWriter._indentPt(spec, text, tp.font, column, doc))
            assert pptWriter.textWidth(text.strip(), font['name'], bool(font.get('bold')),
                                       size) <= inner, (r, text)
    head = doc.rows[tp.header[-1]].cells
    body = [r for r in tp.rows if r not in tp.header]
    for column, cell in head.items():
        below = {(doc.rows[r].cells[column].align or {}).get('horizontal', 'left')
                 for r in body if column in doc.rows[r].cells}
        assert below == {cell.align['horizontal']}, cell.value

    prs = Presentation(io.BytesIO(_deck(includeFees=includeFees)))
    table = next(t for s in prs.slides for t in _tables(s)
                 if any(c.text == 'Categories & Asset Classes' for row in list(t.rows)[:2]
                        for c in row.cells))
    for row in table.rows:
        for cell in row.cells:
            assert cell.text_frame.word_wrap is False
    heads = list(table.rows)[len(tp.header) - 1].cells
    assert heads[0].text_frame.paragraphs[0].alignment == pptWriter.PP_ALIGN.LEFT
    notional = [c.text for c in heads].index('Notional')
    assert heads[notional].text_frame.paragraphs[0].alignment == pptWriter.PP_ALIGN.RIGHT


def test_the_assumptions_table_wraps_nothing_and_sits_under_the_rule():
    """D137. The assumptions table keeps every entry on one line - its two
    two-line heads read as one - across the slide's width, each column at
    least its widest entry; every row is one height, and no taller than its
    text needs (D138); its asset names step in half as far as the risk
    table's; and it moves a little down under the slide's rule where the
    slide has the room."""
    plans = [tp for e in _plan() for tp in e.get('tables', []) if tp.doc.name == 'assumptions']
    assert plans
    for tp in plans:
        assert tp.oneLine and not tp.natural
        widths = tp.widths()
        needs = pptWriter._columnNeeds(tp, tp.font)
        assert all(w >= n - 1e-6 for w, n in zip(widths, needs))
        assert sum(widths) == pytest.approx(tp.box[2])
        assert len(set(tp.heights)) == 1
        assert pptWriter._measure(tp, tp.font)[1]
        tight, _ = pptWriter._measure(tp, tp.font)
        room = pptWriter.FULL_BOX[3] - sum(tight)
        assert tp.box[1] == pytest.approx(pptWriter.FULL_BOX[1] + min(pptWriter._TABLE_DROP_PT, room))
        # and its rows at their own height, not opened into the slide (D138)
        assert tp.heights == tight
        assert sum(tp.heights) <= tp.box[3] + 1e-6
    heads = [tp.doc.rows[2].cells[c].value for c in (2, 7)]
    assert all('\n' in head for head in heads), 'the sheet keeps its two lines'
    prs = Presentation(io.BytesIO(_deck()))
    table = next(t for s in prs.slides for t in _tables(s)
                 if any('Long-Term Estimates' in c.text for c in t.rows[0].cells))
    texts = [c.text for c in table.rows[1].cells]
    assert 'Risk Premia with Estimated Range' in texts
    assert 'Estimated Mean Return (2.5% Risk Free Rate)' in texts
    assert pptWriter._INDENT_EM_BY_DOC['assumptions'] == pptWriter._INDENT_EM / 2


def test_a_one_line_table_fills_between_its_needs_and_its_shares():
    """D137. The widths of a one-line table that fills its box: a column
    that needs more than its share holds its need, the others share what is
    left in proportion, and the whole is the box's width."""
    widths = pptWriter._fillWidths([50.0, 10.0, 10.0], [1.0, 1.0, 2.0], 200.0)
    assert widths[0] == 50.0 and widths[2] == pytest.approx(2 * widths[1])
    assert sum(widths) == pytest.approx(200.0)
    assert pptWriter._fillWidths([150.0, 80.0], [1.0, 1.0], 200.0) == [150.0, 80.0]


def test_the_deck_is_stamped_like_a_workbook_and_a_draft_is_not():
    """dc:identifier carries the UID, so stampedProposalId reads a deck
    exactly as it reads a workbook; a deck with no delivered proposal says
    Draft in every footer and stamps nothing."""
    deck = _deck()
    assert stampedProposalId(deck) == UID
    prs = Presentation(io.BytesIO(deck))
    assert prs.core_properties.title == 'PMG Proposal {} - USD Hedged'.format(UID)

    draft = _deck(proposalId=None)
    assert stampedProposalId(draft) is None
    drafted = Presentation(io.BytesIO(draft))
    texts = ' '.join(shape.text_frame.text for slide in drafted.slides
                     for shape in slide.shapes if shape.has_text_frame)
    assert 'Draft — no delivered proposal for this scenario' in texts
    assert UID not in texts


def test_the_fee_columns_leave_the_deck_with_the_fees():
    """includeFees=False: the two fee columns are absent from the slide as
    they are absent from the sheet, because implColumns drives both."""
    prs = Presentation(io.BytesIO(_deck(includeFees=False)))
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_table:
                continue
            rows = list(shape.table.rows)
            # the header, under the Long-term / Initial row a private-markets
            # book carries above it (D136)
            header = next(([cell.text for cell in row.cells] for row in rows[:2]
                           if 'Products' in [cell.text for cell in row.cells]), [])
            if 'Products' in header:
                assert 'Mgmt fee' not in header
                assert 'Wtd fee (bp)' not in header
                assert 'Initial Wtd fee (bp)' not in header, 'the fee twin leaves with it'
                assert 'Product Cost' in header, 'product cost stays (D52)'
                assert 'Initial (%)' in header and 'Initial Notional' in header
                return
    pytest.fail('no implementation table found')


def test_the_cover_is_a_flag():
    withCover = Presentation(io.BytesIO(_deck()))
    without = Presentation(io.BytesIO(_deck(cover=False)))
    assert len(withCover.slides) == len(without.slides) + 1
    assert not any(shape.has_table for shape in withCover.slides[0].shapes)
    assert any(shape.has_table for shape in without.slides[0].shapes)


# --------------------------------------------------------------------- #
# One delivery, both files, both locked (D123).
# --------------------------------------------------------------------- #

def test_one_export_delivers_both_files_or_neither(monkeypatch):
    """Over HTTP, the D123 story. The export refuses (both files) where the
    workbook always refused; delivered, it is ONE zip holding the workbook
    and the deck, both stamped with one UID, both recorded; the admin
    register hands each back byte for byte; and there is no deck-only route
    left to take one without the other."""
    import zipfile
    from fastapi.testclient import TestClient
    from cyrus_pmg.pmgService.isgPMGService import app
    from cyrus_pmg.pmgService.scenario import proposalRegister, sleeves
    from cyrus_pmg.pmgService.scenario.sleeves import listSleeves
    from cyrus_pmg.pmgService.scenario.types import PortfolioKey
    monkeypatch.setenv('PMG_ALLOWED_KERBEROS', 'bob,alice')
    monkeypatch.setenv('PMG_ADMIN_KERBEROS', 'alice')
    client = TestClient(app)
    pwa, admin = {'X-Kerberos': 'bob'}, {'X-Kerberos': 'alice'}

    made = client.post('/api/v1/scenario', headers=pwa, json={
        'mandate': {'topAccountSize': 1e9, 'mandateSize': 1e9,
                    'primaryPwa': 'A. Castellanos — Madrid'},
        'basis': {'currency': 'USD', 'hedging': 'Hedged'}})
    assert made.status_code == 200, made.text
    sid = made.json()['id']
    assert client.put('/api/v1/scenario/' + sid, headers=pwa,
                      json={'variant': sleeves.VARIANTS[0]}).status_code == 200
    key = {'currency': 'USD', 'riskLevel': 'Moderate', 'allocationType': 'Full',
           'excludeRealAssets': False}
    resolved = client.post('/api/v1/scenario/' + sid + '/portfolio', headers=pwa,
                           json={'key': key, 'role': 'base'})
    assert resolved.status_code == 200, resolved.text

    # refused as ever: no sleeves - and nothing half-delivered
    refused = client.post('/api/v1/scenario/' + sid + '/export', headers=pwa)
    assert refused.status_code == 422
    assert 'sleeve' in refused.json()['error'].lower()

    chosen = {}
    for category in resolved.json()['portfolio']['categories']:
        if category['name'] in rules.AUTO_SLEEVE_CATEGORIES:
            continue
        under = rules.sleeveCategory(category['name'])
        if under not in chosen:
            chosen[under] = listSleeves(under, sleeves.VARIANTS[0],
                                        PortfolioKey.fromDict(key))[0]['name']
    assert client.put('/api/v1/scenario/' + sid, headers=pwa,
                      json={'sleeves': chosen}).status_code == 200

    delivered = client.post('/api/v1/scenario/' + sid + '/export', headers=pwa)
    assert delivered.status_code == 200, delivered.text
    assert delivered.headers['content-type'] == 'application/zip'
    uid = delivered.headers['x-proposal-id']
    assert delivered.headers['content-disposition'].endswith('_{}.zip"'.format(uid))
    archive = zipfile.ZipFile(io.BytesIO(delivered.content))
    byKind = {name.rsplit('.', 1)[1]: archive.read(name) for name in archive.namelist()}
    assert sorted(byKind) == ['pptx', 'xlsx'], 'both files, and only them'
    assert stampedProposalId(byKind['xlsx']) == uid
    assert stampedProposalId(byKind['pptx']) == uid

    # the register holds both, and the admin console hands both back
    for leaf, kind in (('workbook', 'xlsx'), ('deck', 'pptx')):
        fetched = client.get('/api/v1/scenario/repository/proposals/{}/{}'.format(uid, leaf),
                             headers=admin)
        assert fetched.status_code == 200, (leaf, fetched.text)
        assert fetched.content == byKind[kind], leaf
        assert client.get('/api/v1/scenario/repository/proposals/{}/{}'.format(uid, leaf),
                          headers=pwa).status_code == 403
    assert proposalRegister.getProposal(uid)['deckBytes'] == len(byKind['pptx'])

    # there is no way to take the deck alone
    assert client.post('/api/v1/scenario/' + sid + '/export.pptx',
                       headers=pwa).status_code in (404, 405)


def test_the_deck_is_locked_like_the_workbook():
    """D123: the deck carries a password to modify - the workbook's own
    PA55WORD, verified by recomputing the ISO write-protection hash - sits in
    schema order after defaultTextStyle, is marked final, and never carries
    the plaintext. Every deck, draft or delivered."""
    import base64
    import re
    import zipfile
    from cyrus_pmg.pmgService.scenario.workbook import SHEET_PASSWORD
    for deck in (_deck(), _deck(proposalId=None)):
        archive = zipfile.ZipFile(io.BytesIO(deck))
        presentation = archive.read('ppt/presentation.xml').decode('utf-8')
        found = re.search(r'<p:modifyVerifier ([^>]*)/>', presentation)
        assert found, 'a password to modify'
        attrs = dict(re.findall(r'(\w+)="([^"]*)"', found.group(1)))
        assert (attrs['cryptAlgorithmSid'], attrs['spinCount']) == ('14', '100000')
        salt = base64.b64decode(attrs['saltData'])
        stored = base64.b64decode(attrs['hashData'])
        assert pptWriter.modifyHash(SHEET_PASSWORD, salt) == stored
        assert pptWriter.modifyHash('pa55word', salt) != stored, 'case matters'
        tags = [t for t in re.findall(r'<(p:[A-Za-z]+)', presentation) if t.count(':') == 1]
        at = tags.index('p:modifyVerifier')
        assert tags[at - 1] == 'p:defaultTextStyle', 'schema order'
        custom = archive.read('docProps/custom.xml').decode('utf-8')
        assert 'name="_MarkAsFinal"' in custom and '<vt:bool>true</vt:bool>' in custom
        assert 'custom-properties' in archive.read('_rels/.rels').decode('utf-8')
        assert '/docProps/custom.xml' in archive.read('[Content_Types].xml').decode('utf-8')
        assert Presentation(io.BytesIO(deck)).core_properties.content_status == 'Final'
        for plain in (SHEET_PASSWORD.encode('utf-8'), SHEET_PASSWORD.encode('utf-16-le')):
            assert plain not in deck, 'only the hash reaches the file'
        # each deck salts afresh, so two locks never share a hash
    salts = set()
    for _ in range(2):
        text = zipfile.ZipFile(io.BytesIO(_deck())).read('ppt/presentation.xml').decode()
        salts.add(re.search(r'saltData="([^"]+)"', text).group(1))
    assert len(salts) == 2


def test_the_page_offers_one_download_and_awaits_its_writes():
    """The front end, pinned at the source (D122, D123): every
    fire-and-forget scenario write is tracked, the one export runs behind the
    save barrier, and the card offers a single button for both files."""
    root = os.path.join(HERE, '..', '..', 'generator')
    with open(os.path.join(root, 'js', 'core.js'), encoding='utf-8') as handle:
        core = handle.read()
    assert core.count("trackWrite(apiFetch('/scenario/'"
                      " + encodeURIComponent(state.scenarioId), {") == 6, \
        'six fire-and-forget writes, each tracked'
    assert 'writesSettled: writesSettled,' in core
    with open(os.path.join(root, 'js', 'implementation.js'), encoding='utf-8') as handle:
        impl = handle.read()
    assert 'await App.writesSettled();' in impl
    assert "exportProposal(); return;" in impl
    assert 'Download proposal' in impl
    for gone in ('implexportppt', 'Download PowerPoint', "'.pptx'", 'exportFile('):
        assert gone not in impl, gone
    assert 'Assumptions and ' in impl, 'the card names all four sheets (A7)'
