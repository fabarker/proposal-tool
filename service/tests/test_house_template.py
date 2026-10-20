"""The deck built in a template (D174).

The deck opens the template SCENARIO_DECK_TEMPLATE names, adds its slides
from the template's layout, and puts them where the template's marks say.
These hold it to that: the placeholder template end to end, and small
templates made here from it - with slides of their own, fields, conditions,
sections and a sensitivity label - for what a house template may carry.
"""

import io
import json
import os
import zipfile

import pytest
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import Part
from pptx.opc.packuri import PackURI
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from cyrus_pmg.pmgService.scenario import assetEstimates, houseTemplate, pptWriter, rules
from cyrus_pmg.pmgService.scenario.houseTemplate import TemplateError
from cyrus_pmg.pmgService.scenario.pptWriter import writeDeck
from cyrus_pmg.pmgService.scenario.types import BasisInput, MandateInput
from cyrus_pmg.pmgService.scenario.workbook import buildImplementationRows, stampedProposalId

HERE = os.path.dirname(os.path.abspath(__file__))
GOLDEN = os.path.join(HERE, 'golden')
PLACEHOLDER = os.path.join(HERE, '..', '..', 'deckTemplate', 'ProposalDeck_Template.pptx')
UID = 'pr_0123456789ab'
BASIS = BasisInput(currency='USD', hedging='Hedged')
MANDATE = MandateInput(topAccountSize=10_000_000.0, mandateSize=25_000_000.0,
                       primaryPwa='J. Mercer')
_SECTIONS_URI = '{521415D9-36F7-43E2-AB2F-B90AF26B5E84}'
#: the charges a PWA enters for the Fees Summary slide (D177)
CHARGES = {'fxHighEdge': 5e6, 'fxLowEdge': 1e6, 'fxHighBps': 10, 'fxMidBps': 10, 'fxLowBps': 20,
           'equityTradeBps': 35, 'custodyFee': 5000}
_P14 = 'http://schemas.microsoft.com/office/powerpoint/2010/main'


def _golden():
    with open(os.path.join(GOLDEN, 'engine_USD_Hedged_2col.results.json'), encoding='utf-8') as f:
        results = json.load(f)
    with open(os.path.join(GOLDEN, 'engine_USD_Hedged_2col.impl.json'), encoding='utf-8') as f:
        implementation = json.load(f)
    return results, implementation


def _deck(includeFees=True, cover=True, charges=CHARGES, schedule=None):
    results, implementation = _golden()
    return writeDeck(BASIS, MANDATE, results, implementation['sleeves'],
                     rules.AUTO_SLEEVE_CATEGORIES, implementation['variant'],
                     implementation['tacticalTilt'], schedule or implementation['feeSchedule'],
                     implementation['feeLevel'], includeFees, implementation['volPremium'],
                     assets=assetEstimates.forSlice('USD', 'Hedged'), proposalId=UID, cover=cover,
                     otherCharges=charges)


def _model(includeFees=True):
    results, implementation = _golden()
    return buildImplementationRows(
        results[0], implementation['sleeves'], rules.AUTO_SLEEVE_CATEGORIES,
        MANDATE.mandateSize, implementation['variant'], implementation['tacticalTilt'],
        implementation['feeSchedule'] if includeFees else None, implementation['feeLevel'],
        MANDATE.topAccountSize, implementation['volPremium'], 'USD')


def _titles(prs):
    out = []
    for slide in prs.slides:
        title = next((ph for ph in slide.placeholders
                      if ph.placeholder_format.type is not None
                      and 'TITLE' in str(ph.placeholder_format.type)), None)
        out.append(title.text_frame.text if title is not None else None)
    return out


def _notes(slide):
    if not slide.has_notes_slide:
        return ''
    return slide.notes_slide.notes_text_frame.text


def _ownSlide(prs, title, notes=None):
    """A slide of the template's own, from its layout: a title and a text
    box, and *notes* in its speaker notes."""
    layout = next(layout for layout in prs.slide_layouts if layout.name == 'Content')
    slide = prs.slides.add_slide(layout)
    for ph in list(slide.placeholders):
        if ph.placeholder_format.idx == 0:
            ph.text = title
        else:
            ph._element.getparent().remove(ph._element)
    if notes is not None:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def _textbox(slide, name, runs):
    """A text box whose one paragraph is *runs*: (text, bold, colour)."""
    box = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(6), Inches(0.5))
    box.name = name
    paragraph = box.text_frame.paragraphs[0]
    for text, bold, colour in runs:
        run = paragraph.add_run()
        run.text = text
        run.font.bold = bold
        if colour:
            run.font.color.rgb = RGBColor.from_string(colour)
    return box


def _template(tmp_path, edit=None, name='template.pptx'):
    """The placeholder template, changed by *edit*, saved as a file."""
    prs = Presentation(PLACEHOLDER)
    if edit:
        edit(prs)
    path = str(tmp_path / name)
    prs.save(path)
    return path


def _moveSlide(prs, fromIndex, toIndex):
    ids = prs.slides._sldIdLst
    element = list(ids)[fromIndex]
    ids.remove(element)
    ids.insert(toIndex, element)


def _dropMarker(prs, mark):
    for element, slide in zip(list(prs.slides._sldIdLst), list(prs.slides)):
        if _notes(slide).strip() == 'tool: ' + mark:
            rId = element.get(qn('r:id'))
            prs.slides._sldIdLst.remove(element)
            prs.part.drop_rel(rId)


@pytest.fixture
def placeholder(monkeypatch):
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, PLACEHOLDER)
    monkeypatch.delenv(houseTemplate.LAYOUT_SETTING, raising=False)
    return PLACEHOLDER


# ---- the placeholder, end to end -------------------------------------------------

def test_the_check_reads_the_placeholder_clean():
    found = houseTemplate.report(PLACEHOLDER)
    assert found['problems'] == []
    assert found['layout'] == 'Content'
    assert {role: ph['idx'] for role, ph in found['roles'].items()} == {
        'title': 0, 'subtitle': 10, 'content': 11, 'footnotes': 12, 'page': 14}
    assert found['size'] == (720.0, 540.0)                     # 10 x 7.5 in: 4:3
    assert [s['marker'] for s in found['slides'][1:] if s['marker']] == list(houseTemplate.SLIDES)
    own = [s for s in found['slides'] if not s['marker'] and not s['templateOnly']]
    assert [(s['conditions'], s['areas']) for s in own] == [
        (['fees'], []), (['fees'], []), (['casp'], ['fee-ladder'])]
    assert found['slides'][0]['templateOnly']
    assert found['unmarked'] == []
    assert found['theme']['faces'] == {'headings': 'GS Sans', 'body': 'GS Sans'}


def test_without_the_setting_the_deck_is_drawn_as_before(monkeypatch):
    monkeypatch.delenv(houseTemplate.TEMPLATE_SETTING, raising=False)
    prs = Presentation(io.BytesIO(_deck()))
    assert round(prs.slide_width / 914400.0, 3) == pptWriter.SLIDE_W_IN


def test_the_deck_is_built_in_the_template_in_the_marks_order(placeholder):
    data = _deck()
    prs = Presentation(io.BytesIO(data))
    template = Presentation(PLACEHOLDER)
    assert (prs.slide_width, prs.slide_height) == (template.slide_width, template.slide_height)
    assert _titles(prs) == [
        'Investment Proposal', 'Strategic Asset Allocation', 'Strategic Asset Allocation',
        'Implemented Model', 'Composition of the Implemented Model', 'Fees Summary',
        'Portfolio Fee Structure', 'Fee Summary', 'Advisory Fee Schedule',
        'Long-Term Capital Market Assumptions']
    for slide in prs.slides:
        # every slide from the template's layout; no marker, no style slide
        assert slide.slide_layout.name == 'Content'
        assert not _notes(slide).strip()
        assert not any(shape.name.startswith('Style ·') for shape in slide.shapes)
    assert stampedProposalId(data) == UID


def test_placeholder_text_carries_no_formatting_of_its_own(placeholder):
    prs = Presentation(io.BytesIO(_deck()))
    for slide in prs.slides:
        for ph in slide.placeholders:
            for run in ph._element.iter(qn('a:rPr')):
                assert set(run.keys()) <= {'lang', 'dirty'}
                assert len(run) == 0                    # no fill, no face


def test_every_slide_has_its_page_number_and_no_uid(placeholder):
    prs = Presentation(io.BytesIO(_deck()))
    for number, slide in enumerate(prs.slides, 1):
        byIdx = {ph.placeholder_format.idx: ph for ph in slide.placeholders}
        assert 13 not in byIdx
        assert not any(UID in (sh.text_frame.text if sh.has_text_frame else '')
                       for sh in slide.shapes)
        fields = [f for f in byIdx[14]._element.iter(qn('a:fld'))]
        assert [f.get('type') for f in fields] == ['slidenum']
        assert fields[0].find(qn('a:t')).text == str(number)


def test_a_slide_has_a_short_subtitle_and_its_long_line_opens_the_footnotes(placeholder):
    prs = Presentation(io.BytesIO(_deck()))
    allocation = prs.slides[1]
    byIdx = {ph.placeholder_format.idx: ph for ph in allocation.placeholders}
    assert byIdx[10].text_frame.text == 'Comparing Allocations'
    assert byIdx[12].text_frame.paragraphs[0].text.startswith('USD · Hedged — ')


def test_the_tables_fit_the_templates_content_box_and_name_its_body_face(placeholder):
    template = houseTemplate.Template(PLACEHOLDER)
    left, top, width, height = template.box
    prs = Presentation(io.BytesIO(_deck()))
    tables = [shape for slide in prs.slides for shape in slide.shapes if shape.has_table]
    assert tables
    for shape in tables:
        assert shape.left / 12700.0 >= left - 0.01
        assert (shape.left + shape.width) / 12700.0 <= left + width + 0.01
        assert (shape.top + shape.height) / 12700.0 <= top + height + 0.01
        faces = {latin.get('typeface') for latin in shape._element.iter(qn('a:latin'))}
        assert faces == {'+mn-lt'}


def test_the_templates_theme_and_fonts_are_left_alone(placeholder):
    prs = Presentation(io.BytesIO(_deck()))
    template = Presentation(PLACEHOLDER)
    theme = prs.slide_master.part.part_related_by(RT.THEME).blob
    assert theme == template.slide_master.part.part_related_by(RT.THEME).blob
    fonts = prs.part._element.find(qn('p:embeddedFontLst'))
    original = template.part._element.find(qn('p:embeddedFontLst'))
    assert len(fonts) == len(original)                  # no second copy of a face


def test_the_deck_is_locked_in_the_template_too(placeholder):
    data = _deck()
    prs = Presentation(io.BytesIO(data))
    assert prs.part._element.find(qn('p:modifyVerifier')) is not None
    custom = prs.part.package.part_related_by(RT.CUSTOM_PROPERTIES)
    assert b'_MarkAsFinal' in custom.blob
    assert prs.core_properties.content_status == 'Final'


def test_the_cover_is_made_only_at_its_marker(tmp_path, monkeypatch):
    path = _template(tmp_path, lambda prs: _dropMarker(prs, 'cover'))
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, path)
    prs = Presentation(io.BytesIO(_deck()))
    assert 'Investment Proposal' not in _titles(prs)


def test_a_template_with_gaps_in_its_slide_files_makes_a_sound_file(tmp_path, monkeypatch):
    # the cover's marker out: the template's slide files now skip a number,
    # and a new slide named by the count would take an existing one's name
    path = _template(tmp_path, lambda prs: _dropMarker(prs, 'cover'))
    with zipfile.ZipFile(path) as archive:
        assert 'ppt/slides/slide2.xml' not in archive.namelist()
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, path)
    with zipfile.ZipFile(io.BytesIO(_deck())) as archive:
        names = archive.namelist()
    assert len(names) == len(set(names))


# ---- the marks -----------------------------------------------------------------------

def test_the_templates_own_slides_stay_and_the_tool_goes_where_marked(tmp_path, monkeypatch):
    def edit(prs):
        intro = _ownSlide(prs, 'About us', 'Speak to the history.\ntool: when fees')
        _textbox(intro, 'Words', [('We have been here since 1869.', False, None)])
        _moveSlide(prs, len(prs.slides) - 1, 2)          # after the cover's marker
        _ownSlide(prs, 'Closing thoughts')
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, _template(tmp_path, edit))
    prs = Presentation(io.BytesIO(_deck()))
    titles = _titles(prs)
    assert titles[:3] == ['Investment Proposal', 'About us', 'Strategic Asset Allocation']
    assert titles[-1] == 'Closing thoughts'
    about = prs.slides[1]
    assert [s.name for s in about.shapes if s.name == 'Words']
    assert next(s for s in about.shapes if s.name == 'Words').text_frame.text == \
        'We have been here since 1869.'
    # the tool's own line comes off the notes; the presenter's stays
    assert _notes(about).strip() == 'Speak to the history.'


def test_a_slide_whose_condition_fails_is_taken_out(tmp_path, monkeypatch):
    def edit(prs):
        _ownSlide(prs, 'Fees at a glance', 'tool: when fees')
        _ownSlide(prs, 'CASP only', 'tool: when casp')
        _ownSlide(prs, 'Private markets', 'tool: when private markets')
        _ownSlide(prs, 'Never', 'tool: template only')
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, _template(tmp_path, edit))
    priced = _titles(Presentation(io.BytesIO(_deck())))
    assert priced[-3:] == ['Fees at a glance', 'CASP only', 'Private markets']
    unpriced = _titles(Presentation(io.BytesIO(_deck(includeFees=False))))
    assert unpriced[-1] == 'Private markets'
    assert 'Fees at a glance' not in unpriced and 'CASP only' not in unpriced
    assert 'Never' not in priced + unpriced


def test_the_marks_order_is_the_decks_order(tmp_path, monkeypatch):
    # the assumptions' marker moved up to sit before the allocation's
    path = _template(tmp_path, lambda prs: _moveSlide(prs, 10, 2))
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, path)
    titles = _titles(Presentation(io.BytesIO(_deck())))
    assert titles[1] == 'Long-Term Capital Market Assumptions'
    assert titles[2] == 'Strategic Asset Allocation'


def test_a_tool_slide_with_no_marker_goes_at_the_end(tmp_path, monkeypatch):
    def edit(prs):
        _ownSlide(prs, 'Appendix')
        _dropMarker(prs, 'composition')
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, _template(tmp_path, edit))
    titles = _titles(Presentation(io.BytesIO(_deck())))
    assert titles[-2:] == ['Appendix', 'Composition of the Implemented Model']
    assert houseTemplate.report(str(tmp_path / 'template.pptx'))['unmarked'] == ['composition']


def test_a_marker_that_does_not_apply_takes_its_slides_out_with_it(tmp_path, monkeypatch):
    def edit(prs):
        for slide in prs.slides:
            if _notes(slide).strip() == 'tool: composition':
                slide.notes_slide.notes_text_frame.text = 'tool: composition\ntool: when casp'
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, _template(tmp_path, edit))
    titles = _titles(Presentation(io.BytesIO(_deck(includeFees=False))))
    assert 'Composition of the Implemented Model' not in titles
    assert titles[-1] == 'Long-Term Capital Market Assumptions'


def test_fields_take_their_own_formatting_even_across_runs(tmp_path, monkeypatch):
    def edit(prs):
        slide = _ownSlide(prs, 'Fee Summary', 'tool: when fees')
        _textbox(slide, 'Total', [('Total cost ', False, None), ('{{tot', True, 'FFFFFF'),
                                  ('alCost}}', False, None), (' a year', False, None)])
        _textbox(slide, 'Mandate', [('Based on a {{mandateSize}} portfolio for {{primaryPwa}}',
                                     False, None)])
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, _template(tmp_path, edit))
    prs = Presentation(io.BytesIO(_deck()))
    slide = prs.slides[len(prs.slides) - 1]
    total = next(s for s in slide.shapes if s.name == 'Total').text_frame.paragraphs[0]
    expected = pptWriter.fieldValues(_model(), MANDATE, UID)['totalCost']
    assert total.text == 'Total cost {} a year'.format(expected)
    figure = next(r for r in total.runs if r.text == expected)
    assert figure.font.bold and str(figure.font.color.rgb) == 'FFFFFF'
    mandate = next(s for s in slide.shapes if s.name == 'Mandate').text_frame.text
    assert mandate == 'Based on a $25m portfolio for J. Mercer'


def test_the_fee_fields_are_the_models_own_figures():
    model = _model()
    values = pptWriter.fieldValues(model, MANDATE, UID)
    items = [i for g in model['groups'] for i in g['items']]
    fee = sum(i['printedPct'] * i['managementFee'] for i in items) / 100.0
    cost = sum(i['printedPct'] * float(i['productCost']) for i in items) / 100.0
    assert values['advisoryFee'] == '{:.2f}%'.format(fee)
    assert values['productCost'] == '{:.2f}%'.format(cost)
    assert values['totalCost'] == values['weightedCost'] == '{:.2f}%'.format(
        model['total']['wtdFeeBp'] / 100.0)
    assert values['initialCost'] == '{:.2f}%'.format(model['initial']['total']['wtdFeeBp'] / 100.0)
    assert values['mandateSize'] == '$25m' and 'proposalUid' not in values
    unpriced = pptWriter.fieldValues(_model(includeFees=False), MANDATE, UID, priced=False)
    assert 'totalCost' not in unpriced and 'advisoryFee' not in unpriced


def test_a_field_with_no_figure_refuses_the_export(placeholder):
    # the PWA has entered every charge but the equity one: Fees Summary
    # quotes it, so the deck is refused with the slide and the charge named
    short = dict(CHARGES)
    short.pop('equityTradeBps')
    with pytest.raises(TemplateError) as raised:
        _deck(charges=short)
    assert 'Slide 7' in raised.value.message and '{{equityTradeBps}}' in raised.value.message
    assert 'the PWA has not entered the equity transaction charge' in raised.value.message
    # the custody fee stands at its default, the trade sizes are fixed
    _deck(charges={k: v for k, v in CHARGES.items()
                   if k not in ('custodyFee', 'fxHighEdge', 'fxLowEdge')})
    # unpriced, Fees Summary is not delivered and nothing is asked for
    _deck(includeFees=False, charges={})


@pytest.mark.parametrize('notes,words', [
    ('tool: comparing allocations', 'not a slide the tool makes'),
    ('tool: when rdr', 'condition the tool does not know'),
])
def test_a_mark_the_tool_does_not_know_is_refused_with_its_slide(tmp_path, monkeypatch, notes, words):
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING,
                       _template(tmp_path, lambda prs: _ownSlide(prs, 'Odd', notes)))
    with pytest.raises(TemplateError) as raised:
        _deck()
    assert 'Slide 12' in raised.value.message and words in raised.value.message


def test_an_unknown_field_or_area_is_refused(tmp_path):
    def edit(prs):
        slide = _ownSlide(prs, 'Odd')
        _textbox(slide, 'Field', [('{{advisoryFees}}', False, None)])
        _textbox(slide, 'tool: fee ladder', [('', False, None)])
    found = houseTemplate.report(_template(tmp_path, edit))
    assert any('{{advisoryFees}}' in p for p in found['problems'])
    assert any('"tool: fee ladder"' in p for p in found['problems'])


def test_an_area_with_nothing_to_draw_takes_its_slide_out(tmp_path, monkeypatch):
    # the schedule's own condition off: priced by RDR, its tier table has
    # nothing to draw, and the slide goes all the same
    def edit(prs):
        for slide in prs.slides:
            if any(shape.name == 'tool: fee-ladder' for shape in slide.shapes):
                slide.notes_slide.notes_text_frame.text = 'For the presenter.'
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, _template(tmp_path, edit))
    results, implementation = _golden()
    rdr = writeDeck(BASIS, MANDATE, results, implementation['sleeves'],
                    rules.AUTO_SLEEVE_CATEGORIES, implementation['variant'],
                    implementation['tacticalTilt'], 'RDR', implementation['feeLevel'], True,
                    implementation['volPremium'], assets=assetEstimates.forSlice('USD', 'Hedged'),
                    proposalId=UID, otherCharges=CHARGES)
    assert 'Advisory Fee Schedule' not in _titles(Presentation(io.BytesIO(rdr)))
    assert 'Advisory Fee Schedule' in _titles(Presentation(io.BytesIO(_deck())))


# ---- what a house template may carry ----------------------------------------------------

def _addSections(prs):
    """Two sections: the cover's marker alone, then everything else."""
    ids = [e.get('id') for e in prs.slides._sldIdLst]
    sections = ''.join(
        '<p14:section name="{}" id="{{{}}}"><p14:sldIdLst>{}</p14:sldIdLst></p14:section>'.format(
            name, guid, ''.join('<p14:sldId id="{}"/>'.format(i) for i in members))
        for name, guid, members in (('Opening', '6A0A3A57-0000-4000-8000-000000000001', ids[:2]),
                                    ('Body', '6A0A3A57-0000-4000-8000-000000000002', ids[2:])))
    root = prs.part._element
    extLst = root.find(qn('p:extLst'))
    if extLst is None:
        extLst = etree.SubElement(root, qn('p:extLst'))
    extLst.append(parse_xml(
        '<p:ext xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" uri="{}">'
        '<p14:sectionLst xmlns:p14="{}">{}</p14:sectionLst></p:ext>'.format(
            _SECTIONS_URI, _P14, sections)))


def test_the_templates_sections_stay_true(tmp_path, monkeypatch):
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, _template(tmp_path, _addSections))
    prs = Presentation(io.BytesIO(_deck()))
    slideIds = [e.get('id') for e in prs.slides._sldIdLst]
    sections = prs.part._element.findall('.//{%s}section' % _P14)
    members = [[e.get('id') for e in s.iter('{%s}sldId' % _P14)] for s in sections]
    assert sum(members, []) == slideIds                  # every slide once, in order
    assert members[0] == slideIds[:1]                    # the style slide left; the cover stayed


def test_a_templates_own_custom_properties_are_kept(tmp_path, monkeypatch):
    label = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<Properties '
             'xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
             'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
             '<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" '
             'name="Classification"><vt:lpwstr>Internal</vt:lpwstr></property></Properties>')

    def edit(prs):
        package = prs.part.package
        part = Part(PackURI('/docProps/custom.xml'),
                    'application/vnd.openxmlformats-officedocument.custom-properties+xml',
                    package, label.encode('utf-8'))
        package.relate_to(part, RT.CUSTOM_PROPERTIES)
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, _template(tmp_path, edit))
    prs = Presentation(io.BytesIO(_deck()))
    custom = prs.part.package.part_related_by(RT.CUSTOM_PROPERTIES).blob
    assert b'Classification' in custom and b'_MarkAsFinal' in custom
    assert b'pid="3"' in custom
    parts = [str(p.partname) for p in prs.part.package.iter_parts()]
    assert parts.count('/docProps/custom.xml') == 1


def test_a_stock_layout_is_read_by_its_placeholders_types(tmp_path, monkeypatch):
    """Office's own Title and Content layout: no subtitle and no footnotes
    placeholder, a title, an object box, a footer and a slide number. The
    roles are found all the same, and the ones it lacks are left out."""
    stock = Presentation()
    path = str(tmp_path / 'stock.pptx')
    layout = stock.slide_layouts[1]
    marker = stock.slides.add_slide(layout)
    marker.notes_slide.notes_text_frame.text = 'tool: implemented-model'
    stock.save(path)
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, path)
    monkeypatch.setenv(houseTemplate.LAYOUT_SETTING, layout.name)
    found = houseTemplate.report(path)
    assert found['problems'] == []
    assert found['roles']['subtitle'] is None and found['roles']['footnotes'] is None
    assert 'uid' not in found['roles']
    titles = _titles(Presentation(io.BytesIO(_deck())))
    assert titles[0] == 'Implemented Model'
    assert 'Investment Proposal' not in titles           # no cover marker, no cover


def test_a_missing_layout_is_named_with_the_ones_there_are(tmp_path, monkeypatch):
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, PLACEHOLDER)
    monkeypatch.setenv(houseTemplate.LAYOUT_SETTING, 'Content Slide')
    with pytest.raises(TemplateError) as raised:
        _deck()
    assert 'no layout named "Content Slide"' in raised.value.message
    assert '"Content"' in raised.value.message


def test_a_template_that_is_not_there_is_refused(monkeypatch, tmp_path):
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, str(tmp_path / 'missing.pptx'))
    with pytest.raises(TemplateError):
        _deck()


# ---- the strategic slides (step 2 of the plan) -------------------------------------------

def _four(results):
    return list(results) + [dict(r, name=r['name'] + ' (B)', header=r['header']) for r in results]


def _templatePlan(results=None):
    golden, implementation = _golden()
    template = houseTemplate.Template(PLACEHOLDER)
    return pptWriter.planDeck(
        BASIS, MANDATE, results or golden, _model(), assetEstimates.forSlice('USD', 'Hedged'),
        True, implementation['variant'], implementation['feeSchedule'],
        implementation['feeLevel'], UID, box=template.box, restyle=True)


def _labels(doc):
    return [str(doc.rows[r].cells[1].value or '').strip() for r in sorted(doc.rows)
            if 1 in doc.rows[r].cells]


@pytest.mark.parametrize('lineup', ['golden', 'four'])
def test_each_strategic_slide_is_one_slide_of_its_own(lineup):
    golden, _ = _golden()
    results = golden if lineup == 'golden' else _four(golden)
    slides = _templatePlan(results)
    strategic = [s for s in slides if s.get('mark') in ('comparing-allocations', 'risk-return')]
    assert [(s['heading'], s['short']) for s in strategic] == [
        ('Strategic Asset Allocation', 'Comparing Allocations'),
        ('Strategic Asset Allocation', 'Risk and Return Analysis')]
    for slide in strategic:
        plan, = slide['tables']
        assert plan.font >= pptWriter.MIN_PT
        assert sum(plan.heights) <= plan.box[3] + 1e-6


def test_the_allocation_is_the_scenarios_own_groups_and_lines():
    from cyrus_pmg.pmgService.scenario import sheetDoc
    golden, _ = _golden()
    doc = sheetDoc.buildComparingDoc(golden)
    order, assets = sheetDoc._categoryOrder(golden)
    labels = _labels(doc)
    assert labels[:3] == ['', '', '']                       # the three header rows
    expected = []
    for name in order:
        expected.append(name)
        if assets[name] != [name]:
            expected += assets[name]
    assert labels[3:3 + len(expected)] == expected
    assert labels[3 + len(expected):] == ['Total', sheetDoc.MEAN_RETURN_LABEL, 'Volatility',
                                          'Sharpe Ratio', sheetDoc.VAR_YEAR_LABEL]
    assert 'Hybrid Fixed Income' not in labels             # the table's groups, not the model's
    # every group's lines add up to it as printed, in every column
    rows = sorted(doc.rows)
    groups = [r for r in rows if r in doc.sections][:len(order)]
    for start, end in zip(groups, groups[1:] + [rows[3 + len(expected)]]):
        for column in range(2, 2 + len(golden)):
            group = round(doc.rows[start].cells[column].value * 100, 1)
            lines = [doc.rows[r].cells[column].value * 100 for r in range(start + 1, end)]
            if lines:
                assert round(sum(lines), 1) == group
    names = [doc.rows[2].cells[c].value for c in range(2, 2 + len(golden))]
    kinds = [doc.rows[3].cells[c].value for c in range(2, 2 + len(golden))]
    assert names == ['Moderate', 'Conservative'] and kinds == ['Full', 'Full']
    assert doc.rows[1].cells[2].value == 'Proposed Allocation'
    assert doc.rows[1].cells[3].value == 'For Comparison'
    var = next(r for r in rows if 1 in doc.rows[r].cells
               and doc.rows[r].cells[1].value == sheetDoc.VAR_YEAR_LABEL)
    premium = next(p for p in golden[0]['premia'] if p['group'].startswith('Value at Risk')
                   and p['horizon'] == 'Over 1 Year')
    assert doc.rows[var].cells[2].value == pytest.approx(abs(premium['nominalPct']) / 100.0)


def test_a_group_whose_one_line_is_itself_prints_once():
    from cyrus_pmg.pmgService.scenario import sheetDoc
    golden, _ = _golden()
    lonely = json.loads(json.dumps(golden))
    for result in lonely:
        result['categories'].append({'name': 'Asset Allocation Strategies', 'weightPct': 0.0,
                                     'assets': [{'reportingName': 'Asset Allocation Strategies',
                                                 'weightPct': 0.0}]})
    lonely[0]['categories'][-1]['weightPct'] = 2.0
    lonely[0]['categories'][-1]['assets'][0]['weightPct'] = 2.0
    labels = _labels(sheetDoc.buildComparingDoc(lonely))
    assert labels.count('Asset Allocation Strategies') == 1


def test_the_risk_slide_has_the_groups_metrics_every_period_and_three_horizons():
    from cyrus_pmg.pmgService.scenario import sheetDoc
    golden, _ = _golden()
    doc = sheetDoc.buildReturnRiskDoc(golden)
    labels = _labels(doc)
    order, _ = sheetDoc._categoryOrder(golden)
    assert labels[3:3 + len(order)] == order
    periods = [s['period'] for s in golden[0]['stress']]
    shown = [l for l in labels if l.split(' (')[0] in periods]
    assert [l.split(' (')[0] for l in shown] == periods       # every period, in order
    assert 'Oil Embargo (Jan 73 – Sep 74)' in shown and 'Bond Crisis' in shown
    assert labels[-4:] == ['Value at Risk with 99% Confidence', 'Over 1 Month', 'Over 1 Year',
                           'Over 3 Years']
    assert not any('Conditional' in l or 'Probability' in l for l in labels)
    assert doc.signRules and doc.signRules[0][4:] == (sheetDoc.STRESS_NEG, sheetDoc.STRESS_POS)


def test_the_strategic_tables_take_the_templates_colours(placeholder):
    prs = Presentation(io.BytesIO(_deck()))
    allocation = next(shape for shape in prs.slides[1].shapes if shape.has_table)
    xml = etree.tostring(allocation._element).decode()
    assert '<a:srgbClr' not in xml                        # every colour the theme's
    assert 'schemeClr val="accent1"><a:lumMod val="20000"/><a:lumOff val="80000"/>' in xml
    risk = next(shape for shape in prs.slides[2].shapes if shape.has_table)
    inks = {e.get('val') for e in risk._element.iter(qn('a:srgbClr'))}
    from cyrus_pmg.pmgService.scenario import sheetDoc
    assert inks <= {sheetDoc.STRESS_NEG, sheetDoc.STRESS_POS}   # but losses and gains



# ---- the fee slides from what the tool holds (step 3 of the plan) --------------------------

def _rows(table):
    return [[cell.text_frame.text for cell in row.cells] for row in table.rows]


def test_the_implemented_model_has_no_fee_columns_in_a_template(placeholder):
    prs = Presentation(io.BytesIO(_deck()))
    model = next(s for s in prs.slides if _titles([s])[0] == 'Implemented Model') \
        if False else prs.slides[3]
    table = next(shape for shape in model.shapes if shape.has_table).table
    heads = [cell.text_frame.text for cell in table.rows[0].cells]
    assert 'Product Cost' not in heads and 'Mgmt fee' not in heads
    assert not any(h.startswith('Wtd') for h in heads)
    assert 'Allocation (%)' in heads


def test_portfolio_fee_structure_closes_on_100_and_carries_the_fees(placeholder):
    from cyrus_pmg.pmgService.scenario import sheetDoc
    model = _model()
    doc = sheetDoc.buildFeeStructureDoc(model)
    heads = [doc.rows[2].cells[c].value for c in sorted(doc.rows[2].cells)]
    assert heads == [None, None, 'Initial Allocation', 'Target Allocation', 'Product Cost',
                     'Portfolio Management Fee', 'Weighted Average Cost¹', 'Weighted Average Cost¹']
    lines = [r for r in sorted(doc.rows) if r > 2 and 1 in doc.rows[r].cells
             and str(doc.rows[r].cells[1].value or '').startswith('  ')]
    for column in (3, 4):
        total = sum(round(doc.rows[r].cells[column].value * 100, 1)
                    for r in lines if doc.rows[r].cells.get(column) is not None
                    and doc.rows[r].cells[column].value is not None)
        assert round(total, 1) == 100.0
    last = doc.rows[max(doc.rows)]
    assert last.cells[1].value == 'Total Portfolio'
    assert last.cells[7].value == pytest.approx(model['initial']['total']['wtdFeeBp'] / 1e4)
    assert last.cells[8].value == pytest.approx(model['total']['wtdFeeBp'] / 1e4)
    prs = Presentation(io.BytesIO(_deck()))
    titles = _titles(prs)
    assert 'Implemented Portfolio: Initial Allocation' not in titles
    # at its marker, between the template's own Fees Summary and Fee Summary
    assert titles.index('Portfolio Fee Structure') == titles.index('Fees Summary') + 1
    assert titles.index('Fee Summary') == titles.index('Portfolio Fee Structure') + 1


def test_without_fees_it_is_the_weights_alone_or_not_there():
    from cyrus_pmg.pmgService.scenario import sheetDoc
    unpriced = _model(includeFees=False)
    doc = sheetDoc.buildFeeStructureDoc(unpriced, includeFees=False)
    heads = [doc.rows[1].cells[c].value for c in sorted(doc.rows[1].cells)]
    assert heads == [None, None, 'Initial Allocation', 'Target Allocation']
    long = dict(unpriced)
    long.pop('initial')
    assert sheetDoc.buildFeeStructureDoc(long, includeFees=False) is None
    priced = dict(_model())
    priced.pop('initial')
    single = sheetDoc.buildFeeStructureDoc(priced)
    assert [single.rows[1].cells[c].value for c in sorted(single.rows[1].cells)] == [
        None, None, 'Allocation', 'Product Cost', 'Portfolio Management Fee',
        'Weighted Average Cost¹']


def test_the_fee_ladder_is_the_cards_tiers_at_the_proposals_level():
    from cyrus_pmg.pmgService.scenario import fees, sheetDoc
    _, implementation = _golden()
    model = _model()
    level = implementation['feeLevel']
    doc = sheetDoc.buildFeeLadderDoc(model, 'CASP', level)
    tiers = fees.tiers()
    source, point = fees.splitLevel(level)
    labels = [doc.rows[r].cells[1].value for r in sorted(doc.rows)][1:]
    rates = [doc.rows[r].cells[2].value for r in sorted(doc.rows)][1:]
    assert labels == sheetDoc.ladderLabels(tiers, model['currency'])
    assert labels[0].startswith('First $') and labels[-1].startswith('>$')
    assert rates == [pytest.approx(fees._rate('CASP', None, t['id'], source, point) / 100.0)
                     for t in tiers]
    assert doc.rows[1].cells[2].value == 'CASP Fee'
    assert sheetDoc.buildFeeLadderDoc(model, 'RDR', level) is None
    assert sheetDoc.buildFeeLadderDoc(model, None, level) is None
    custom = sheetDoc.buildFeeLadderDoc(model, 'CASP', fees.CUSTOM_LEVEL)
    assert [custom.rows[r].cells[1].value for r in sorted(custom.rows)][1:] == ['Every dollar']


def test_the_tier_table_is_drawn_in_its_area_and_the_figures_in_their_fields(placeholder):
    prs = Presentation(io.BytesIO(_deck()))
    titles = _titles(prs)
    schedule = prs.slides[titles.index('Advisory Fee Schedule')]
    assert not any(shape.name.startswith('tool:') for shape in schedule.shapes)
    table = next(shape for shape in schedule.shapes if shape.has_table)
    template = Presentation(PLACEHOLDER)
    area = next(shape for slide in template.slides for shape in slide.shapes
                if shape.name == 'tool: fee-ladder')
    assert (table.left, table.top) == (area.left, area.top)
    assert table.width <= area.width + 12700 and table.height <= area.height + 12700
    assert _rows(table.table)[1][0].startswith('First $')
    summary = prs.slides[titles.index('Fee Summary')]
    text = ' '.join(f.text for f in houseTemplate._textFrames(summary.shapes))
    values = pptWriter.fieldValues(_model(), MANDATE, UID)
    for name in ('advisoryFee', 'productCost', 'totalCost'):
        assert values[name] in text
    assert 'Based on a $25m portfolio' in text and '{{' not in text
    assert UID not in text


def test_the_fee_slides_follow_the_proposals_pricing(placeholder):
    unpriced = _titles(Presentation(io.BytesIO(_deck(includeFees=False))))
    assert 'Fee Summary' not in unpriced and 'Advisory Fee Schedule' not in unpriced
    assert 'Portfolio Allocation' in unpriced                # a private-markets book
    results, implementation = _golden()
    rdr = writeDeck(BASIS, MANDATE, results, implementation['sleeves'],
                    rules.AUTO_SLEEVE_CATEGORIES, implementation['variant'],
                    implementation['tacticalTilt'], 'RDR', implementation['feeLevel'], True,
                    implementation['volPremium'], assets=assetEstimates.forSlice('USD', 'Hedged'),
                    proposalId=UID, otherCharges=CHARGES)
    titles = _titles(Presentation(io.BytesIO(rdr)))
    assert 'Fee Summary' in titles and 'Advisory Fee Schedule' not in titles


def test_an_area_too_small_for_its_table_is_refused(tmp_path, monkeypatch):
    def edit(prs):
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.name == 'tool: fee-ladder':
                    shape.height = Inches(0.4)
    monkeypatch.setenv(houseTemplate.TEMPLATE_SETTING, _template(tmp_path, edit))
    with pytest.raises(TemplateError) as raised:
        _deck()
    assert '"tool: fee-ladder" is too small' in raised.value.message



# ---- the other charges (step 4 of the plan) ---------------------------------------------

def test_the_fees_summary_quotes_the_charges_the_pwa_entered(placeholder):
    prs = Presentation(io.BytesIO(_deck()))
    titles = _titles(prs)
    summary = prs.slides[titles.index('Fees Summary')]
    text = '\n'.join(f.text for f in houseTemplate._textFrames(summary.shapes))
    assert '▪ Trades over $5M (or currency equivalent): 10 bps' in text
    assert '▪ Trades between $1M - $5M (or currency equivalent): 10 bps' in text
    assert '▪ Trades under $1M (or currency equivalent): 20 bps' in text
    assert 'we charge 35 bps' in text and 'annual custody fee of $5,000.' in text
    assert pptWriter.fieldValues(_model(), MANDATE, UID)['weightedCost'] in text


def test_the_charges_are_checked_on_the_way_in():
    from cyrus_pmg.pmgService.scenario import otherCharges
    from cyrus_pmg.pmgService.scenario.types import ValidationError
    # emptied, the custody fee goes back to its default; GSBZ Custody is $0
    assert otherCharges.clean(dict(CHARGES, custodyFee=None)) == CHARGES
    assert otherCharges.clean(dict(CHARGES, custodyFee=0))['custodyFee'] == 0
    # the trade sizes are the house's, whatever is sent
    assert otherCharges.clean({'fxHighEdge': 1e6, 'fxLowEdge': 5e6}) == {
        'fxHighEdge': 5e6, 'fxLowEdge': 1e6, 'custodyFee': 5000}
    for bad in ({'fxHighBps': -1}, {'custodyFee': 'a lot'}, {'brokerage': 3},
                {'equityTradeBps': 900}, {'custodyFee': -1}):
        with pytest.raises(ValidationError):
            otherCharges.clean(bad)
    # what is left for the PWA: the commissions
    assert otherCharges.missing({}) == [
        f['label'] for f in otherCharges.FIELDS if f['id'] in (
            'fxHighBps', 'fxMidBps', 'fxLowBps', 'equityTradeBps')]
    assert otherCharges.required({}, includeFees=False)['custodyFee'] == 5000
    with pytest.raises(ValidationError) as raised:
        otherCharges.required({'fxHighBps': 10}, includeFees=True)
    assert raised.value.field == 'otherCharges' and 'Buying or selling' in str(raised.value.message)
    described = {f['id']: f for f in otherCharges.describe()}
    assert described['fxHighEdge']['fixed'] == 5e6 and described['fxLowEdge']['fixed'] == 1e6
    assert described['custodyFee']['default'] == 5000
    assert described['custodyFee']['zero'] == 'Remove custody for GSBZ Custody'


def test_the_charges_go_from_the_scenario_to_the_deck_and_the_register(monkeypatch, placeholder):
    """Over HTTP: the PWA's charges are checked on the way in, kept with the
    scenario, quoted on the Fees Summary slide and recorded with the
    proposal; a priced export whose deck quotes a charge still unset is
    refused with it named."""
    import zipfile
    from fastapi.testclient import TestClient
    from cyrus_pmg.pmgService.isgPMGService import app
    from cyrus_pmg.pmgService.scenario import proposalRegister, sleeves
    from cyrus_pmg.pmgService.scenario.sleeves import listSleeves
    from cyrus_pmg.pmgService.scenario.types import PortfolioKey
    monkeypatch.setenv('PMG_ALLOWED_KERBEROS', 'bob')
    client = TestClient(app)
    pwa = {'X-Kerberos': 'bob'}
    sid = client.post('/api/v1/scenario', headers=pwa, json={
        'mandate': {'topAccountSize': 1e9, 'mandateSize': 1e9, 'primaryPwa': 'A. Castellanos — Madrid'},
        'basis': {'currency': 'USD', 'hedging': 'Hedged'}}).json()['id']
    assert client.put('/api/v1/scenario/' + sid, headers=pwa,
                      json={'variant': sleeves.VARIANTS[0]}).status_code == 200
    key = {'currency': 'USD', 'riskLevel': 'Moderate', 'allocationType': 'Full',
           'excludeRealAssets': False}
    resolved = client.post('/api/v1/scenario/' + sid + '/portfolio', headers=pwa,
                           json={'key': key, 'role': 'base'})
    chosen = {}
    for category in resolved.json()['portfolio']['categories']:
        if category['name'] in rules.AUTO_SLEEVE_CATEGORIES:
            continue
        under = rules.sleeveCategory(category['name'])
        chosen.setdefault(under, listSleeves(under, sleeves.VARIANTS[0],
                                             PortfolioKey.fromDict(key))[0]['name'])
    assert client.put('/api/v1/scenario/' + sid, headers=pwa, json={
        'sleeves': chosen, 'includeFees': True, 'feeSchedule': 'CASP'}).status_code == 200

    bad = client.put('/api/v1/scenario/' + sid, headers=pwa,
                     json={'otherCharges': {'custodyFee': -5}})
    assert bad.status_code == 422 and bad.json()['field'] == 'otherCharges'
    # a new scenario stands at the house's trade sizes and the default custody
    assert client.get('/api/v1/scenario/' + sid, headers=pwa).json()['otherCharges'] == {
        'fxHighEdge': 5e6, 'fxLowEdge': 1e6, 'custodyFee': 5000}
    refused = client.post('/api/v1/scenario/' + sid + '/export', headers=pwa)
    assert refused.status_code == 422, refused.text
    assert refused.json()['field'] == 'otherCharges'
    assert 'Enter the other charges before downloading' in refused.json()['error']

    stored = client.put('/api/v1/scenario/' + sid, headers=pwa, json={'otherCharges': CHARGES})
    assert stored.status_code == 200
    assert client.get('/api/v1/scenario/' + sid, headers=pwa).json()['otherCharges'] == {
        k: (round(v) if k in ('fxHighEdge', 'fxLowEdge', 'custodyFee') else v)
        for k, v in CHARGES.items()}
    delivered = client.post('/api/v1/scenario/' + sid + '/export', headers=pwa)
    assert delivered.status_code == 200, delivered.text
    uid = delivered.headers['x-proposal-id']
    archive = zipfile.ZipFile(io.BytesIO(delivered.content))
    deck = next(archive.read(n) for n in archive.namelist() if n.endswith('.pptx'))
    prs = Presentation(io.BytesIO(deck))
    titles = _titles(prs)
    text = ' '.join(f.text for f in houseTemplate._textFrames(prs.slides[titles.index('Fees Summary')].shapes))
    assert 'Trades over $5M' in text and 'annual custody fee of $5,000.' in text
    conn = proposalRegister._connect()
    try:
        held = conn.execute('SELECT otherCharges FROM proposals WHERE proposalId = ?', (uid,)).fetchone()[0]
    finally:
        conn.close()
    assert json.loads(held)['custodyFee'] == 5000


def test_the_check_takes_a_template_without_a_cover_marker_to_have_its_own(tmp_path):
    _template(tmp_path, lambda prs: _dropMarker(prs, 'cover'))
    assert houseTemplate.report(str(tmp_path / 'template.pptx'))['unmarked'] == []
