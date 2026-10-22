"""The product catalogue kept in the repository and maintained from the
console (D165).

Each test has a repository file of its own, loaded from the packaged files
on its first open exactly as an upgraded host's is."""
import base64
import csv
import json
import sqlite3

import pytest

from cyrus_pmg.pmgService import dashboardRouter
from cyrus_pmg.pmgService.scenario import fees, products, sleeveRepo, sleeveTools
from cyrus_pmg.pmgService.scenario.types import ValidationError


@pytest.fixture()
def store(tmp_path, monkeypatch):
    monkeypatch.setenv('SCENARIO_SLEEVES_DB', str(tmp_path / 'sleeves.db'))
    monkeypatch.setenv('SCENARIO_REGISTER_DB', str(tmp_path / 'proposals.db'))
    products.reload()
    yield tmp_path
    products.reload()


def _caller(kerberos='alice'):
    from cyrus_pmg.pmgService.core import accessControl
    return accessControl.UserData(kerberos, accessControl._rolesFor(kerberos))


def _placed():
    return {row['productId'] for e in sleeveRepo.listAll() for row in e['products']}


def _unplaced():
    held = _placed()
    return next(p['productId'] for p in products.all() if p['productId'] not in held)


def _new(**over):
    fields = {'productId': 'probe-etf', 'name': 'Probe Global Equity ETF', 'ticker': 'PRB',
              'assetClass': 'Global Equity', 'style': 'Passive', 'vehicle': 'ETF',
              'shareClass': 'Accumulating', 'source': 'External', 'liquidity': 'Daily',
              'exposureCurrency': 'USD', 'productCost': '0.12', 'feeGroup': 'Passive',
              'distributionYield': '', 'minimumInvestment': ''}
    fields.update(over)
    return fields


def _upload(path):
    with open(path, 'rb') as fh:
        return base64.b64encode(fh.read()).decode('ascii')


def _fileWith(tmp_path, rows, name='products.csv'):
    path = tmp_path / name
    with open(path, 'w', newline='', encoding='utf-8') as fh:
        csv.writer(fh).writerows(rows)
    return str(path)


# ------------------------------------------------------------ the first load ---

def test_the_first_open_loads_the_file_once_and_serves_what_the_file_served(store, tmp_path,
                                                                           monkeypatch):
    """The catalogue moves in on the first open after the upgrade (schema 7):
    every product the file carried, in its order and its served shape, one
    *seeded* revision each - and once moved in, the file is never read again."""
    byId, order = products.readFile(products.sourcePath(), strict=False)
    served = products.all()
    assert [p['productId'] for p in served] == order
    for p in served:
        record = dict(byId[p['productId']])
        record['ticker'] = record['ticker'] or products.NO_TICKER
        record.pop('category')                   # the console's, not served (D178)
        assert p == record
    info = products.describeSource()
    assert info['products'] == 73 and info['withdrawn'] == 0 and info['revision'] == 1
    assert info['lastEditedBy'] == 'seed' and info['seededFrom'] == products.sourcePath()
    history = products.history(order[0])
    assert [h['action'] for h in history] == ['seeded'] and history[0]['source'] == 'products.csv'
    # the file is not read again: pointing the service at another changes nothing
    other = _fileWith(tmp_path, [products.FILE_COLUMNS])
    monkeypatch.setenv('SCENARIO_PRODUCTS_SOURCE', other)
    products.reload()
    assert len(products.all()) == 73


def test_a_file_that_cannot_be_read_leaves_the_catalogue_unloaded_and_says_why(tmp_path,
                                                                             monkeypatch):
    """Nothing half-loaded: the reads say what was wrong, the sleeves and the
    rest of the repository still open, and the next open tries again."""
    monkeypatch.setenv('SCENARIO_SLEEVES_DB', str(tmp_path / 'sleeves.db'))
    bad = _fileWith(tmp_path, [products.FILE_COLUMNS,
                               ['p-1', 'Probe', '', 'Equity', 'Active', 'ETF', '', 'Internal',
                                'Daily', 'USD', 'free', 'Passive', '', '']], name='bad.csv')
    monkeypatch.setenv('SCENARIO_PRODUCTS_SOURCE', bad)
    products.reload()
    try:
        with pytest.raises(products.BadCatalogue) as caught:
            products.all()
        assert 'bad.csv' in str(caught.value) and 'row 2' in str(caught.value)
        assert sleeveRepo.describe()['schemaVersion'] == sleeveRepo.SCHEMA_VERSION
        monkeypatch.delenv('SCENARIO_PRODUCTS_SOURCE')
        assert len(products.all()) == 73, 'fixed, the next open loads it'
    finally:
        products.reload()


# ------------------------------------------------------------------ writes ---

def test_a_product_is_added_with_its_stamps_and_a_first_version(store):
    before = products.describeSource()['revision']
    made = products.create(_new(), user='alice', note='New ETF for the ESG sleeves')
    assert made['productId'] == 'probe-etf' and made['revision'] == 1
    assert made['createdBy'] == made['lastEditedBy'] == 'alice'
    assert made['productCost'] == 0.12 and made['distributionYield'] is None
    assert products.get('probe-etf')['name'] == 'Probe Global Equity ETF'
    assert products.describeSource()['revision'] == before + 1
    assert products.describeSource()['lastEditedBy'] == 'alice'
    trail = products.history('probe-etf')
    assert trail[0]['action'] == 'created' and trail[0]['note'] == 'New ETF for the ESG sleeves'
    # placed in a sleeve at once: the library's own check sees it
    variant, category = sleeveRepo.VARIANTS[0], sleeveRepo.categories()[0]
    sleeve = sleeveRepo.createSleeve(variant, category, 'Probe Sleeve',
                                     [{'productId': 'probe-etf', 'weight': 1}], user='alice')
    assert sleeve['products'][0]['product']['name'] == 'Probe Global Equity ETF'


@pytest.mark.parametrize('over, field, words', [
    ({'name': '  '}, 'name', 'Give the product a name'),
    ({'productCost': 'free'}, 'productCost', 'not a number'),
    ({'productCost': '-0.1'}, 'productCost', 'cannot be negative'),
    ({'productCost': '45'}, 'productCost', 'looks wrong'),
    ({'feeGroup': 'Made Up'}, 'feeGroup', 'not on the fee card'),
    ({'feeGroup': ''}, 'feeGroup', 'Choose a fee group'),
    ({'shareClass': 'Distribution'}, 'shareClass', 'not one of'),
    ({'style': ''}, 'style', 'Style is needed'),
    ({'minimumInvestment': '5m'}, 'minimumInvestment', 'not a number'),
    ({'productId': ''}, 'productId', 'Give the product an ID'),
])
def test_a_product_that_is_not_one_is_refused_naming_the_field(store, over, field, words):
    with pytest.raises(ValidationError) as caught:
        products.create(_new(**over), user='alice')
    assert caught.value.field == field and words in caught.value.message
    assert products.get('probe-etf') is None


def test_an_id_is_never_reused_and_typed_values_are_read_as_people_type_them(store):
    taken = products.all()[0]
    with pytest.raises(ValidationError) as caught:
        products.create(_new(productId=taken['productId']), user='alice')
    assert caught.value.field == 'productId' and taken['name'] in caught.value.message
    made = products.create(_new(productCost='0.25%', minimumInvestment='5,000,000',
                                shareClass='Dis', name='Probe   Fund'), user='alice')
    assert made['productCost'] == 0.25 and made['minimumInvestment'] == 5000000.0
    assert made['shareClass'] == 'Distributing' and made['name'] == 'Probe Fund'
    products.withdraw('probe-etf', user='alice')
    with pytest.raises(ValidationError) as caught:
        products.create(_new(), user='alice')
    assert 'never reused' in caught.value.message


def test_a_change_is_a_new_version_field_by_field_and_a_stale_one_is_refused(store):
    pid = _unplaced()
    first = products.entry(pid)
    fields = dict(first, productCost=first['productCost'] + 0.05, liquidity='Weekly')
    saved = products.update(pid, fields, user='bob', note='Repriced', base=first['revision'])
    assert saved['revision'] == first['revision'] + 1 and saved['lastEditedBy'] == 'bob'
    top = products.history(pid)[0]
    assert top['action'] == 'updated' and top['note'] == 'Repriced' and top['current']
    assert {c['field'] for c in top['changes']} == {'productCost', 'liquidity'}
    assert any(line.startswith('Liquidity ') and line.endswith('→ Weekly') for line in top['changeWords'])
    # the same save again, made against the version it replaced
    with pytest.raises(products.StaleError) as caught:
        products.update(pid, dict(fields, liquidity='Monthly'), user='carol', base=first['revision'])
    assert 'bob' in caught.value.message and 'Reload' in caught.value.message
    with pytest.raises(ValidationError) as caught:
        products.update(pid, fields, user='bob', base=saved['revision'])
    assert 'Nothing has changed' in caught.value.message
    with pytest.raises(ValidationError) as caught:
        products.update(pid, dict(fields, productId='something-else'), user='bob')
    assert caught.value.field == 'productId' and 'cannot be changed' in caught.value.message


def test_a_product_in_a_sleeve_cannot_be_withdrawn_and_one_that_is_not_comes_back(store):
    placed = sorted(_placed())[0]
    with pytest.raises(ValidationError) as caught:
        products.withdraw(placed, user='alice')
    assert 'sleeve' in caught.value.message and 'cannot be withdrawn' in caught.value.message
    assert products.has(placed)

    pid = _unplaced()
    gone = products.withdraw(pid, user='alice', note='Closed to new money')
    assert gone['withdrawn'] and gone['withdrawnBy'] == 'alice'
    assert not products.has(pid) and products.get(pid) is None
    assert pid in [p['productId'] for p in products.entries(withdrawn=True)]
    assert len(products.all()) == 72
    # a withdrawn product cannot be placed
    with pytest.raises(ValidationError):
        sleeveRepo.createSleeve(sleeveRepo.VARIANTS[0], sleeveRepo.categories()[0], 'Gone',
                                [{'productId': pid, 'weight': 1}])
    with pytest.raises(ValidationError) as caught:
        products.update(pid, products.entry(pid), user='alice')
    assert 'Restore it first' in caught.value.message
    back = products.restore(pid, user='bob')
    assert not back['withdrawn'] and products.has(pid)
    assert [h['action'] for h in products.history(pid)][:2] == ['restored', 'withdrawn']


def test_an_earlier_version_is_put_back_as_a_new_one(store):
    pid = _unplaced()
    original = products.entry(pid)
    products.update(pid, dict(original, name=original['name'] + ' II'), user='alice')
    reverted = products.revert(pid, 1, user='bob')
    assert reverted['name'] == original['name'] and reverted['revision'] == 3
    top = products.history(pid)[0]
    assert top['action'] == 'reverted' and top['restoredFrom'] == 1
    assert top['note'] == 'Restored version 1'
    with pytest.raises(ValidationError) as caught:
        products.revert(pid, 1, user='bob')
    assert 'in force already' in caught.value.message


def test_a_write_from_another_process_is_read_at_once(store):
    """The cache is the file's: a write that did not pass through this
    process - another worker's - moves the file, and the next read sees it."""
    pid = _unplaced()
    assert products.get(pid)['name'] != 'Renamed elsewhere'
    conn = sqlite3.connect(sleeveRepo.dbPath())
    conn.execute("UPDATE products SET name = 'Renamed elsewhere' WHERE productId = ?", (pid,))
    conn.commit()
    conn.close()
    assert products.get(pid)['name'] == 'Renamed elsewhere'


# ------------------------------------------------------------------ import ---

def test_a_file_is_previewed_then_imported_and_what_it_leaves_alone_is_untouched(store, tmp_path):
    rows = products.exportRows()
    head, body = rows[0], rows[1:]
    at = {c: i for i, c in enumerate(head)}
    held = _placed()
    dropped = next(r for r in body if r[0] not in held)[0]
    keptHeld = next(r for r in body if r[0] in held)[0]
    changed = next(r for r in body if r[0] not in (dropped, keptHeld))
    changed[at['ProductCost']] = '0.99'
    newRow = ['probe-etf', 'Probe Global Equity ETF', 'PRB', 'Global Equity', 'Public Equity',
              'Passive', 'ETF', 'Accumulating', 'External', 'Daily', 'USD', '', '0.12', 'Passive', '', '']
    out = [head] + [r for r in body if r[0] not in (dropped, keptHeld)] + [newRow]
    path = _fileWith(tmp_path, out, name='delivery.csv')
    untouched = products.entry(body[-1][0])

    preview = products.importPreview('delivery.csv', _upload(path))
    assert [a['productId'] for a in preview['added']] == ['probe-etf']
    assert [c['productId'] for c in preview['changed']] == [changed[0]]
    assert preview['changed'][0]['changeWords'][0].startswith('Product cost ')
    missing = {m['productId']: m for m in preview['missing']}
    assert set(missing) == {dropped, keptHeld}
    assert missing[dropped]['canWithdraw'] and not missing[keptHeld]['canWithdraw']
    assert missing[keptHeld]['sleeves']
    assert preview['unchanged'] == len(body) - 3

    # a product a sleeve holds cannot be withdrawn by an import either
    with pytest.raises(ValidationError):
        products.importFile('delivery.csv', _upload(path), user='alice', withdraw=[keptHeld],
                            base=preview['revision'])
    done = products.importFile('delivery.csv', _upload(path), user='alice', withdraw=[dropped],
                               note='October delivery', base=preview['revision'])
    assert (done['added'], done['changed'], done['withdrawn']) == (1, 1, 1)
    assert products.get(changed[0])['productCost'] == 0.99
    assert products.has('probe-etf') and not products.has(dropped) and products.has(keptHeld)
    top = products.history(changed[0])[0]
    assert top['action'] == 'imported' and top['source'] == 'delivery.csv'
    assert top['note'] == 'October delivery'
    after = products.entry(untouched['productId'])
    assert (after['revision'], after['lastEditedAt']) == (untouched['revision'], untouched['lastEditedAt'])
    # and against the preview it was made from, a second apply is stale
    with pytest.raises(products.StaleError):
        products.importFile('delivery.csv', _upload(path), user='alice', base=preview['revision'])


def test_a_file_with_one_bad_row_is_refused_whole(store, tmp_path):
    rows = products.exportRows()
    rows[5][rows[0].index('FeeGroup')] = 'Made Up'
    rows[2][rows[0].index('ProductCost')] = '0.77'
    path = _fileWith(tmp_path, rows, name='bad.csv')
    before = products.describeSource()['revision']
    with pytest.raises(products.BadCatalogue) as caught:
        products.importPreview('bad.csv', _upload(path))
    assert 'row 6' in str(caught.value) and 'fee group' in str(caught.value)
    with pytest.raises(products.BadCatalogue):
        products.importFile('bad.csv', _upload(path), user='alice')
    assert products.describeSource()['revision'] == before
    assert products.get(rows[2][0])['productCost'] != 0.77


def test_an_export_reads_back_as_the_catalogue_and_a_workbook_imports_too(store, tmp_path):
    from openpyxl import Workbook
    rows = products.exportRows()
    path = _fileWith(tmp_path, rows, name='same.csv')
    preview = products.importPreview('same.csv', _upload(path))
    assert not (preview['added'] or preview['changed'] or preview['missing'])
    with pytest.raises(ValidationError) as caught:
        products.importFile('same.csv', _upload(path), user='alice')
    assert 'nothing to import' in caught.value.message
    book = Workbook()
    book.active.title = 'Read me'
    sheet = book.create_sheet('products')
    sheet.append(rows[0])
    for r in rows[1:]:
        sheet.append([float(v) if v and c in ('ProductCost', 'DistributionYield', 'MinimumInvestment')
                      else v for v, c in zip(r, rows[0])])
    sheet.cell(row=3, column=rows[0].index('Name') + 1).value = 'Renamed In A Workbook'
    xlsx = tmp_path / 'products.xlsx'
    book.save(str(xlsx))
    preview = products.importPreview('products.xlsx', _upload(str(xlsx)))
    assert [c['productId'] for c in preview['changed']] == [rows[2][0]]


# ------------------------------------------------------- the category (D178) ---

def test_a_product_has_a_category_of_its_own_from_the_file_under_any_of_its_headings(store,
                                                                                  tmp_path):
    """The column the catalogue bands by: read under Category, Catagory or
    Group; previewed and imported field by field like any other; kept by the
    console's products and the export, not served beside a sleeve product."""
    rows = products.exportRows()
    head = rows[0]
    at = head.index('Category')
    assert head[at - 1] == 'AssetClass', 'the export writes it after the asset class'
    pid = rows[1][0]
    assert products.entry(pid)['category'] == rows[1][at] == 'Investment Grade Fixed Income'
    assert 'category' not in products.get(pid), 'a sleeve product takes its sleeve\'s'
    for heading in ('Catagory', 'Group'):
        out = [list(r) for r in rows]
        out[0][at] = heading
        out[1][at] = 'Hedge Funds'
        path = _fileWith(tmp_path, out, name='delivery.csv')
        preview = products.importPreview('delivery.csv', _upload(path))
        assert [c['productId'] for c in preview['changed']] == [pid], heading
        assert preview['changed'][0]['changeWords'] == [
            'Category Investment Grade Fixed Income → Hedge Funds']
    products.importFile('delivery.csv', _upload(path), user='alice', base=preview['revision'])
    assert products.entry(pid)['category'] == 'Hedge Funds'
    assert products.history(pid)[0]['changeWords'] == [
        'Category Investment Grade Fixed Income → Hedge Funds']
    # a file without the column at all still imports, and blanks every category
    plain = [[v for i, v in enumerate(r) if i != at] for r in products.exportRows()]
    path = _fileWith(tmp_path, plain, name='plain.csv')
    preview = products.importPreview('plain.csv', _upload(path))
    assert len(preview['changed']) == len(plain) - 1
    assert preview['changed'][0]['changeWords'][0].startswith('Category ')


@pytest.mark.parametrize('typed, words', [
    ('public equity', "is spelled Public Equity"),
    ('Equities', "is not one the sleeves are built for: Investment Grade Fixed Income"),
])
def test_a_category_must_be_one_the_library_knows(store, tmp_path, typed, words):
    with pytest.raises(ValidationError) as caught:
        products.create(_new(category=typed), user='alice')
    assert caught.value.field == 'category' and words in caught.value.message
    rows = products.exportRows()
    rows[3][rows[0].index('Category')] = typed
    path = _fileWith(tmp_path, rows, name='typo.csv')
    with pytest.raises(products.BadCatalogue) as caught:
        products.importPreview('typo.csv', _upload(path))
    assert 'row 4' in str(caught.value) and words in str(caught.value)


def test_a_category_is_given_and_changed_in_the_form_and_blank_is_none(store):
    made = products.create(_new(category='Public Equity'), user='alice')
    assert made['category'] == 'Public Equity'
    moved = products.update('probe-etf', _new(category=''), user='alice', base=made['revision'])
    assert moved['category'] == ''
    assert products.history('probe-etf')[0]['changeWords'] == ['Category Public Equity → —']


def test_a_catalogue_from_before_the_category_gains_the_column_blank(store):
    """8 -> 9: the column is added to the products already there, each blank,
    and nothing else about them moves."""
    before = products.entries()
    conn = sleeveRepo.connection()
    conn.execute('ALTER TABLE products DROP COLUMN category')
    conn.execute("UPDATE meta SET value = '8' WHERE key = 'schemaVersion'")
    conn.commit()
    conn.close()
    products.reload()
    assert sleeveRepo.describe()['schemaVersion'] == 14
    after = products.entries()
    assert [p['category'] for p in after] == [''] * len(before)
    strip = lambda ps: [{k: v for k, v in p.items() if k != 'category'} for p in ps]
    assert strip(after) == strip(before)


# ---------------------------------------------------------- the console's API ---

def test_the_endpoints_write_as_the_signed_in_admin_and_refuse_as_the_page_expects(store, monkeypatch):
    monkeypatch.setenv('PMG_ALLOWED_KERBEROS', 'alice')
    monkeypatch.setenv('PMG_ADMIN_KERBEROS', 'alice')
    alice = _caller('alice')
    body = dashboardRouter.getRepository(caller=alice)
    assert len(body['products']) == 73 and body['withdrawnProducts'] == []
    assert body['products'][0]['revision'] == 1 and body['products'][0]['lastEditedBy'] == 'seed'
    assert 'ETF' in body['productVocabulary']['vehicle'] and body['feeGroups'] == fees.feeGroups()
    made = dashboardRouter.createRepositoryProduct({'product': _new(), 'note': 'n'}, caller=alice)
    assert made['product']['createdBy'] == 'alice' and made['catalogue']['products'] == 74
    refused = dashboardRouter.createRepositoryProduct({'product': _new(productId='x-1', name='')},
                                                      caller=alice)
    assert refused.status_code == 422 and json.loads(refused.body)['field'] == 'name'
    stale = dashboardRouter.updateRepositoryProduct(
        {'productId': 'probe-etf', 'product': _new(name='Changed'), 'base': 7}, caller=alice)
    assert stale.status_code == 409
    saved = dashboardRouter.updateRepositoryProduct(
        {'productId': 'probe-etf', 'product': _new(name='Changed'), 'base': 1}, caller=alice)
    assert saved['product']['name'] == 'Changed'
    trail = dashboardRouter.getRepositoryProductHistory(productId='probe-etf', caller=alice)
    assert [h['action'] for h in trail['history']] == ['updated', 'created']
    gone = dashboardRouter.withdrawRepositoryProduct({'productId': 'probe-etf', 'base': 2}, caller=alice)
    assert gone['product']['withdrawn']
    back = dashboardRouter.restoreRepositoryProduct({'productId': 'probe-etf'}, caller=alice)
    assert not back['product']['withdrawn']
    csvBody = dashboardRouter.exportProducts(caller=alice).body.decode('utf-8')
    assert csvBody.splitlines()[0] == ','.join(products.EXPORT_COLUMNS)
    assert len(csvBody.strip().splitlines()) == 75
    bad = dashboardRouter.previewProductImport({'name': 'x.txt', 'content': 'eA=='}, caller=alice)
    assert bad.status_code == 422 and json.loads(bad.body)['field'] == 'file'


# ------------------------------------------------------------- the command line ---

def test_the_command_line_exports_imports_and_reads_a_history(store, tmp_path, capsys):
    out = tmp_path / 'out.csv'
    assert sleeveTools.main(['--products-export', str(out)]) == 0
    rows = list(csv.reader(open(str(out), encoding='utf-8')))
    assert rows[0] == products.EXPORT_COLUMNS and len(rows) == 74
    pid = rows[3][0]
    rows[3][rows[0].index('Liquidity')] = 'Weekly'
    with open(str(out), 'w', newline='', encoding='utf-8') as fh:
        csv.writer(fh).writerows(rows)
    capsys.readouterr()
    assert sleeveTools.main(['--products-import', str(out), '--note', 'from the desk']) == 0
    said = capsys.readouterr().out
    assert '0 new, 1 changed' in said and 'Liquidity' in said
    assert products.get(pid)['liquidity'] == 'Weekly'
    assert sleeveTools.main(['--product-history', pid]) == 0
    said = capsys.readouterr().out
    assert 'imported' in said and 'seeded' in said and '"from the desk"' in said
    assert sleeveTools.main(['--census']) == 0
    assert 'catalogue  in the store · version 2 · 73 products' in capsys.readouterr().out
