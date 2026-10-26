"""D211: the MAP benchmarks - the benchmark of a multi-asset portfolio and
its mnemonic, by the portfolio's risk level, base currency and hedging
policy.

One table, a row a benchmark: kept in the console's Account Opening tab,
edited on screen, pasted from Excel or loaded whole from a file. A portfolio
on offer with no row is a gap. A strategic proposal's workbook and deck name
its portfolio's benchmark, and the register records it with its mnemonic; a
custom portfolio's is chosen from the ones the table holds.

D216: a portfolio under Equity Not Hedged takes the Unhedged row; the table
holds no rows for it.
"""

import base64
import csv
import io
import json
import os
import shutil
import subprocess
import zipfile

import pytest

from cyrus_pmg.pmgService import dashboardRouter
from cyrus_pmg.pmgService.scenario import portfolioBenchmarks as pb, proposalRegister, repositoryFeed, sleeveRepo
from cyrus_pmg.pmgService.scenario.types import ValidationError
from test_proposal_register import _caller

HEAD = list(pb.COLUMNS)


@pytest.fixture()
def store(tmp_path, monkeypatch):
    monkeypatch.setenv('SCENARIO_SLEEVES_DB', str(tmp_path / 'sleeves.db'))
    monkeypatch.setenv('SCENARIO_REGISTER_DB', str(tmp_path / 'proposals.db'))
    monkeypatch.setenv('SCENARIO_STORE_DIR', str(tmp_path / 'store'))
    sleeveRepo.describe()
    yield tmp_path


def _upload(tmp_path, rows, name='benchmarks.csv'):
    path = tmp_path / name
    with open(path, 'w', newline='', encoding='utf-8') as fh:
        csv.writer(fh).writerows(rows)
    return base64.b64encode(path.read_bytes()).decode('ascii')


def _row(benchmark, currency, risk, hedging, mnemonic=''):
    return {'benchmark': benchmark, 'currency': currency, 'risk': risk, 'hedging': hedging, 'mnemonic': mnemonic}


def _full():
    """A row for every portfolio on offer."""
    return {'rows': [_row('MAP {} {} {}'.format(*c.split('|')), c.split('|')[1], c.split('|')[0], c.split('|')[2],
                          'M-{}'.format(c.split('|')[1])) for c in pb.cells()]}


def test_a_portfolio_is_a_risk_level_a_currency_and_a_hedging_policy(store):
    cells = pb.cells()
    assert len(cells) == len(set(cells)) == 3 * len(pb._offered()), 'Equity Not Hedged takes Unhedged'
    assert cells[0] == 'LowVol|USD|Hedged', 'by currency, then risk, then hedging'
    assert pb.offeredRisks()[-1] == 'All Equity' and pb.offeredCurrencies() == ['USD', 'GBP', 'CHF', 'EUR']
    # the allocation plays no part: every allocation at a risk level reads one portfolio
    assert pb.keyOf('USD|Moderate|Full|1', 'Hedged') == pb.keyOf('USD|Moderate|Core|0', 'Hedged') == 'Moderate|USD|Hedged'
    assert pb.keyOf('GBP|All Equity|NA|NA', 'Unhedged') == 'All Equity|GBP|Unhedged'
    assert pb.current()['revision'] == 0
    assert pb.gaps() == cells and pb.choices() == []
    assert pb.forPortfolio('USD|Moderate|Full|0', 'Hedged') == ''


def test_rows_edited_on_screen_name_each_portfolios_benchmark(store):
    pb.save({'rows': [_row('  60/40  USD ', 'usd', 'Moderate', 'hedged', 'BM60US'),
                      _row('ACWI GBP', 'GBP', 'All Equity', 'Unhedged'),
                      _row('', '', '', '', '')]}, 'First', user='alice', base=0)
    now = pb.current()
    assert now['revision'] == 1 and len(now['rows']) == 2, 'a wholly blank row is left out'
    assert pb.entryFor('Moderate|USD|Hedged') == {'benchmark': '60/40 USD', 'mnemonic': 'BM60US'}
    assert pb.forPortfolio('USD|Moderate|ex-Alts|0', 'Hedged') == '60/40 USD'
    assert pb.entryForPortfolio('GBP|All Equity|NA|NA', 'Unhedged') == {'benchmark': 'ACWI GBP', 'mnemonic': ''}
    assert pb.forPortfolio('USD|Moderate|Full|0', 'Unhedged') == ''
    assert pb.choices() == ['60/40 USD', 'ACWI GBP'] and pb.mnemonicOf('60/40 USD') == 'BM60US'
    assert len(pb.gaps()) == len(pb.cells()) - 2
    # risk by its label too
    pb.save({'rows': now['rows'] + [_row('Cons-Mod EUR', 'EUR', 'Conservative-Moderate', 'ISG Hedged', 'CM')]},
            user='alice', base=1)
    assert pb.benchmarkFor('ConsMod|EUR|ISG Hedged') == 'Cons-Mod EUR'
    assert pb.history()[0]['changes'] == ['Conservative-Moderate · EUR · ISG Hedged: added Cons-Mod EUR (CM)']
    feed = [e for e in repositoryFeed.activity(limit=50)['entries'] if e['area'] == 'benchmarks']
    assert {e['name'] for e in feed} == {'Portfolio benchmarks'}


def test_equity_not_hedged_takes_the_unhedged_row(store, tmp_path):
    assert pb.hedgings() == ['Hedged', 'ISG Hedged', 'Unhedged']
    assert not any(c.endswith('|Equity Not Hedged') for c in pb.cells()), 'no portfolio of its own, so no gap'
    assert pb.keyOf('USD|Moderate|Full|0', 'Equity Not Hedged') == 'Moderate|USD|Unhedged'
    pb.save({'rows': [_row('60/40 USD', 'USD', 'Moderate', 'Unhedged', 'U60')]}, user='alice', base=0)
    assert pb.entryForPortfolio('USD|Moderate|Core|1', 'Equity Not Hedged') == {'benchmark': '60/40 USD',
                                                                                'mnemonic': 'U60'}
    assert pb.forPortfolio('USD|Moderate|Full|0', 'Hedged') == ''
    from cyrus_pmg.pmgService.scenario import accountRequests
    assert accountRequests.mandateBenchmark({'baseKey': 'USD|Moderate|Full|0', 'hedging': 'Equity Not Hedged'}) == (
        '60/40 USD', 'U60', 'table'), 'an account opening request for one delivered before reads it too'
    # a row for it is refused on screen, and left out of a file
    with pytest.raises(ValidationError, match='Row 1: Equity Not Hedged takes the Unhedged benchmark'):
        pb.save({'rows': [_row('X', 'USD', 'Moderate', 'equity not hedged')]}, user='alice', base=1)
    rows = [HEAD, ['A', 'USD', 'Moderate', 'Unhedged', 'U60'], ['X', 'GBP', 'Moderate', 'Equity Not Hedged', ''],
            ['B', 'GBP', 'Aggressive', 'Hedged', '']]
    preview = pb.preview('b.csv', _upload(tmp_path, rows))
    assert preview['problems'] == [] and preview['leftOut'] == 1 and preview['rows'] == 2
    pb.importFile('b.csv', _upload(tmp_path, rows), user='alice', base=1)
    assert [r['benchmark'] for r in pb.current()['rows']] == ['A', 'B']
    only = pb.preview('b.csv', _upload(tmp_path, [HEAD, ['X', 'GBP', 'Moderate', 'Equity Not Hedged', '']]))
    assert only['problems'] == ['b.csv holds no rows but ones for Equity Not Hedged.']
    # a table saved with rows for it reads without them
    conn = sleeveRepo.connection()
    conn.execute("UPDATE policies SET body = ? WHERE kind = 'portfolioBenchmarks'",
                 (json.dumps({'rows': [_row('A', 'USD', 'Moderate', 'Unhedged'),
                                       _row('Z', 'USD', 'Moderate', 'Equity Not Hedged')]}),))
    conn.commit()
    conn.close()
    assert [r['benchmark'] for r in pb.current()['rows']] == ['A']
    assert pb.forPortfolio('USD|Moderate|Full|0', 'Equity Not Hedged') == 'A'
    assert dashboardRouter.getPortfolioBenchmarks(caller=_caller('alice'))['readsAs'] == {
        'Equity Not Hedged': 'Unhedged'}


def test_the_console_leaves_a_pasted_row_for_equity_not_hedged_out():
    """The page's paste (pbmPasteInto) and its check of a row (pbmChecked)
    read Equity Not Hedged as the server does."""
    node = shutil.which('node')
    if not node:
        pytest.skip('node not available')
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', '..', 'generator', 'js', 'repository.js')
    with open(path, encoding='utf-8') as fh:
        source = fh.read()

    def fn(name):
        start = source.index('function ' + name + '(')
        return source[start:source.index('\n}\n', start)] + '\n}\n'
    data = {'currencies': pb.offeredCurrencies(), 'hedgings': pb.hedgings(), 'readsAs': pb.READS_AS,
            'risks': [{'risk': r, 'label': pb.riskLabel(r)} for r in pb.offeredRisks()], 'rows': []}
    script = ('var PBM_FIELDS = %s; var pbm = { data: %s, draft: null, confirming: false };\n'
              'function render() {}\nfunction aocTidy(v) { return String(v == null ? "" : v).trim(); }\n' % (
                  json.dumps(list(pb.FIELDS)), json.dumps(data))
              + ''.join(fn(n) for n in ('pbmSquash', 'pbmMatch', 'pbmTakesAnother', 'pbmRiskLabel', 'pbmRows',
                                        'pbmEnsureDraft', 'pbmBlank', 'pbmCellKey', 'pbmCellWords', 'pbmChecked',
                                        'pbmPut', 'pbmPasteInto')) + '''
pbmPasteInto({ dataset: { pbmrow: '0', pbmf: 'benchmark' } },
  'Benchmark\\tCurrency\\tRisk\\tHedging\\nA\\tUSD\\tModerate\\tUnhedged\\nX\\tUSD\\tModerate\\tEquity Not Hedged\\nB\\tGBP\\tModerate\\tHedged');
var pasted = pbm.draft.map(function (r) { return r.benchmark; });
pbm.draft.push({ benchmark: 'Y', currency: 'USD', risk: 'Moderate', hedging: 'Equity Not Hedged', mnemonic: '' });
process.stdout.write(JSON.stringify({ pasted: pasted, note: pbm.pasted, bad: pbmChecked()[2].bad }));
''')
    got = json.loads(subprocess.run([node, '-e', script], capture_output=True, text=True, check=True).stdout)
    assert got['pasted'] == ['A', 'B']
    assert got['note'].endswith('1 row for Equity Not Hedged left out: it takes the Unhedged benchmark.')
    assert got['bad'] == {'hedging': 'Equity Not Hedged takes the Unhedged benchmark: delete the row'}


@pytest.mark.parametrize('row, words', [
    (_row('', 'USD', 'Moderate', 'Hedged'), 'no Benchmark'),
    (_row('B', 'JPY', 'Moderate', 'Hedged'), 'Currency "JPY" is not one on offer'),
    (_row('B', 'USD', 'Spicy', 'Hedged'), 'Risk "Spicy" is not one on offer'),
    (_row('B', 'USD', 'Moderate', ''), 'no Hedging'),
    (_row('X' * 201, 'USD', 'Moderate', 'Hedged'), 'keep the benchmark to 200 characters'),
    (_row('B', 'USD', 'Moderate', 'Hedged', 'M' * 41), 'keep the mnemonic to 40 characters'),
])
def test_a_row_that_is_not_what_it_should_be_is_refused(store, tmp_path, row, words):
    with pytest.raises(ValidationError, match=words):
        pb.save({'rows': [row]}, user='alice')
    preview = pb.preview('b.csv', _upload(tmp_path, [HEAD, [row[f] for f in pb.FIELDS]]))
    assert any(words in p for p in preview['problems']), preview['problems']
    assert pb.current()['revision'] == 0


def test_two_rows_for_one_portfolio_are_named(store, tmp_path):
    rows = [_row('A', 'USD', 'Moderate', 'Hedged'), _row('B', 'USD', 'Moderate', 'Hedged')]
    with pytest.raises(ValidationError, match='Row 1 and row 2 are both for Moderate · USD · Hedged'):
        pb.save({'rows': rows}, user='alice')
    problems = pb.preview('b.csv', _upload(tmp_path, [HEAD] + [[r[f] for f in pb.FIELDS] for r in rows]))['problems']
    assert problems == ['Row 2 and row 3 are both for Moderate · USD · Hedged.'], 'the file\'s own row numbers'


def test_a_file_seeds_the_whole_table_and_an_export_loads_back(store, tmp_path):
    rows = [['Mnemonic', 'Hedging', 'Risk Level', 'Currency', 'Benchmark Name'],
            ['BM-1', 'Hedged', 'Low Vol', 'USD', 'Low Vol USD'],
            ['', 'Unhedged', 'Aggressive', 'gbp', 'Agg GBP'],
            ['', '', '', '', '']]
    preview = pb.preview('b.csv', _upload(tmp_path, rows))
    assert preview['problems'] == [] and preview['rows'] == 2 and preview['changed'] == 2
    assert preview['gaps'] == len(pb.cells()) - 2
    pb.importFile('b.csv', _upload(tmp_path, rows), user='alice', base=preview['base'])
    assert pb.entryFor('LowVol|USD|Hedged') == {'benchmark': 'Low Vol USD', 'mnemonic': 'BM-1'}
    assert pb.history()[0]['note'] == 'Loaded b.csv'
    with pytest.raises(ValidationError, match='missing column'):
        pb.preview('b.csv', _upload(tmp_path, [['Benchmark', 'Currency'], ['X', 'USD']]))
    pb.save(_full(), user='alice', base=1)
    assert pb.gaps() == []
    exported = pb.exportRows()
    assert exported[0] == HEAD and len(exported) == len(pb.cells()) + 1
    assert exported[1] == ['MAP LowVol USD Hedged', 'USD', 'Low Vol', 'Hedged', 'M-USD']
    again = pb.preview('b.csv', _upload(tmp_path, exported))
    assert again['problems'] == [] and again['changed'] == 0, 'an export loads back with nothing to do'
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(pb.exportWorkbook()))
    assert book.sheetnames == ['MAP Benchmarks'] and book.active.max_row == len(pb.cells()) + 1


def test_a_change_a_restore_and_a_table_from_before(store):
    pb.save({'rows': [_row('A', 'USD', 'Moderate', 'Hedged', 'MA')]}, user='alice', base=0)
    with pytest.raises(ValidationError, match='already in force'):
        pb.save({'rows': [_row('A', 'USD', 'Moderate', 'Hedged', 'MA')]}, user='alice')
    with pytest.raises(sleeveRepo.StaleError):
        pb.save({'rows': []}, user='bob', base=0)
    pb.save({'rows': [_row('B', 'USD', 'Moderate', 'Hedged', 'MB')]}, user='bob', base=1)
    assert pb.history()[0]['changes'] == ['Moderate · USD · Hedged: A → B', 'Moderate · USD · Hedged: mnemonic MA → MB']
    pb.save({'rows': []}, user='bob', base=2)
    assert pb.history()[0]['changes'] == ['Moderate · USD · Hedged: removed B (MB)']
    pb.revert(1, user='bob', base=3)
    assert pb.benchmarkFor('Moderate|USD|Hedged') == 'A'
    conn = sleeveRepo.connection()
    conn.execute("UPDATE policies SET body = ? WHERE kind = 'portfolioBenchmarks'",
                 ('{"fixedIncome":{"USD":"Agg"},"weights":{"Moderate":40}}',))
    conn.commit()
    conn.close()
    assert pb.current()['rows'] == [], 'a table in another shape reads as no rows'


def test_the_view_is_the_admins_and_answers_over_its_routes(store, tmp_path):
    got = dashboardRouter.getPortfolioBenchmarks(caller=_caller('alice'))
    assert got['revision'] == 0 and got['cells'] == len(pb.cells()) and got['portfolios'] == pb.cells()
    assert got['mnemonicMax'] == 40 and got['risks'][0] == {'risk': 'LowVol', 'label': 'Low Vol'}
    saved = dashboardRouter.savePortfolioBenchmarks(dict(_full(), base=0), caller=_caller('alice'))
    assert saved['portfolioBenchmarks']['revision'] == 1 and saved['portfolioBenchmarks']['gaps'] == []
    assert dashboardRouter.savePortfolioBenchmarks(dict(_full(), base=0), caller=_caller('alice')).status_code == 409
    assert dashboardRouter.savePortfolioBenchmarks({'rows': [_row('', 'USD', 'Moderate', 'Hedged')], 'base': 1},
                                                   caller=_caller('alice')).status_code == 422
    assert dashboardRouter.getPortfolioBenchmarksHistory(caller=_caller('alice'))['history'][0]['revision'] == 1
    rows = [HEAD, ['W', 'GBP', 'Low Vol', 'Unhedged', '']]
    preview = dashboardRouter.previewPortfolioBenchmarks({'name': 'b.csv', 'content': _upload(tmp_path, rows)},
                                                         caller=_caller('alice'))
    # every other row removed, and the one kept changes its benchmark and its mnemonic
    assert preview['changed'] == len(pb.cells()) - 1 + 2
    done = dashboardRouter.importPortfolioBenchmarks({'name': 'b.csv', 'content': _upload(tmp_path, rows), 'base': 1},
                                                     caller=_caller('alice'))
    assert [r['benchmark'] for r in done['portfolioBenchmarks']['rows']] == ['W']
    back = dashboardRouter.revertPortfolioBenchmarks({'revision': 1, 'base': 2}, caller=_caller('alice'))
    assert back['portfolioBenchmarks']['revision'] == 3 and back['portfolioBenchmarks']['gaps'] == []
    assert dashboardRouter.exportPortfolioBenchmarks(caller=_caller('alice')).body[:2] == b'PK'
    from cyrus_pmg.pmgService.core.accessControl import requireAdmin
    seen = 0
    for route in dashboardRouter.router.routes:
        if route.path.startswith('/scenario/repository/portfolio-benchmarks'):
            seen += 1
            assert requireAdmin in [d.call for d in route.dependant.dependencies], route.path
    assert seen == 7


# ------------------------------------------------------- the delivery ---

def _files(response):
    bundle = zipfile.ZipFile(io.BytesIO(response.content if hasattr(response, 'content') else response.body))
    from openpyxl import load_workbook
    book = load_workbook(io.BytesIO(bundle.read(next(n for n in bundle.namelist() if n.endswith('.xlsx')))))
    deck = bundle.read(next(n for n in bundle.namelist() if n.endswith('.pptx')))
    impl = book['Implementation']
    preamble = {impl.cell(row=r, column=1).value: impl.cell(row=r, column=2).value for r in range(1, 10)}
    from pptx import Presentation
    words = ' '.join(shape.text_frame.text for slide in Presentation(io.BytesIO(deck)).slides
                     for shape in slide.shapes if shape.has_text_frame)
    return preamble, words


def test_a_strategic_proposal_names_its_portfolios_benchmark_and_records_it(store, tmp_path):
    from test_entities import MA, USD_MOD_FULL, _scenario
    done = dashboardRouter.exportScenario(_scenario('', MA, USD_MOD_FULL), caller=_caller())
    assert done.status_code == 200, done.body
    preamble, _ = _files(done)
    assert 'Benchmark' not in preamble, 'none in the table: the files name none'
    assert proposalRegister.getProposal(done.headers['X-Proposal-Id'])['benchmark'] is None
    pb.save(_full(), user='alice')
    done = dashboardRouter.exportScenario(_scenario('', MA, USD_MOD_FULL), caller=_caller())
    assert done.status_code == 200, done.body
    preamble, deck = _files(done)
    assert preamble['Benchmark'] == 'MAP Moderate USD Hedged'
    assert 'Benchmark MAP Moderate USD Hedged' in deck, 'the Implemented Model slide names it'
    kept = proposalRegister.getProposal(done.headers['X-Proposal-Id'])
    assert kept['benchmark'] == 'MAP Moderate USD Hedged'
    from cyrus_pmg.pmgService.scenario import accountRequests
    assert kept['benchmarkMnemonic'] == 'M-USD'
    assert accountRequests.mandateBenchmark(kept) == ('MAP Moderate USD Hedged', 'M-USD', 'proposal')


def test_a_custom_portfolio_chooses_its_benchmark_from_the_table(store, monkeypatch, tmp_path):
    from test_custom_portfolio import VARIANT, _asBuild, _client, _put, _scenario as _customScenario, _strategic
    client, pwa, admin, _ = _client(monkeypatch)
    sid = _customScenario(client, pwa)
    build = _asBuild(_strategic()[1])
    assert _put(client, admin, sid, build).status_code == 200
    listed = client.get('/api/v1/scenario/custom-portfolio', headers=pwa, params={'variant': VARIANT}).json()
    assert listed['benchmarks'] == []
    refused = _put(client, pwa, sid, dict(build, benchmark='Made up'))
    assert refused.status_code == 422 and refused.json()['field'] == 'benchmark'
    pb.save(_full(), user='alice')
    listed = client.get('/api/v1/scenario/custom-portfolio', headers=pwa, params={'variant': VARIANT}).json()
    assert listed['benchmarks'] == pb.choices() and len(listed['benchmarks']) > 20
    # anyone who may edit the proposal chooses it; the build is not changed by it
    chosen = _put(client, pwa, sid, dict(build, benchmark='MAP Agg GBP Unhedged'))
    assert chosen.status_code == 200, chosen.text
    held = chosen.json()['customPortfolio']
    assert held['benchmark'] == 'MAP Agg GBP Unhedged' and held['builtBy'] == 'alice'
    # a save of the build without it keeps it
    assert _put(client, admin, sid, build).json()['customPortfolio']['benchmark'] == 'MAP Agg GBP Unhedged'
    client.put('/api/v1/scenario/' + sid, headers=pwa, json={'includeFees': False})
    client.put('/api/v1/scenario/' + sid, headers=pwa,
               json={'noticeAction': {'notice': 'n1', 'action': 'confirm', 'ticked': True}})
    delivered = client.post('/api/v1/scenario/' + sid + '/export', headers=pwa)
    assert delivered.status_code == 200, delivered.text
    preamble, _ = _files(delivered)
    assert preamble['Portfolio'] == 'Custom, Moderate risk profile'
    assert preamble['Benchmark'] == 'MAP Agg GBP Unhedged'
    kept = proposalRegister.getProposal(delivered.headers['x-proposal-id'])
    assert (kept['benchmark'], kept['benchmarkMnemonic']) == ('MAP Agg GBP Unhedged', 'M-GBP')
    # cleared, none is named
    assert _put(client, pwa, sid, dict(build, benchmark='')).json()['customPortfolio']['benchmark'] == ''
