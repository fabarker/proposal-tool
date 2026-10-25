"""The MAP benchmarks: the benchmark of a multi-asset portfolio and its
mnemonic, by the portfolio's risk level, base currency and hedging policy
(D211).

A ROW is a benchmark, its mnemonic, and the currency, risk level and hedging
policy it is the benchmark for. A PORTFOLIO is a risk level, a currency and a
hedging policy - the allocation and the implementation type play no part -
and has at most one row. The risk levels and currencies are the ones the
strategic universe holds a portfolio at, each under every hedging policy. A
portfolio on offer with no row is a GAP: the console counts them, its
proposal is delivered without a benchmark, and an account opening request
for it cannot be submitted until the gap is filled.

What reads it. A strategic proposal's benchmark is its portfolio's row: the
Implementation sheet's first rows and the deck's Implemented Model slide name
it, and the register records it with its mnemonic; the account opening
request names the benchmark recorded for the mandate's own row. A custom
portfolio (D171) has no portfolio of its own: the person preparing it
chooses its benchmark from the ones the table holds.

The admin keeps the table in the console's Account Opening tab, MAP
Benchmarks: rows edited, added and deleted on screen, a block pasted from
Excel, or a file loaded whole - Benchmark, Currency, Risk, Hedging and
Mnemonic. Every change is a revision with who, when, an optional note and
what moved, and an earlier revision can be put back.

Storage. The ``policies`` and ``policyHistory`` tables (sleeveRepo, schema 4)
under the kind 'portfolioBenchmarks' and the one scope '*', as {rows:
[{currency, risk, hedging, benchmark, mnemonic}]} in the table's order. A
revision from before, in another shape, reads as no rows. Until the first
save there is no table. A read never writes.
"""

from __future__ import annotations

import datetime
import io
import json
import os
import threading
from functools import lru_cache

from . import saaKeys, sleeveRepo
from .types import PortfolioKey, ValidationError

KIND = 'portfolioBenchmarks'
SCOPE = '*'
BENCHMARK_MAX = 200
MNEMONIC_MAX = 40
NOTE_MAX = 240
MAX_ROWS = 2000
#: changes listed for one revision before the rest are counted
CHANGES_SHOWN = 40
CURRENCIES = saaKeys.CURRENCIES
#: the file's columns, matched by name whatever the case or spacing
COLUMNS = ('Benchmark', 'Currency', 'Risk', 'Hedging', 'Mnemonic')
_HEADS = {'benchmark': 'benchmark', 'benchmarkname': 'benchmark', 'name': 'benchmark',
          'currency': 'currency', 'basecurrency': 'currency', 'ccy': 'currency',
          'risk': 'risk', 'risklevel': 'risk', 'riskprofile': 'risk',
          'hedging': 'hedging', 'hedgingpolicy': 'hedging', 'currencyhedgingpolicy': 'hedging',
          'mnemonic': 'mnemonic', 'benchmarkmnemonic': 'mnemonic', 'ticker': 'mnemonic', 'code': 'mnemonic'}
_REQUIRED = ('benchmark', 'currency', 'risk', 'hedging')
FIELDS = ('benchmark', 'currency', 'risk', 'hedging', 'mnemonic')


def _now() -> str:
    return datetime.datetime.now().isoformat(timespec='seconds')


def _squash(text) -> str:
    return ''.join(ch for ch in str(text or '').lower() if ch.isalnum())


def _text(value) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return ' '.join(str(value if value is not None else '').split())


# ------------------------------------------------------------- the cells ---

def hedgings() -> list:
    from . import rules
    return list(rules.HEDGING_POLICIES)


def risks() -> list:
    return list(saaKeys.RISK_LEVELS)


def riskLabel(risk) -> str:
    from .rules import RISK_LEVEL_LABELS
    return RISK_LEVEL_LABELS.get(risk, risk)


def cell(risk, currency, hedging) -> str:
    return '|'.join((risk, currency, hedging))


def cellWords(key) -> str:
    risk, currency, hedging = key.split('|')
    return '{} · {} · {}'.format(riskLabel(risk), currency, hedging)


@lru_cache(maxsize=None)
def _offered() -> tuple:
    """(risk, currency) for every risk level and currency the strategic
    universe holds a portfolio at, least risky first, then by currency."""
    from . import universe
    held = {(key.riskLevel, key.currency) for key in universe.keys()}
    order = {r: i for i, r in enumerate(risks())}
    return tuple(sorted(held, key=lambda rc: (order.get(rc[0], 99), CURRENCIES.index(rc[1])
                                              if rc[1] in CURRENCIES else 99)))


def offeredRisks() -> list:
    out = []
    for risk, _ in _offered():
        if risk not in out:
            out.append(risk)
    return out


def offeredCurrencies() -> list:
    held = {c for _, c in _offered()}
    return [c for c in CURRENCIES if c in held]


def cells() -> list:
    """Every portfolio on offer, in the table's order: currency, risk,
    hedging."""
    order = {r: i for i, r in enumerate(risks())}
    return [cell(r, c, h) for r, c in sorted(_offered(), key=lambda rc: (CURRENCIES.index(rc[1]), order[rc[0]]))
            for h in hedgings()]


def keyOf(keyStr, hedging) -> str:
    """The portfolio a strategic proposal reads, by its key and its hedging
    policy."""
    key = PortfolioKey.fromStr(keyStr) if isinstance(keyStr, str) else keyStr
    return cell(key.riskLevel, key.currency, hedging)


def rowKey(row) -> str:
    return cell(row['risk'], row['currency'], row['hedging'])


# --------------------------------------------------------------- storage ---

def _row(conn):
    return conn.execute('SELECT * FROM policies WHERE kind = ? AND scope = ?',
                        (KIND, SCOPE)).fetchone()


def _parse(text) -> dict:
    """A stored body; one in another shape reads as no rows."""
    body = json.loads(text) if isinstance(text, str) else (text or {})
    rows = body.get('rows') if isinstance(body.get('rows'), list) else []
    return {'rows': [{f: str((r or {}).get(f) or '') for f in FIELDS} for r in rows if isinstance(r, dict)]}


_lock = threading.Lock()
_cache = None            # (file stamp, the row as read, or None for none)


def _stored(path):
    """The row in force, kept against the file's stamp: every delivery
    reads it."""
    global _cache
    stamp = sleeveRepo.fileStamp(path)
    if stamp is None:
        return None
    with _lock:
        if _cache is not None and _cache[0] == stamp:
            return _cache[1]
    found = None
    conn = sleeveRepo.readOnly(path)
    if conn is not None:
        try:
            row = _row(conn)
            found = ({k: row[k] for k in ('body', 'revision', 'updatedBy', 'updatedAt')}
                     if row is not None else None)
        except Exception:                      # noqa: BLE001 - an older file: nothing yet
            found = None
        finally:
            conn.close()
    with _lock:
        _cache = (stamp, found)
    return found


def current() -> dict:
    """The table in force: {revision, rows, updatedBy, updatedAt}, a fresh
    copy on every call. Revision 0 is no table. A read never writes."""
    path = sleeveRepo.dbPath()
    row = _stored(path) if os.path.exists(path) else None
    if row is None:
        return {'rows': [], 'revision': 0, 'updatedBy': '', 'updatedAt': ''}
    return dict(_parse(row['body']), revision=row['revision'], updatedBy=row['updatedBy'],
                updatedAt=row['updatedAt'])


def _byCell(table) -> dict:
    return {rowKey(r): r for r in table['rows']}


def entryFor(cellKey, table=None):
    """{benchmark, mnemonic} of a portfolio, or None where it has no row."""
    table = current() if table is None else table
    found = _byCell(table).get(cellKey)
    return {'benchmark': found['benchmark'], 'mnemonic': found['mnemonic']} if found else None


def benchmarkFor(cellKey, table=None) -> str:
    """The benchmark of a portfolio, blank where it has none."""
    return (entryFor(cellKey, table) or {}).get('benchmark') or ''


def forPortfolio(keyStr, hedging, table=None) -> str:
    """A strategic portfolio's benchmark, by its key and hedging policy;
    blank where it has none."""
    if not keyStr:
        return ''
    return benchmarkFor(keyOf(keyStr, hedging), table)


def entryForPortfolio(keyStr, hedging, table=None):
    if not keyStr:
        return None
    return entryFor(keyOf(keyStr, hedging), table)


def mnemonicOf(benchmark, table=None) -> str:
    """The mnemonic the table gives a benchmark by its name, blank where it
    holds none: a custom portfolio's, chosen by name."""
    table = current() if table is None else table
    for r in table['rows']:
        if r['benchmark'] == benchmark and r['mnemonic']:
            return r['mnemonic']
    return ''


def choices(table=None) -> list:
    """The benchmarks the table holds, each once, in alphabetical order: what
    a custom portfolio chooses from."""
    table = current() if table is None else table
    return sorted({r['benchmark'] for r in table['rows'] if r['benchmark']}, key=str.lower)


def gaps(table=None) -> list:
    """The portfolios on offer with no row."""
    held = _byCell(current() if table is None else table)
    return [c for c in cells() if c not in held]


# ------------------------------------------------------------ the table ---

def _pick(raw, options, labels=None):
    """The option *raw* names - by its name, its label, or a name it is the
    start of when only one is - or None."""
    said = _squash(raw)
    if not said:
        return None
    for o in options:
        if said == _squash(o) or (labels and said == _squash(labels.get(o, ''))):
            return o
    starts = [o for o in options if _squash(o).startswith(said)]
    return starts[0] if len(starts) == 1 else None


def _clean(raw, where) -> tuple:
    """(row, [problems]) for one row as given: each part named, the
    currency, risk level and hedging policy ones on offer, the benchmark and
    mnemonic within their lengths."""
    from .rules import RISK_LEVEL_LABELS
    raw = raw or {}
    problems = []
    currency = _pick(raw.get('currency'), offeredCurrencies())
    risk = _pick(raw.get('risk'), offeredRisks(), RISK_LEVEL_LABELS)
    hedging = _pick(raw.get('hedging'), hedgings())
    benchmark, mnemonic = _text(raw.get('benchmark')), _text(raw.get('mnemonic'))
    for word, said, got in (('Currency', raw.get('currency'), currency), ('Risk', raw.get('risk'), risk),
                            ('Hedging', raw.get('hedging'), hedging)):
        if got is None:
            problems.append('{}: {}.'.format(where, 'no {}'.format(word) if not _text(said)
                                             else '{} "{}" is not one on offer'.format(word, _text(said))))
    if not benchmark:
        problems.append('{}: no Benchmark.'.format(where))
    elif len(benchmark) > BENCHMARK_MAX:
        problems.append('{}: keep the benchmark to {} characters.'.format(where, BENCHMARK_MAX))
    if len(mnemonic) > MNEMONIC_MAX:
        problems.append('{}: keep the mnemonic to {} characters.'.format(where, MNEMONIC_MAX))
    return ({'benchmark': benchmark, 'currency': currency or '', 'risk': risk or '', 'hedging': hedging or '',
             'mnemonic': mnemonic}, problems)


def _check(rows, where) -> tuple:
    """(rows in the table's order, [problems]): every row clean, and no two
    for one portfolio. Rows wholly blank are left out."""
    if not isinstance(rows, list):
        raise ValidationError('benchmarks', 'Send the benchmarks as a list of rows.')
    if len(rows) > MAX_ROWS:
        raise ValidationError('benchmarks', 'At most {} rows.'.format(MAX_ROWS))
    out, problems, seen = [], [], {}
    for n, raw in enumerate(rows, start=1):
        if not isinstance(raw, dict):
            raise ValidationError('benchmarks', 'Each row is a benchmark, a currency, a risk level, a hedging '
                                                'policy and a mnemonic.')
        if not any(_text(raw.get(f)) for f in FIELDS):
            continue
        row, bad = _clean(raw, where(n))
        problems.extend(bad)
        if bad:
            continue
        key = rowKey(row)
        if key in seen:
            problems.append('{} and {} are both for {}.'.format(where(seen[key]), where(n).lower(), cellWords(key)))
            continue
        seen[key] = n
        out.append(row)
    order = {c: i for i, c in enumerate(cells())}
    return sorted(out, key=lambda r: order[rowKey(r)]), problems


def normalise(body) -> dict:
    """The canonical table, or a ValidationError naming ``benchmarks``."""
    if not isinstance(body, dict):
        raise ValidationError('benchmarks', 'Send the benchmarks as a list of rows.')
    rows, problems = _check(body.get('rows'), lambda n: 'Row {}'.format(n))
    if problems:
        raise ValidationError('benchmarks', ' '.join(problems[:6]) + (
            '' if len(problems) <= 6 else ' And {} more.'.format(len(problems) - 6)))
    return {'rows': rows}


# ------------------------------------------------------------- the file ---

def _fileName(name) -> str:
    name = os.path.basename(str(name or '').replace('\\', '/')).strip()
    if not name.lower().endswith(('.csv', '.xlsx')):
        raise ValidationError('file', 'A benchmarks file is a .csv or an .xlsx.')
    return name


def _readSheet(name, data):
    """(header, rows) of the sheet that holds the table: a CSV's only one, or
    the first sheet of a workbook whose header names every required column."""
    if name.lower().endswith('.csv'):
        import csv
        from .csvText import unwrapRows
        try:
            text = data.decode('utf-8-sig')
        except UnicodeDecodeError:
            text = data.decode('cp1252')
        rows = list(csv.reader(unwrapRows(io.StringIO(text, newline=''))))
        return (rows[0] if rows else []), rows[1:]
    from openpyxl import load_workbook
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:                          # noqa: BLE001 - not a workbook
        raise ValidationError('file', '{} could not be read as a workbook.'.format(name))
    for sheet in book.worksheets:
        rows = sheet.iter_rows(values_only=True)
        head = next(rows, None) or []
        heads = {_HEADS.get(_squash(h)) for h in head}
        if all(r in heads for r in _REQUIRED):
            return ([('' if h is None else str(h)) for h in head],
                    [['' if v is None else v for v in r] for r in rows])
    raise ValidationError('file', 'No sheet in {} has the columns {}.'.format(name, ', '.join(COLUMNS[:4])))


def readFile(name, content) -> dict:
    """The table a file holds, with what is wrong with it: {table, rows,
    problems: [words]}."""
    from . import products
    name = _fileName(name)
    head, rows = _readSheet(name, products._decode(content))
    at = {}
    for i, h in enumerate(head):
        field = _HEADS.get(_squash(h))
        if field is not None:
            at.setdefault(field, i)
    missing = [c for c, f in zip(COLUMNS, FIELDS) if f in _REQUIRED and f not in at]
    if missing:
        raise ValidationError('file', '{}: missing column(s) {}.'.format(name, ', '.join(missing)))
    given = [{f: (r[at[f]] if f in at and at[f] < len(r) else '') for f in FIELDS} for r in rows]
    # the file's own row numbers, the header being row 1
    checked, problems = _check(given, lambda n: 'Row {}'.format(n + 1))
    count = sum(1 for g in given if any(_text(v) for v in g.values()))
    if count == 0:
        problems.append('{} holds no rows.'.format(name))
    return {'name': name, 'table': {'rows': checked}, 'rows': count, 'problems': problems}


def _words(row) -> str:
    return row['benchmark'] + (' ({})'.format(row['mnemonic']) if row['mnemonic'] else '')


def _changes(before, after) -> list:
    """What moved between two tables, in the words the console shows."""
    was, now = _byCell(before), _byCell(after)
    out = []
    for c in cells():
        a, b = was.get(c), now.get(c)
        if a == b:
            continue
        if a is None:
            out.append('{}: added {}'.format(cellWords(c), _words(b)))
        elif b is None:
            out.append('{}: removed {}'.format(cellWords(c), _words(a)))
        else:
            if a['benchmark'] != b['benchmark']:
                out.append('{}: {} → {}'.format(cellWords(c), a['benchmark'], b['benchmark']))
            if a['mnemonic'] != b['mnemonic']:
                out.append('{}: mnemonic {} → {}'.format(cellWords(c), a['mnemonic'] or '—', b['mnemonic'] or '—'))
    return out


def _summary(changes) -> list:
    if len(changes) <= CHANGES_SHOWN:
        return changes
    return changes[:CHANGES_SHOWN] + ['and {} more'.format(len(changes) - CHANGES_SHOWN)]


def preview(name, content) -> dict:
    """What a file would do, before anything moves."""
    read = readFile(name, content)
    now = current()
    changes = _changes(now, read['table'])
    return {'name': read['name'], 'rows': read['rows'], 'problems': read['problems'],
            'cells': len(read['table']['rows']), 'changes': _summary(changes), 'changed': len(changes),
            'gaps': len(gaps(read['table'])), 'base': now['revision']}


# ------------------------------------------------------------ the record ---

def _note(note) -> str:
    if note is not None and not isinstance(note, str):
        raise ValidationError('note', 'A note is words.')
    note = ' '.join((note or '').split())
    if len(note) > NOTE_MAX:
        raise ValidationError('note', 'Keep the note to {} characters.'.format(NOTE_MAX))
    return note


def _write(conn, body, action, note, user) -> None:
    top = conn.execute('SELECT COALESCE(MAX(revision), 0) FROM policyHistory WHERE kind = ? '
                       'AND scope = ?', (KIND, SCOPE)).fetchone()[0]
    revision = top + 1
    text = json.dumps(body, separators=(',', ':'))
    stamp = _now()
    conn.execute('INSERT OR REPLACE INTO policies (kind, scope, body, revision, updatedBy, '
                 'updatedAt) VALUES (?, ?, ?, ?, ?, ?)',
                 (KIND, SCOPE, text, revision, user or '', stamp))
    conn.execute('INSERT INTO policyHistory (kind, scope, revision, action, body, note, actor, at) '
                 'VALUES (?, ?, ?, ?, ?, ?, ?, ?)',
                 (KIND, SCOPE, revision, action, text, note, user or '', stamp))


def _refuseStale(conn, base) -> None:
    if base is None:
        return
    try:
        base = int(base)
    except (TypeError, ValueError):
        raise ValidationError('base', 'The revision the change was made against is a number.')
    row = _row(conn)
    now = row['revision'] if row is not None else 0
    if now != base:
        raise sleeveRepo.StaleError(
            'base', 'The MAP benchmarks were changed by {} since you opened them. Reload to see the '
            'change.'.format((row['updatedBy'] if row is not None else '') or 'someone'))


def _previous(conn) -> dict:
    existing = _row(conn)
    return _parse(existing['body']) if existing is not None else {'rows': []}


def _put(cleaned, action, note, user, base) -> dict:
    conn = sleeveRepo.connection()
    try:
        conn.commit()
        conn.execute('BEGIN IMMEDIATE')
        try:
            _refuseStale(conn, base)
            if cleaned['rows'] == _previous(conn)['rows']:
                raise ValidationError('benchmarks', 'Those are the MAP benchmarks already in force.')
            _write(conn, cleaned, action, note, user)
            conn.commit()
        except Exception:
            conn.rollback()
            raise
        return current()
    finally:
        conn.close()


def save(body, note='', user='', base=None) -> dict:
    """Put a table edited on screen in force. Refused, naming the field, when
    a row is not what it should be, someone saved in between (*base*) or
    nothing would change."""
    return _put(normalise(body), 'updated', _note(note), user, base)


def importFile(name, content, note='', user='', base=None) -> dict:
    """Put a file's table in force, whole: a portfolio the file leaves out
    loses its row, and it stays in the history. Refused while the file has
    a problem."""
    note = _note(note)
    read = readFile(name, content)
    if read['problems']:
        raise ValidationError('file', '{}: {}'.format(read['name'], ' '.join(read['problems'][:6])
                                                       + ('' if len(read['problems']) <= 6 else
                                                          ' And {} more.'.format(len(read['problems']) - 6))))
    return _put(read['table'], 'imported', note or 'Loaded {}'.format(read['name']), user, base)


def _rows(conn) -> list:
    return conn.execute('SELECT * FROM policyHistory WHERE kind = ? AND scope = ? '
                        'ORDER BY revision', (KIND, SCOPE)).fetchall()


def history() -> list:
    """Every revision, newest first, each with what it changed."""
    conn = sleeveRepo.connection()
    try:
        rows = _rows(conn)
    finally:
        conn.close()
    out, previous = [], {'rows': []}
    top = rows[-1]['revision'] if rows else None
    for r in rows:
        body = _parse(r['body'])
        changes = _changes(previous, body)
        out.append({'revision': r['revision'], 'action': r['action'], 'note': r['note'],
                    'actor': r['actor'], 'at': r['at'], 'changes': _summary(changes),
                    'changed': len(changes), 'cells': len(body['rows']),
                    'current': r['revision'] == top})
        previous = body
    return list(reversed(out))


def feedRows(conn) -> list:
    """The record as the Activity feed reads it."""
    out, previous = [], {'rows': []}
    for r in _rows(conn):
        body = _parse(r['body'])
        out.append({'id': r['id'], 'revision': r['revision'], 'action': r['action'], 'note': r['note'],
                    'actor': r['actor'], 'at': r['at'], 'changes': _summary(_changes(previous, body))})
        previous = body
    return out


def revert(revision, note='', user='', base=None) -> dict:
    """Put an earlier revision's table back in force, as a new revision. A
    row for a portfolio no longer on offer is left out."""
    note = _note(note)
    try:
        number = int(revision)
    except (TypeError, ValueError):
        raise ValidationError('revision', 'Name the revision by its number.')
    conn = sleeveRepo.connection()
    try:
        row = conn.execute('SELECT * FROM policyHistory WHERE kind = ? AND scope = ? AND '
                           'revision = ?', (KIND, SCOPE, number)).fetchone()
    finally:
        conn.close()
    if row is None:
        raise ValidationError('revision', 'There is no revision {} of the MAP benchmarks.'.format(number))
    offered = set(cells())
    body = _parse(row['body'])
    return _put(normalise({'rows': [r for r in body['rows'] if rowKey(r) in offered]}), 'reverted',
                note or 'Restored revision {}'.format(number), user, base)


def exportRows() -> list:
    """The table in force as the file it loads from."""
    return [list(COLUMNS)] + [[r['benchmark'], r['currency'], riskLabel(r['risk']), r['hedging'], r['mnemonic']]
                              for r in current()['rows']]


def exportWorkbook() -> bytes:
    """``exportRows`` as an .xlsx with one sheet, MAP Benchmarks, ready to
    edit and load again."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    book = Workbook()
    sheet = book.active
    sheet.title = 'MAP Benchmarks'
    for row in exportRows():
        sheet.append(row)
    for c in sheet[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='0B2545')
    for letter, width in zip('ABCDE', (70, 10, 22, 18, 18)):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


def describe() -> dict:
    """Everything the console's MAP Benchmarks view shows."""
    table = current()
    return dict(table, currencies=offeredCurrencies(), hedgings=hedgings(),
                risks=[{'risk': r, 'label': riskLabel(r)} for r in offeredRisks()],
                gaps=gaps(table), cells=len(cells()), portfolios=cells(), benchmarkMax=BENCHMARK_MAX,
                mnemonicMax=MNEMONIC_MAX, noteMax=NOTE_MAX, maxRows=MAX_ROWS)
