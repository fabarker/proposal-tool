#!/usr/bin/env python3
"""The sleeve seed workbook (D159): a template to seed the library from
scratch, and the live library exported in the same shape.

Run:  python3 build_sleeve_template.py [--db PATH]

Writes, beside the seed in sleeveSource/:
  sleeve-seed-template.xlsx   empty Sleeves and Rules sheets, ready to fill
  sleeve-seed-library.xlsx    the library in --db (default: the service's
                              var/sleeves.db, opened read-only) as a worked
                              example that loads back to the same library

Both carry the sheets the loader reads by name - Sleeves (SEED_COLUMNS_EDITIONS)
and Rules (RULE_COLUMNS) - plus Instructions, Catalogue and Lists. Weights are
stored as fractions, as the loader reads them, and shown as percentages.
"""
import argparse
import os
import shutil
import sqlite3
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SERVICE = os.path.join(ROOT, 'service')
OUT_DIR = os.path.join(ROOT, 'sleeveSource')
LIVE_DB = os.path.join(SERVICE, 'var', 'sleeves.db')
sys.path.insert(0, SERVICE)

from openpyxl import Workbook                                       # noqa: E402
from openpyxl.formatting.rule import FormulaRule                    # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation       # noqa: E402

NAVY = '16243A'
INPUT = 'FFFDF2'
BAD = 'FDE2DF'
ROWS = 2000                       # how far the drop-downs and checks reach
HEAD_FONT = Font(name='Calibri', bold=True, color='FFFFFF')
HEAD_FILL = PatternFill('solid', fgColor=NAVY)
INPUT_FILL = PatternFill('solid', fgColor=INPUT)
THIN = Side(style='thin', color='C9D2DC')


def vocabulary():
    """Every value a sheet may hold, from the code that validates it."""
    from cyrus_pmg.pmgService.scenario import products, sleeveRepo, sleeveRules
    from cyrus_pmg.pmgService.scenario.rules import RISK_LEVEL_LABELS
    from cyrus_pmg.pmgService.scenario.sleeves import VARIANTS
    vocab = sleeveRules.vocabulary()
    return {
        'variants': list(VARIANTS),
        'categories': list(sleeveRepo.categories()),
        'fixed': list(sleeveRepo.fixedCategories()),
        'currency': list(vocab['currency']),
        'riskLevel': list(vocab['riskLevel']),
        'allocationType': list(vocab['allocationType']),
        'riskLabels': dict(RISK_LEVEL_LABELS),
        'products': products.all(),
    }


def libraryRows(dbPath):
    """(sleeve rows, rule rows) of the library at *dbPath*, read from a copy
    so the live file is never opened for writing."""
    tmp = tempfile.mkdtemp()
    copy = os.path.join(tmp, 'sleeves.db')
    src = sqlite3.connect('file:{}?mode=ro'.format(dbPath), uri=True)
    dst = sqlite3.connect(copy)
    src.backup(dst)
    src.close()
    dst.close()
    os.environ['SCENARIO_SLEEVES_DB'] = copy
    from cyrus_pmg.pmgService.scenario import sleeveRepo
    try:
        return sleeveRepo.exportRows(), sleeveRepo.exportRuleRows()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def _head(ws, names, widths):
    for i, (name, width) in enumerate(zip(names, widths), start=1):
        c = ws.cell(row=1, column=i, value=name)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(vertical='center')
        ws.column_dimensions[c.column_letter].width = width
    ws.freeze_panes = 'A2'
    ws.row_dimensions[1].height = 20


def _list(ws, col, values, title):
    ws.cell(row=1, column=col, value=title).font = Font(bold=True)
    for i, v in enumerate(values, start=2):
        ws.cell(row=i, column=col, value=v)
    letter = ws.cell(row=1, column=col).column_letter
    return "Lists!${0}$2:${0}${1}".format(letter, len(values) + 1)


def _dropdown(ws, ref, cols, prompt):
    dv = DataValidation(type='list', formula1=ref, allow_blank=True, showErrorMessage=True,
                        errorTitle='Not in the list', error=prompt)
    ws.add_data_validation(dv)
    for col in cols:
        dv.add('{0}2:{0}{1}'.format(col, ROWS))


def _inputs(ws, cols):
    for col in cols:
        for r in range(2, ROWS + 1):
            ws['{}{}'.format(col, r)].fill = INPUT_FILL


def build(path, sleeveRows=(), ruleRows=(), example=False):
    v = vocabulary()
    wb = Workbook()

    # ---- Sleeves: the sheet the loader reads first, by name
    ws = wb.active
    ws.title = 'Sleeves'
    _head(ws, ['Variant', 'Category', 'Sleeve', 'Edition', 'ProductId', 'Weight',
               'Product name', 'Edition total', 'Check'], [28, 36, 34, 16, 38, 10, 40, 13, 44])
    rows = list(sleeveRows)
    if example and not rows:
        rows = _exampleRows(v)
    for i, row in enumerate(rows, start=2):
        for j, value in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=value)
    for r in range(2, ROWS + 1):
        ws['F{}'.format(r)].number_format = '0.00%'
        ws['G{}'.format(r)] = ('=IF(E{0}="","",IFERROR(VLOOKUP(E{0},Catalogue!$A:$B,2,FALSE),'
                               '"Not in the catalogue"))').format(r)
        ws['H{}'.format(r)] = ('=IF(A{0}="","",SUMIFS($F:$F,$A:$A,A{0},$B:$B,B{0},$C:$C,C{0},'
                               '$D:$D,D{0}&""))').format(r)
        ws['H{}'.format(r)].number_format = '0.00%'
        ws['I{}'.format(r)] = (
            '=IF(A{0}="","",IF(COUNTIF(Catalogue!$A:$A,E{0})=0,"Unknown product id",'
            'IF(ABS(H{0}-1)>0.000001,"Edition weights must total exactly 100%",'
            'IF(COUNTIFS($A:$A,A{0},$B:$B,B{0},$C:$C,C{0},$D:$D,D{0}&"",$E:$E,E{0})>1,'
            '"Product listed twice in this edition","OK"))))').format(r)
    _inputs(ws, 'ABCDEF')
    red = PatternFill('solid', fgColor=BAD)
    ws.conditional_formatting.add('A2:I{}'.format(ROWS), FormulaRule(
        formula=['AND($A2<>"",$I2<>"OK")'], fill=red))
    lists = wb.create_sheet('Lists')
    refV = _list(lists, 1, v['variants'], 'Implementation type (Variant)')
    refC = _list(lists, 2, v['categories'], 'Category')
    refCur = _list(lists, 3, v['currency'], 'Currency')
    refRisk = _list(lists, 4, v['riskLevel'], 'RiskLevel (stored form)')
    for i, risk in enumerate(v['riskLevel'], start=2):
        lists.cell(row=i, column=5, value=v['riskLabels'].get(risk, risk))
    lists.cell(row=1, column=5, value='Risk level as the desk says it').font = Font(bold=True)
    refAlloc = _list(lists, 6, v['allocationType'], 'AllocationType')
    lists.cell(row=1, column=7, value='Fixed categories (one sleeve per type)').font = Font(bold=True)
    for i, c in enumerate(v['fixed'], start=2):
        lists.cell(row=i, column=7, value=c)
    for col, w in zip('ABCDEFG', [30, 36, 10, 24, 30, 16, 38]):
        lists.column_dimensions[col].width = w
    _dropdown(ws, refV, 'A', 'Choose an implementation type from the list.')
    _dropdown(ws, refC, 'B', 'Choose a category from the list.')
    _dropdown(ws, 'Catalogue!$A$2:$A${}'.format(len(v['products']) + 1), 'E',
              'Use a product id from the Catalogue sheet.')

    # ---- Rules
    wr = wb.create_sheet('Rules', 1)
    _head(wr, ['Variant', 'Category', 'Sleeve', 'Edition', 'Currency', 'RiskLevel',
               'AllocationType', 'Check'], [28, 36, 34, 16, 18, 28, 22, 46])
    rrows = list(ruleRows)
    if example and not rrows:
        rrows = _exampleRules(v)
    for i, row in enumerate(rrows, start=2):
        for j, value in enumerate(row, start=1):
            wr.cell(row=i, column=j, value=value)
    for r in range(2, ROWS + 1):
        wr['H{}'.format(r)] = (
            '=IF(A{0}="","",IF(D{0}="","An edition needs a label; the fallback has no rules",'
            'IF(COUNTIFS(Sleeves!$A:$A,A{0},Sleeves!$B:$B,B{0},Sleeves!$C:$C,C{0},'
            'Sleeves!$D:$D,D{0})=0,"No Sleeves rows for this edition",'
            'IF(AND(E{0}="",F{0}="",G{0}=""),"A rule must name at least one value","OK"))))').format(r)
    _inputs(wr, 'ABCDEFG')
    wr.conditional_formatting.add('A2:H{}'.format(ROWS), FormulaRule(
        formula=['AND($A2<>"",$H2<>"OK")'], fill=red))
    _dropdown(wr, refV, 'A', 'Choose an implementation type from the list.')
    _dropdown(wr, refC, 'B', 'Choose a category from the list.')
    # single values from the list; several are typed as Full|Core, so these
    # warn rather than refuse
    for ref, col in ((refCur, 'E'), (refRisk, 'F'), (refAlloc, 'G')):
        dv = DataValidation(type='list', formula1=ref, allow_blank=True, showErrorMessage=True,
                            errorStyle='warning', errorTitle='Several values?',
                            error='Several values are written with | between them, e.g. GBP|EUR. '
                                  'Each value must come from the Lists sheet.')
        wr.add_data_validation(dv)
        dv.add('{0}2:{0}{1}'.format(col, ROWS))

    # ---- Catalogue
    wc = wb.create_sheet('Catalogue')
    _head(wc, ['ProductId', 'Name', 'Asset class', 'Vehicle', 'Style', 'Source', 'Currency',
               'Liquidity', 'Cost', 'Minimum'], [38, 40, 20, 14, 10, 10, 10, 11, 9, 13])
    for i, p in enumerate(v['products'], start=2):
        for j, key in enumerate(['productId', 'name', 'assetClass', 'vehicle', 'style', 'source',
                                 'exposureCurrency', 'liquidity', 'productCost',
                                 'minimumInvestment'], start=1):
            wc.cell(row=i, column=j, value=p.get(key))
        wc.cell(row=i, column=9).number_format = '0.00"%"'
        wc.cell(row=i, column=10).number_format = '#,##0'

    # ---- Instructions, first in the book
    wi = wb.create_sheet('Instructions', 0)
    wi.column_dimensions['A'].width = 120
    for i, line in enumerate(_instructions(v), start=1):
        c = wi.cell(row=i, column=1, value=line)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        if line and not line.startswith(' ') and line == line.upper():
            c.font = Font(bold=True, color=NAVY, size=12)
    wi.cell(row=1, column=1).font = Font(bold=True, color=NAVY, size=15)
    for sheet in (lists, wc):
        sheet.sheet_properties.tabColor = '8494A4'
    wb.active = 0
    wb.save(path)


def _exampleRows(v):
    """Two illustrative editions of one sleeve, marked as examples."""
    ids = [p['productId'] for p in v['products']]
    a, b = ids[0], ids[1]
    var, cat = v['variants'][0], v['categories'][0]
    return [(var, cat, 'Example sleeve - replace me', '', a, 0.6),
            (var, cat, 'Example sleeve - replace me', '', b, 0.4),
            (var, cat, 'Example sleeve - replace me', 'GBP', a, 1.0)]


def _exampleRules(v):
    return [(v['variants'][0], v['categories'][0], 'Example sleeve - replace me', 'GBP',
             'GBP', '', '')]


def _instructions(v):
    risks = ', '.join('{} ({})'.format(r, v['riskLabels'].get(r, r)) for r in v['riskLevel'])
    return [
        'Sleeve seed workbook',
        'Fill the Sleeves sheet (and the Rules sheet for editions), then load the workbook into the '
        'sleeve library. Yellow cells are for you; grey sheets are reference.',
        '',
        'THE SHEETS',
        '  Sleeves - one row per product in an edition: Variant (implementation type), Category, '
        'Sleeve (its name), Edition (blank for the fallback), ProductId, Weight. The last three '
        'columns are checks, filled in for you.',
        '  Rules - which strategic portfolios each labelled edition is for. Only editions with a '
        'label have rules; the fallback never does.',
        '  Catalogue - every product id you may use, with its name, vehicle, cost and minimum.',
        '  Lists - the values the drop-downs offer.',
        '',
        'SLEEVES AND EDITIONS',
        '  A sleeve is a name in a category under an implementation type. Its fallback edition '
        '(Edition left blank) is what every portfolio gets unless a labelled edition claims it.',
        '  A labelled edition (e.g. GBP) is a variant of the same sleeve for particular portfolios. '
        'Give it its own product rows under the same Variant, Category and Sleeve, with the label '
        'in Edition, and at least one row on the Rules sheet.',
        '  Every sleeve needs its fallback. Two editions of one sleeve may not claim the same '
        'portfolio, and they may not leave the fallback with no portfolio at all.',
        '',
        'RULES',
        '  Each Rules row is one rule. A portfolio gets the edition if ANY of its rules matches '
        '(OR between rows).',
        '  Within a row, every field you fill must match (AND across Currency, RiskLevel and '
        'AllocationType). A blank field means any value.',
        '  A field may list several values with | between them, e.g. GBP|EUR, and matches any of '
        'them (OR within a field).',
        '  Currencies: {}.'.format(', '.join(v['currency'])),
        '  Risk levels, as stored (and as the desk says them): {}.'.format(risks),
        '  Allocation types: {} (NA is the all-equity portfolio).'.format(
            ', '.join(v['allocationType'])),
        '',
        'WEIGHTS',
        '  Type weights as percentages; they are stored as fractions (60% is 0.6). Each edition\'s '
        'weights must total exactly 100% - the Edition total and Check columns show it.',
        '  A product may appear once in an edition.',
        '',
        'NAMES',
        '  Names and labels are tidied when loaded: spaces at the ends removed and runs of spaces '
        'made one. Names that differ only in capitals or spacing are the same name - "active '
        'passive" and "Active Passive" cannot both exist in one category under one type.',
        '  The same name under several implementation types is one sleeve offered more widely: '
        'repeat its rows under each Variant.',
        '  "fallback" cannot be used as an edition label.',
        '',
        'FIXED CATEGORIES',
        '  {} hold exactly one sleeve per implementation type; the overlays attach it '
        'automatically. Give each type one sleeve there, with no editions.'.format(
            ' and '.join(v['fixed'])),
        '',
        'LOADING IT',
        '  1. Back up the library first: copy service/var/sleeves.db somewhere safe (with the '
        'service stopped).',
        '  2. To ADD to the library (editions of the same sleeve and label are overwritten, '
        'everything else stays):',
        '       cd service && PYTHONPATH=. python -m cyrus_pmg.pmgService.scenario.sleeveTools '
        '--import ../sleeveSource/your-file.xlsx',
        '  3. To SEED FROM SCRATCH (every existing sleeve is archived first, with its history kept):',
        '       cd service && PYTHONPATH=. python -m cyrus_pmg.pmgService.scenario.sleeveTools '
        '--import ../sleeveSource/your-file.xlsx --replace',
        '     Alternatively, for a brand new library: stop the service, move sleeves.db aside, set '
        'SCENARIO_SLEEVES_SEED to this workbook and start the service - an empty store is seeded '
        'from it once.',
        '  4. The import checks everything a save in the console checks (catalogue ids, totals, '
        'rules, overlaps) and loads nothing if any row fails; the message names the row.',
        '  5. Restart the service so open pages pick the library up.',
    ]


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('\n')[0])
    ap.add_argument('--db', default=LIVE_DB, help='library to export (opened read-only)')
    ap.add_argument('--out', default=OUT_DIR)
    args = ap.parse_args(argv)
    template = os.path.join(args.out, 'sleeve-seed-template.xlsx')
    build(template)
    print('wrote', template)
    if os.path.exists(args.db):
        rows, rules = libraryRows(args.db)
        library = os.path.join(args.out, 'sleeve-seed-library.xlsx')
        build(library, rows, rules)
        print('wrote {} ({} product rows, {} rules)'.format(library, len(rows), len(rules)))


if __name__ == '__main__':
    main()
